from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "AI证券合规风险管理与智能审阅平台_当前阶段评估报告.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "1F6B45"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"
BORDER = "D9DEE7"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    """Apply fixed DXA geometry so Word and LibreOffice agree on table layout."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER, size="6", inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right", "insideH", "insideV"]
    for name in names:
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single" if (inside or name in {"top", "left", "bottom", "right"}) else "nil")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_paragraph_border_bottom(paragraph, color=BLUE, size="12", space="6"):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def set_cell_text(cell, text, *, bold=False, color="222222", size=9.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    set_run_font(r, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, body_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        shade_cell(hdr.cells[idx], header_fill)
        set_cell_text(hdr.cells[idx], value, bold=True, color=INK, size=9.3)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            set_cell_text(cells[idx], value, size=body_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text, *, bold_prefix=None, color="222222", after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=color, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11, color="222222")
    return p


def add_note(doc, label, text, fill=CALLOUT, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="C9D2DF", size="6")
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.1
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, color=label_color, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color="25313D")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_page(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("AI证券合规风险管理与智能审阅平台")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r2 = hp.add_run("    |    当前阶段评估报告")
    set_run_font(r2, size=8.5, color=MUTED)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    r3 = fp.add_run("内部使用  |  评估日期：2026年8月31日  |  页码 ")
    set_run_font(r3, size=8.5, color=MUTED)
    add_page_field(fp)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AI证券合规风险管理与智能审阅平台")
    set_run_font(r, size=23, color="111827", bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(16)
    r2 = p2.add_run("当前阶段评估报告")
    set_run_font(r2, size=14, color=MUTED, bold=False)

    metadata = [
        ("评估对象", "C:\\Users\\PC\\Documents\\ChatGPT\\创业黑马 当前版本"),
        ("评估范围", "本地 RAG、文档解析、智能审阅、财务勾稽、账号与审计等已实现模块"),
        ("部署形态", "单机本地运行，可在同一局域网访问；当前默认模式为 local_rag"),
        ("评估日期", "2026年8月31日"),
        ("总体结论", "个人及少量用户可用的本地内测版，已具备真实应用基础，尚未达到企业正式生产标准"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        a = p.add_run(label + "：")
        set_run_font(a, size=10.5, color=INK, bold=True)
        b = p.add_run(value)
        set_run_font(b, size=10.5, color="303846")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(8)
    set_paragraph_border_bottom(rule, color=BLUE, size="12", space="5")


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    configure_header_footer(doc.sections[0])

    # Core document properties, kept generic to avoid exposing personal metadata.
    doc.core_properties.title = "AI证券合规风险管理与智能审阅平台当前阶段评估报告"
    doc.core_properties.subject = "当前能力、验证结果与提升路线"
    doc.core_properties.author = "项目评估"
    doc.core_properties.comments = "基于本地项目运行检查生成"

    add_title_block(doc)
    add_note(
        doc,
        "结论摘要",
        "当前系统已经跨过简单 Demo 阶段：本地文档解析、真实本地 Embedding、FAISS 检索、规则与资料引用、风险定位、财务勾稽和审计记录均可运行。当前更准确的定位是“个人及少量用户的本地内测版”。要进入上市公司证券部正式使用，还需要完成真实 LLM 验收、复杂文档质量验证、审核闭环、生产级部署与安全治理。",
    )

    doc.add_heading("一、当前版本定位", level=1)
    add_body(doc, "系统目前围绕“本地知识检索 + 证券材料智能审阅”构建，核心数据和向量索引保存在本机，不依赖外部 Embedding API。默认运行在 local_rag 模式，适合个人内测、敏感资料查询和审阅流程验证。")
    add_body(doc, "当前版本可以辅助发现问题、给出规则和原文出处，并保留人工复核入口；它仍不能替代董秘、证券部、法务或审计机构的最终专业判断。")

    doc.add_heading("二、当前已具备的能力", level=1)
    capability_rows = [
        ("账号与隔离", "注册、登录、密码修改、账号注销；不同账号数据隔离；管理员可管理账号和共享知识库。", "当前为小规模内部使用设计，尚未接入企业统一身份认证。"),
        ("材料接入", "支持 PDF、DOCX、XLS、XLSX、CSV、TXT、MD、JSON；支持公告、报告、制度、合同和会议材料。", "复杂扫描件、跨页表格和合并单元格仍需更多真实样本验证。"),
        ("本地解析与 OCR", "提取正文、章节、页码、段落、表格、工作表及 CSV 行号；Windows Media OCR 可用。", "OCR 质量受扫描清晰度、字体、印章和版式影响。"),
        ("本地 RAG", "BAAI/bge-m3 本地 Embedding，1024 维；FAISS 语义检索、关键词检索和混合检索并存。", "向量检索能提供证据片段，但检索质量仍需持续评估和治理。"),
        ("证券知识库", "公司资料、法规、深交所规则、证监会规则、历史公告、历史案例和上传材料统一进入共享知识库。", "资料可检索不等于已经全部结构化为可执行规则。"),
        ("风险审阅", "输出风险等级、风险类型、描述、原文位置、页码、章节、法规/制度/案例依据和处理建议。", "正式报告签署、整改闭环和多级复核仍需加强。"),
        ("财务审阅", "结构化提取、Decimal 复算、跨表勾稽、同比边界判断，并区分同比百分比与百分点变化。", "复杂表格识别和更多真实财务案例仍是主要提升方向。"),
        ("对比与溯源", "支持与上年同期、上年度报告、历史公告及摘要对比；风险可回带原文位置。", "跨版本文档的版面级差异和批注导出仍需继续打磨。"),
        ("Memory 与审计", "任务 Memory、对话 Memory、摘要 Memory；人工确认、误报、补充材料和撤销；操作审计日志。", "审核人分配、多级复核、版本冻结和正式签署能力尚不完整。"),
        ("数据更新与访问", "可同步巨潮公告、深交所规则和证监会规则；支持局域网和手机端响应式访问。", "跨不同网络访问还需要 VPN、Tailscale 或 Cloudflare Tunnel，并配置 HTTPS。"),
    ]
    add_table(doc, ["能力域", "当前可用能力", "当前边界"], capability_rows, [1600, 4700, 3060], body_size=8.9)

    doc.add_heading("三、已验证的运行指标", level=1)
    add_body(doc, "以下数据来自当前项目的健康检查、自动化测试、离线 RAG 验证和源码扫描，不代表未来生产环境的容量承诺。")
    metric_rows = [
        ("创业黑马公告", "1,297 条", "最新公告 ID 1225529658；最新披露日期 2026-08-29"),
        ("深交所规则", "1,577 条 / 30 个分类", "已纳入本地共享知识库索引"),
        ("证监会规则", "579 条", "已纳入本地共享知识库索引"),
        ("共享知识库", "4,436 份有效资料", "包含公司资料、法规、规则、公告、案例和上传文件"),
        ("Embedding 状态", "4,436 已完成 / 0 待处理 / 0 失败", "内容哈希缓存和增量处理可用"),
        ("向量索引", "275,868 个已索引 Chunk", "FAISS：web/runtime/faiss/local-BAAI-bge-m3-38f252d1c729"),
        ("本地 Embedding", "BAAI/bge-m3，1024 维", "CPU 可运行；AI 增强模式与本地检索模式共用"),
        ("本地 OCR", "可用", "Windows Media OCR"),
        ("自动化测试", "118 项全部通过", "包含本地 RAG、财务规则和边界场景"),
        ("DOM 契约检查", "241 个 ID 选择器全部存在", "前端核心交互契约通过"),
        ("离线安全验证", "通过", "EXTERNAL_EMBEDDING_API: BLOCKED_AND_NOT_USED"),
        ("密钥扫描", "2,504 个源码文件无疑似硬编码凭据", "仍需在部署环节使用服务器环境变量"),
    ]
    add_table(doc, ["指标", "结果", "说明"], metric_rows, [2100, 2300, 4960], body_size=9.0)

    doc.add_heading("四、目前仍需提升的部分", level=1)
    add_note(doc, "判断边界", "当前报告把“代码存在”与“生产验收完成”分开。阿里云 Provider 已预留，但没有真实公司配置和正式验收时，不应把云端 LLM 推理写成已完成能力。", fill="FFF8E8", label_color=GOLD)
    gap_rows = [
        ("P0", "真实 LLM 推理验收", "当前默认 local_rag 只做本地检索，不提供真正的 AI 推理。", "配置公司阿里云模型后，以脱敏材料完成结构化风险报告、追问、失败重试和费用记录验收。"),
        ("P0", "复杂表格与扫描件", "扫描 PDF、跨页表格、合并单元格、印章和负号识别仍可能影响财务结论。", "补充真实财务报表与批注样本，建立表格识别准确率、负号召回率和人工抽检门槛。"),
        ("P0", "定期报告规则结构化", "规则库可被检索，但并非所有深交所定期报告规则和会计准则都已转成可执行检查项。", "将适用范围、触发条件、例外、期限、披露位置和引用条款结构化，并保留版本和生效日期。"),
        ("P0", "历史样本规模", "财务案例库目前只有 1 个已确认案例、19 条人工批注，不能代表充分训练或校准数据。", "持续收集正确版/错误版/最终公告成套材料，区分误报、漏报和规则适用争议。"),
        ("P1", "人工审核闭环", "已有确认、误报、补充材料和撤销，但缺少审核人分配、多级复核、版本锁定、整改跟踪和正式签署。", "增加状态流转、责任人、截止时间、复核意见、整改凭证和报告版本快照。"),
        ("P1", "报告与批注导出", "系统可展示风险和出处，但 Word 原文批注、对比报告和正式归档仍需统一。", "以原文为底稿插入可追踪批注，保留页码、段落、风险 ID 和依据链接。"),
        ("P1", "统计口径治理", "数据库 Chunk 数与 FAISS 向量数需要统一统计口径，避免管理页面出现数量不一致。", "建立索引健康检查、孤儿向量清理、失败重试和重建前后校验。"),
        ("P2", "可靠部署与安全", "当前为 SQLite + FAISS + 本地文件的单机形态，不具备高可靠生产部署能力。", "按需要升级 PostgreSQL/pgvector、对象存储、备份恢复、加密、权限审计、限流和监控。"),
        ("P2", "跨网络访问", "局域网可访问，不同网络访问尚未完成，直接暴露端口不适合作为正式方案。", "使用 VPN、Tailscale 或 Cloudflare Tunnel，配置 HTTPS、访问控制和失联恢复策略。"),
    ]
    add_table(doc, ["优先级", "待提升项", "现状与风险", "建议动作"], gap_rows, [820, 1900, 3100, 3540], body_size=8.7)

    doc.add_heading("五、分阶段实施建议", level=1)
    phase_rows = [
        ("第一阶段：本地内测巩固", "保持 local_rag 安全模式；补充真实定期报告、财务报表和 XLS 对照案例；完成复杂表格与扫描件专项测试。", "可稳定复现审阅、引用、勾稽和对比结果。"),
        ("第二阶段：云端增强验收", "接入公司提供的阿里云 LLM 配置；仅发送必要 RAG 上下文；完成结构化输出校验、重试、成本统计和断网回退。", "在可控范围内验证真正的 AI 合规推理。"),
        ("第三阶段：审阅流程产品化", "增加审核分配、多级复核、版本冻结、整改闭环、Word 批注和正式报告归档。", "形成可追踪、可复核的业务流程。"),
        ("第四阶段：生产部署治理", "HTTPS、VPN/专网、备份恢复、加密、监控、限流、企业登录和长期运行演练。", "满足公司内部正式使用的可靠性与安全要求。"),
    ]
    add_table(doc, ["阶段", "重点工作", "阶段目标"], phase_rows, [2200, 4850, 2310], body_size=9.0)

    doc.add_heading("六、后续需要提供的资料", level=1)
    add_body(doc, "资料不需要一次性全部提供。优先提供带版本、日期和来源的原始 PDF、DOCX、XLSX 或 CSV；涉及敏感信息可以先脱敏。")
    materials_rows = [
        ("优先", "最新公司章程、信息披露制度、关联交易制度、对外投资制度、对外担保制度、重大交易制度", "审批门槛、披露义务、关联方识别和条款溯源"),
        ("优先", "最新股权结构、控股股东/实际控制人、董监高和关联方汇总表", "公司画像、控制关系和关联交易判断"),
        ("优先", "最近一期定期报告及上年同期报告，附对应 XLS/XLSX 财务底稿", "表格解析、跨表勾稽、同比和历史对比"),
        ("优先", "一套完整正确案例及其草稿、批注版或最终公告", "学习公司实际审阅口径，校准误报和漏报"),
        ("建议", "近三年董事会/股东会决议、历史公告、监管问询及回复", "历史审批口径、监管关注点和文书生成参考"),
        ("云端接入时", "阿里云 Endpoint、模型名称、是否允许材料出域、日志/保留/训练政策和预算上限", "完成真实 LLM 验收；API Key 由部署平台注入服务器，不在聊天中提供"),
    ]
    add_table(doc, ["级别", "资料", "主要用途"], materials_rows, [1000, 5100, 3260], body_size=8.9)

    doc.add_heading("七、成本与技术可行性概览", level=1)
    add_body(doc, "技术上可行：当前已经验证本地解析、Embedding、FAISS、RAG 和规则检查可以在 Windows 本机离线运行。后续只需把云端 LLM 作为独立 Provider 接入，不改变本地知识库和检索链路。")
    cost_rows = [
        ("当前 local_rag", "外部 Embedding 调用 0；外部 LLM 调用 0", "主要成本是本机 CPU/内存、磁盘和维护时间；适合个人内测和敏感资料查询。"),
        ("云端增强单次估算", "按每次约 8,000-20,000 输入 token、1,000-3,000 输出 token 预算", "实际金额取决于百炼模型和公司合同单价；系统应按“输入 token × 输入单价 + 输出 token × 输出单价”计费。"),
        ("个人日常使用", "示例：5 次分析/日约 45,000-115,000 token/日", "这是容量预算示例，不是阿里云报价；正式接入后以用量日志校准。"),
        ("小团队使用", "示例：30 次分析/日约 270,000-690,000 token/日", "可通过限制上下文长度、缓存检索结果和按需调用 LLM 控制成本。"),
    ]
    add_table(doc, ["场景", "用量口径", "说明"], cost_rows, [1900, 3500, 3960], body_size=9.0)
    add_note(doc, "成本提示", "以上 token 为工程预算区间，不能替代阿里云正式报价。当前未配置真实阿里云 API，因此本地测试阶段不会产生阿里云调用费用。", fill="FFF8E8", label_color=GOLD)

    doc.add_heading("八、验收建议", level=1)
    add_body(doc, "建议把“可用”与“可上线”分成两套门槛：")
    add_bullet(doc, "可用门槛：上传 PDF/DOCX/XLSX/CSV 后能完成解析、分块、本地 Embedding、FAISS 检索，并返回可点击的原文出处。")
    add_bullet(doc, "审阅门槛：风险能够回带原文，财务项目能够与 XLS 勾稽，同比百分比与百分点变化不混用，人工确认可撤销并留痕。")
    add_bullet(doc, "上线门槛：完成真实阿里云 LLM 验收、复杂表格抽检、权限与备份演练、报告版本冻结、整改闭环和 HTTPS/专网访问。")

    doc.add_heading("附录：本次评估依据", level=1)
    add_body(doc, "本报告依据当前工作目录中的项目代码、运行时健康检查、118 项 Python 自动化测试、241 项 DOM 契约检查、本地离线 RAG 验证、数据索引统计和源码密钥扫描结果编制。报告中的数量为评估时点快照，公告与规则同步后可能变化。", after=4)
    add_body(doc, "重要限制：知识库检索结果只能说明当前本地资料中是否存在相关依据；对法规适用性、事实认定和最终披露责任的判断，仍应由公司证券部、董秘办、法务或其他授权人员确认。", color=RED, italic=False, after=0)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(str(OUT))


if __name__ == "__main__":
    build()
