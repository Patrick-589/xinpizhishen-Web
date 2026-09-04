from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI证券合规风险管理与智能审阅平台_可行性研究报告.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "263548"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F4F7FB"
PALE_GOLD = "FFF7E6"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "176B61"
WHITE = "FFFFFF"
BORDER = "D7DEE8"


def set_font(run, size=None, bold=None, color=None, italic=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    cell.width = Twips(width_dxa)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    if sum(widths) != 9360:
        raise ValueError(f"Table widths must total 9360 DXA, got {sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def table_text(cell, text, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    alignments = alignments or [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    for i, header in enumerate(headers):
        shade_cell(table.rows[0].cells[i], LIGHT_GRAY)
        table_text(table.rows[0].cells[i], header, bold=True, color=NAVY, size=9.2, align=alignments[i])
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            if row_index % 2 == 1:
                shade_cell(row.cells[col_index], "FAFBFC")
            table_text(row.cells[col_index], value, align=alignments[col_index])
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(5)
    return table


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p._p.get_or_add_pPr().append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, color=INK)
    else:
        run = p.add_run(text)
        set_font(run, color=INK)
    return p


def add_para(doc, text="", bold=False, color=INK, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
             before=0, after=6, line=1.10, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    r1 = p.add_run(f"{label}：")
    set_font(r1, bold=True, color=accent)
    r2 = p.add_run(text)
    set_font(r2, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_font(suffix, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    settings = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.08


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("信披智审 | 可行性研究报告")
    set_font(run, size=8.5, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)
    first_footer = section.first_page_footer
    first_p = first_footer.paragraphs[0]
    first_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_run = first_p.add_run("内部可编辑工作稿 | V1.0")
    set_font(first_run, size=8.5, color=MUTED)


def add_cover(doc):
    add_para(doc, "FEASIBILITY STUDY", color=BLUE, size=10, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=86, after=18, keep=True)
    add_para(doc, "AI证券合规风险管理与智能审阅平台", color=NAVY, size=26, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.05, keep=True)
    add_para(doc, "可行性研究报告", color=NAVY, size=22, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=18, keep=True)
    add_para(doc, "面向创业黑马科技集团股份有限公司证券部\n10人以内本地使用场景",
             color=DARK_BLUE, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=52, line=1.4)
    add_para(doc, "项目周期", color=MUTED, size=9.5, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_para(doc, "2周开发 + 2周Review与验收", color=BLUE, size=15, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=22)
    add_para(doc, "版本：V1.0 可编辑工作稿\n编制日期：2026年8月19日\n适用范围：个人/小团队本地MVP，不含登录及企业级部署",
             color=MUTED, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.45)
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    doc.core_properties.title = "AI证券合规风险管理与智能审阅平台可行性研究报告"
    doc.core_properties.subject = "上市公司证券部本地MVP技术可行性、成本与四周计划"
    doc.core_properties.author = "信披智审项目组"
    doc.core_properties.keywords = "证券合规, 信息披露, RAG, 智能审阅, 可行性研究"

    bullet_id = add_numbering_definition(doc, "bullet")
    number_id = add_numbering_definition(doc, "decimal")
    add_cover(doc)

    add_heading(doc, "执行摘要", 1)
    add_callout(
        doc,
        "可研结论",
        "项目在技术上可行，且当前网页、公告/规则索引、Windows本地服务和确定性检查已形成可继续开发的基础。"
        "在公司于第2个工作日前提供P0资料、于第5个工作日前确认模型接入条件的前提下，可用2周完成可评审的本地RAG MVP，"
        "再用2周进行证券部Review、案例校准、缺陷修复和验收。"
    )
    add_para(doc, "本报告采用“小范围、本地优先、保留现有系统、AI结论必须可溯源并由人工确认”的实施口径。当前阶段不建设登录、复杂权限、企业级部署、移动端或多级审批流；不把公开公告/规则目录误称为已完成的正文知识库，也不把尚未接入的LLM、Embedding、向量数据库、RAG和长期Memory写成现有能力。")
    add_table(
        doc,
        ["决策项", "建议结论", "边界/前提"],
        [
            ["是否立项", "建议按四周MVP验证项目立项", "第1-2周开发，第3-4周Review与验收"],
            ["技术路线", "延续现有Web + Python本地服务，增量加入SQLite、向量检索与AI Provider", "两周内不迁移Next.js，不引入企业级中间件"],
            ["使用范围", "10人以内、可信本机或受控局域网", "不开放公网；无登录仅适用于受控环境"],
            ["预算口径", "一次性内部投入折算约4.3万-8.7万元", "规划估算，不是供应商报价；不含正式采购和税费"],
            ["运行成本", "现金支出建议按300-1,900元/月预留", "模型、Embedding、OCR与备份按量；以实际厂商账单为准"],
        ],
        [1500, 3300, 4560],
    )

    add_heading(doc, "一、项目背景与建设目标", 1)
    add_heading(doc, "1.1 项目定位", 2)
    add_para(doc, "建设面向上市公司证券部实际工作的“AI证券合规风险管理与智能审阅平台”。系统不是普通聊天机器人，也不是简单PDF总结工具，而是围绕材料上传、场景识别、公司知识检索、监管规则检索、历史案例匹配、结构化风险输出、人工确认和审查留痕形成闭环。")
    add_para(doc, "核心目标是同时解决两个问题：第一，基于公司章程、制度、股权结构和历史审批口径进行公司个性化合规推断；第二，在审阅工作台中把每项风险定位到原文，并展示可核验的法规、公司制度和历史案例来源。")

    add_heading(doc, "1.2 本次可研范围", 2)
    for text in [
        "保留当前网页、公告/规则索引、Windows本地服务和确定性检查，不推翻现有成果。",
        "目标用户不超过10人，不建设用户登录和复杂权限体系；任务与复核人暂采用手工选择或文本记录。",
        "第1-2周完成可演示、可上传真实/脱敏材料、可检索来源并形成结构化风险的本地MVP。",
        "第3-4周由证券部使用真实/脱敏样本Review，校准规则、提示词、检索和风险等级。",
        "最终合规结论保留人工确认；无直接依据时必须输出“当前知识库未检索到直接依据，请人工进一步确认”。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "二、当前项目基础与差距", 1)
    add_heading(doc, "2.1 已具备基础（截至2026年8月19日）", 2)
    add_table(
        doc,
        ["能力/数据", "当前状态", "真实边界", "可复用价值"],
        [
            ["本地网页与Windows启动", "已完成并可运行", "本机127.0.0.1服务", "可直接延展为审阅工作台"],
            ["文档上传与解析", "支持PDF、DOCX、TXT、MD，单文件20MB", "扫描PDF暂不支持全面OCR", "可用于首批文本型样本验证"],
            ["确定性检查", "金额/比例、绝对化表述、信息缺失等本地规则", "analysisMode为local_deterministic_rules，aiUsed=false", "适合作为AI前后的可信校验层"],
            ["创业黑马公告", "1,281条索引，2017年至2026-08-18", "索引及官方PDF入口，不等于完整正文知识库", "公司公开资料基线"],
            ["对标公司公告", "传智教育712条、行动教育493条", "仅作公开信息对比，不作为创业黑马审批依据", "公告样式与同业披露参考"],
            ["监管规则", "深交所508条；证监会579条目录", "证监会日常同步为目录、效力状态、文号及原文入口", "法规来源清单与后续正文导入入口"],
        ],
        [1500, 2200, 3000, 2660],
    )

    add_heading(doc, "2.2 尚未具备的关键能力", 2)
    for text in [
        "真实LLM调用与统一AI Provider；当前没有模型生成的合规结论。",
        "Embedding、向量数据库、混合检索、重排和RAG来源约束。",
        "法规、公告及公司制度的正文分块知识库；当前大量数据仍是索引。",
        "SQLite任务库、审查记录、任务上下文Memory和可恢复的长期会话。",
        "公司画像知识库，包括章程、股权、关联方、董监高、审批流程和内部制度。",
        "以证券部已确认案例为基准的准确率、误报率、漏报率和引用有效性评估。",
    ]:
        add_list_item(doc, text, bullet_id)
    add_callout(doc, "判断原则", "现有系统可以证明产品形态和本地流程可运行，但不能据此宣称已经具备真实AI、RAG或公司个性化合规判断。", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "三、技术可行性分析", 1)
    add_heading(doc, "3.1 总体判断", 2)
    add_para(doc, "从工程实现看，本项目的核心组件均有成熟技术路径：PDF/DOCX解析、文本分块、Embedding、向量检索、关键词检索、结构化模型输出、原文定位、人工复核和本地持久化均可在Windows本机完成。项目的主要不确定性不在“能否开发”，而在公司资料是否及时、条款版本是否准确、样本是否有最终人工结论，以及模型是否允许处理相关材料。")
    add_table(
        doc,
        ["维度", "可行性", "依据", "主要约束"],
        [
            ["产品可行性", "高", "已有可运行页面、任务、审阅和数据中心", "需用真实工作流验证交互效率"],
            ["技术可行性", "高", "本地Python可承载解析、检索与模型代理", "两周内需控制范围，避免技术迁移"],
            ["数据可行性", "中", "公开索引充足，公司内部资料尚待提供", "正文质量、版本和条款元数据决定RAG可靠性"],
            ["进度可行性", "中高", "一名熟悉现有项目的全栈AI开发者可增量完成", "P0资料必须在D2前到位"],
            ["运营可行性", "高", "10人以内可采用本地文件、SQLite和环境变量", "无登录时不得暴露公网"],
            ["合规可行性", "有条件", "所有AI输出可强制引用来源并人工确认", "需公司确认材料能否发送至阿里云"],
        ],
        [1400, 1100, 3500, 3360],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "3.2 推荐技术路线（两周MVP）", 2)
    for text in [
        "前端：继续使用现有HTML/CSS/JavaScript工作台，完善上传、来源引用、任务恢复和复核操作；两周内不迁移框架。",
        "本地服务：Python 3.12+，将现有server.py按解析、知识库、检索、AI和任务模块拆分；对外仅提供本地HTTP API。",
        "文档解析：python-docx、pypdf；保留段落、页码和章节映射。扫描件OCR作为条件功能，不列入必达验收。",
        "本地数据：SQLite保存任务、文档元数据、对话、风险和复核结果；原始文件与解析结果使用本地目录保存。",
        "向量检索：采用成熟的本地向量库（建议Chroma持久化模式）保存Embedding及元数据，不自研向量索引。",
        "混合检索：关键词/BM25 + 向量召回 + 元数据过滤；优先返回现行有效、公司适用且版本明确的内容。",
        "AI Provider：服务端封装chat、analyze、embedding接口；默认支持阿里云百炼或企业OpenAI兼容网关，开发时可使用mock。",
        "结构化输出：使用JSON Schema/Pydantic约束风险字段；来源不足时降级为待人工确认，禁止生成无法核验的条款。",
        "任务Memory：按任务ID保存对话摘要、上传文件、风险、用户修改和复核结论；不建设用户身份数据库。",
        "测试：保留Node/Playwright页面测试，新增Python单元测试、RAG来源测试、提示词回归样本和安全扫描。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "3.3 RAG与个性化推理流程", 2)
    process_steps = [
        "材料解析：提取页码、章节、段落、表格及原文位置，生成文档哈希。",
        "知识标准化：记录title、source、date、version、effective_status、page、chunk_id等元数据。",
        "分块与Embedding：中文正文建议按500-800字分块，重叠80-120字；具体参数在Review阶段调优。",
        "场景识别：识别合同、股权交易、关联交易、担保、投资、公告或制度等业务场景。",
        "多库检索：分别检索公司画像、内部制度、监管法规、历史公告和确认案例，并保留各自来源类型。",
        "规则与确定性计算：先计算金额、比例、累计口径和主体关系，再将结果交给模型解释。",
        "模型综合：仅基于候选证据输出结构化风险；每项风险绑定source_documents。",
        "人工复核：证券部确认、误报、补充材料或完成；结果回写任务Memory和评估样本。",
    ]
    for step in process_steps:
        add_list_item(doc, step, number_id)

    add_heading(doc, "四、技术需求说明", 1)
    add_heading(doc, "4.1 本地运行环境", 2)
    add_table(
        doc,
        ["项目", "最低建议", "说明"],
        [
            ["操作系统", "Windows 11 64位", "沿用当前使用环境"],
            ["CPU/内存", "4核CPU、16GB内存", "文档解析与本地向量检索；大批量OCR需更高配置"],
            ["磁盘", "20GB可用空间起步", "原文、解析文本、向量索引、备份按数据量增长"],
            ["网络", "可访问巨潮、深交所、证监会及获批模型网关", "无外网时保留本地mock和已同步资料"],
            ["浏览器", "Edge或Chrome近期版本", "本地访问，不开放公网"],
            ["使用方式", "可信本机或受控局域网，10人以内", "不登录模式依赖Windows账号、目录权限及网络边界"],
        ],
        [1800, 2600, 4960],
    )

    add_heading(doc, "4.2 AI与数据安全要求", 2)
    for text in [
        "真实API Key只能通过服务器环境变量注入，禁止写入HTML、前端JavaScript、源码、报告、聊天或截图。",
        "公司需确认ALIYUN_BASE_URL、ALIYUN_MODEL、ALIYUN_EMBEDDING_MODEL及是否允许材料发送至阿里云。",
        "前端只调用本地服务，不能直接访问模型厂商；日志不记录完整敏感正文、密钥或上游请求头。",
        "公司内部资料进入知识库前需脱敏、标注保密级别、有效状态、来源部门和版本。",
        "无登录版本仅用于可信本机或受控网络；若后续开放跨部门网络访问，应单独增加身份认证和权限控制。",
        "每次AI回答必须保存检索来源、模型、提示词版本和时间；引用不足时不得给出确定性合规结论。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "4.3 建议的数据结构", 2)
    add_para(doc, "两周MVP不建设企业级用户库，但应建立以下本地实体：Document、KnowledgeDocument、Chunk、AnalysisTask、Risk、Review、Conversation、TaskMemory和AuditEvent。User实体暂不启用；处理人可由复核时手工选择或填写。")
    add_callout(doc, "长期Memory边界", "本阶段保留任务上下文和企业知识Memory，不做跨人员画像或隐式用户记忆；共享电脑上的历史内容按任务ID恢复。")

    add_heading(doc, "五、MVP功能范围与非目标", 1)
    add_table(
        doc,
        ["类别", "两周必须完成", "本期明确不做"],
        [
            ["文档", "PDF/DOCX/TXT/MD上传、解析、原文段落定位", "复杂Excel/PPT、全面扫描件OCR"],
            ["知识库", "公司资料、法规正文和案例的元数据导入、分块、检索", "自动理解所有历史材料并直接下结论"],
            ["AI", "Provider、Embedding、RAG、多轮任务上下文、结构化风险", "无依据回答、自动形成最终合规意见"],
            ["审阅", "风险列表、来源、建议、原文跳转、确认/误报/备注", "复杂多级审批流、移动端"],
            ["记录", "任务、对话、风险、复核结果本地持久化", "用户账号、LDAP、企业微信/钉钉登录"],
            ["数据更新", "现有公告/规则索引继续同步；知识文档手工导入", "自动抓取全部法规/公告正文并无人工校验入库"],
        ],
        [1400, 4300, 3660],
    )

    add_heading(doc, "六、公司需要提供的资料与确认事项", 1)
    add_heading(doc, "6.1 P0资料（直接影响两周开发）", 2)
    add_table(
        doc,
        ["优先级", "资料/确认项", "最晚时间", "用途与要求"],
        [
            ["P0", "最新公司章程及仍可能适用的历史版本", "D2", "董事会、股东会权限及审批门槛；需标版本、生效日、有效状态"],
            ["P0", "信息披露、关联交易、对外投资、对外担保、重大交易/重大事项制度", "D2", "至少5项内部制度；原始PDF/DOCX优先，标注替代/废止关系"],
            ["P0", "当前股权结构、控股股东、实际控制人、关联方清单", "D2", "用于控制权和关联关系判断；敏感字段可脱敏"],
            ["P0", "现任董事、监事、高管名单及任职/控制关系", "D2", "用于关联方识别、回避表决和公司画像"],
            ["P0", "证券部现行审批流程与业务口径", "D2", "合同预审、含税口径、12个月累计、同类交易、紧急补审、最终确认人"],
            ["P0", "至少3套完整真实或脱敏案例", "D3", "优先股权交易、合同、关联交易；含协议、审批、决议、公告和最终判断"],
            ["P0", "模型接入与数据使用确认", "D5", "确认服务方式、Base URL、模型、Embedding模型、材料是否允许发送；不要在文档/聊天提供Key"],
            ["P0", "证券部Review负责人及验收人", "D1", "负责人：[待确认]；验收人：[待确认]"],
        ],
        [760, 2710, 900, 4990],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "6.2 P1资料（第3-4周校准）", 2)
    for text in [
        "累计10套历史交易材料，并提供对应董事会/股东会决议、最终公告、证券部最终判断和实际审批路径。",
        "最近三年董事会/股东会决议及5份监管问询回复。",
        "证券部确认过的典型公告、公告模板、内部检查清单。",
        "历史误报、漏报和重大修改案例，用于评估与校准。",
        "每条业务口径均附“制度名称 + 条款号 + 生效日期”，不能只提供口头结论。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "6.3 文件元数据要求", 2)
    add_para(doc, "每份资料至少记录：标题、资料类型、形成/发布日期、生效日期、版本号、当前是否有效、来源部门、保密级别、替代或废止版本。扫描PDF应具备可OCR质量；身份证号、账号、手机号等信息可先脱敏。只有协议而无最终处理结果的材料可用于风险分析，但不能用于验证公司历史审批口径。")

    doc.add_page_break()
    add_heading(doc, "七、四周实施时间表", 1)
    add_callout(doc, "排期口径", "共4周：前2周完成可评审MVP，后2周由证券部Review并完成修正和验收。按1名全栈AI开发者主导、1名证券部负责人兼职Review估算。")
    add_table(
        doc,
        ["阶段", "时间", "关键工作", "公司输入", "交付物/退出条件"],
        [
            ["启动与冻结范围", "D1", "核对现状、冻结MVP边界、建立资料清单和验收样本", "指定负责人；确认不登录、不部署", "范围基线、资料登记表"],
            ["知识导入基础", "D2-D3", "SQLite元数据、文档哈希、解析与分块、资料版本状态", "P0制度、股权、董监高、流程", "公司画像和知识文档可导入"],
            ["检索与Embedding", "D4-D5", "向量库、关键词检索、来源过滤、检索结果结构", "确认Embedding模型及数据权限", "可按问题返回真实来源片段"],
            ["模型与结构化风险", "D6-D7", "AI Provider、RAG提示词、JSON风险Schema、无依据降级", "环境变量注入模型配置", "风险含source_documents"],
            ["审阅闭环", "D8-D9", "原文跳转、任务Memory、确认/误报/备注、报告输出", "3套完整案例及预期结论", "端到端审阅可完成"],
            ["开发验收版", "D10", "集成测试、安全扫描、Windows启动、问题清单", "证券部冒烟测试", "两周MVP候选版"],
            ["Review第1轮", "第3周", "10套案例逐项核对，记录误报、漏报、引用和定位问题", "补齐P1案例、问询、模板", "Review问题清单、指标基线"],
            ["修正与回归", "第4周D1-D3", "调整分块、检索、规则、提示词和风险等级；回归测试", "确认业务口径和争议项", "候选发布版"],
            ["验收与移交", "第4周D4-D5", "复测、使用说明、资料更新说明、遗留项分级", "负责人签署验收意见", "MVP验收记录和下一阶段计划"],
        ],
        [850, 1050, 3150, 1850, 2460],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_heading(doc, "7.1 关键进度门槛", 2)
    for text in [
        "D2未收到P0公司制度和画像资料：继续开发技术框架，但个性化判断顺延，不能作为真实公司口径演示。",
        "D5未确认模型/Embedding接入：使用mock完成流程，真实RAG模型效果验收顺延。",
        "第3周未提供带最终结论的历史案例：只能评估功能完整性，无法评估公司口径准确率。",
        "Review期间发现法规/制度版本冲突：该结论标记待确认，不由AI自动选择效力。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "八、成本预测", 1)
    add_heading(doc, "8.1 测算假设", 2)
    add_para(doc, "成本按10人以内本地MVP估算，不含企业级服务器、SSO、LDAP、PostgreSQL集群、Redis、对象存储、移动端和正式等保建设。人工单价、模型单价均为预算参数，不代表劳动合同成本或供应商正式报价。")
    for text in [
        "开发：1名全栈AI开发者投入20个工作日，规划单价1,500-2,500元/人日。",
        "证券业务Review：1名证券部人员累计8-12人日，规划单价800-1,500元/人日。",
        "资料整理：3-7人日，规划单价800-1,500元/人日。",
        "AI用量：20-50次分析/日，每次约3万输入Token、5千输出Token，每月22个工作日。",
        "模型单价按规划区间测算，最终以公司选定的阿里云模型和当期官方账单为准。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "8.2 一次性建设成本（内部投入折算）", 2)
    add_table(
        doc,
        ["成本项", "工作量/数量", "单价假设", "估算区间", "说明"],
        [
            ["开发与集成", "20人日", "1,500-2,500元/人日", "30,000-50,000元", "含两周开发、两周修正支持"],
            ["证券业务Review", "8-12人日", "800-1,500元/人日", "6,400-18,000元", "样本标注、条款核对、验收"],
            ["资料清洗与录入", "3-7人日", "800-1,500元/人日", "2,400-10,500元", "版本、条款、来源、有效状态"],
            ["本地软硬件", "现有Windows电脑", "开源组件0元", "0-1,200元", "可选移动硬盘/备份介质"],
            ["风险预备金", "约10%", "按前述投入", "4,000-8,000元", "处理格式异常、接口变化和返工"],
            ["合计", "-", "-", "约42,800-87,700元", "建议管理预算按4.3万-8.7万元"],
        ],
        [1500, 1500, 1900, 1700, 2760],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "8.3 月度运行成本", 2)
    add_table(
        doc,
        ["成本项", "正常使用估算", "建议预算上限", "计费说明"],
        [
            ["LLM + Embedding API", "约50-500元/月（纯Token理论区间）", "300-1,500元/月", "包含重试、长文、峰值和价格变动余量"],
            ["OCR（可选）", "0-300元/月", "300元/月", "文本型PDF/DOCX不产生OCR费用"],
            ["本地存储/备份", "0-100元/月", "100元/月", "使用现有磁盘时现金支出可为0"],
            ["维护人工", "1-2人日/月", "1,500-5,000元/月折算", "规则更新、样本回归、故障处理"],
            ["现金支出小计", "约50-900元/月", "建议300-1,900元/月", "不含内部维护人工机会成本"],
        ],
        [1900, 2500, 1900, 3060],
    )
    add_callout(doc, "成本控制", "系统已具备AI用量页面，正式接入后应按日记录分析次数、输入Token、输出Token和估算成本，并设置月度预算预警。", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "8.4 API成本计算公式", 2)
    add_para(doc, "月度模型成本 = 输入Token（百万）× 输入单价 + 输出Token（百万）× 输出单价 + Embedding Token（百万）× Embedding单价。按本报告工作量假设，月输入约1,320万-3,300万Token，月输出约220万-550万Token。由于模型、上下文长度、缓存和批处理价格会变化，报告中的金额只能用于预算，不作为厂商报价。")

    add_heading(doc, "九、主要风险与应对", 1)
    add_table(
        doc,
        ["风险", "影响", "概率", "应对措施"],
        [
            ["公司资料不完整或版本不清", "AI引用过期制度或无法判断门槛", "高", "D2建立来源/版本/有效状态清单；冲突一律待确认"],
            ["只有协议、没有最终处理结果", "无法学习公司历史审批口径", "高", "Review样本必须尽量成套，区分分析样本与验证样本"],
            ["公开索引被误当正文知识库", "检索依据不足或引用错误", "中高", "正文单独导入并校验哈希、页码和版本；索引仅作入口"],
            ["模型编造条款或案例", "形成严重合规风险", "中", "强制source_documents；无来源不得输出确定结论；人工复核"],
            ["原文定位不准确", "降低审阅效率和可追溯性", "中", "文本型PDF/DOCX优先；保存段落与页码映射；扫描件另行标记"],
            ["无登录且局域网开放", "非授权人员访问敏感材料", "中", "仅限可信本机/受控IP；使用Windows目录权限；禁止公网暴露"],
            ["官方接口变化", "公告/规则更新中断", "中", "原子更新、旧数据保留、失败告警和手工更新入口"],
            ["两周范围膨胀", "无法按期交付", "中高", "冻结MVP非目标；复杂OCR、企业部署和多级审批进入后续阶段"],
        ],
        [1800, 2800, 900, 3860],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    add_heading(doc, "十、Review与验收标准", 1)
    add_heading(doc, "10.1 第3-4周Review方法", 2)
    for text in [
        "由证券部对每个样本给出预期风险、适用条款、审批路径、披露结论和最终处理结果。",
        "逐项记录命中、误报、漏报、来源错误、原文定位错误和建议不可执行等问题。",
        "争议项不强行纳入正确/错误统计，先标记“制度/口径待确认”。",
        "所有修正进入固定回归集，防止调整一个规则后破坏其他案例。",
    ]:
        add_list_item(doc, text, bullet_id)

    add_heading(doc, "10.2 建议验收门槛", 2)
    add_table(
        doc,
        ["验收项", "建议标准", "说明"],
        [
            ["文件上传与解析", "支持范围内文件成功解析；失败有明确提示", "扫描件和加密PDF不静默跳过"],
            ["原文定位", "文本型样本定位成功率建议不低于90%", "以证券部选定样本统计"],
            ["来源完整性", "AI生成风险100%包含可打开的来源；无来源明确降级", "不要求每个问题都必须有结论"],
            ["引用真实性", "抽样中不得出现伪造法规、条款或案例", "发现即作为阻断缺陷"],
            ["任务恢复", "重新进入系统可恢复任务材料、风险和复核结果", "不依赖用户账号"],
            ["人工复核", "支持确认、误报、补充材料、备注和完成状态", "操作可在任务记录中追溯"],
            ["安全", "源码和前端无API Key；密钥扫描通过", "Key仅由服务器环境变量读取"],
            ["页面质量", "桌面和手机宽度无横向溢出、关键内容不重叠", "保留自动化页面测试"],
        ],
        [1900, 3300, 4160],
    )
    add_para(doc, "准确率、召回率和误报率的最终目标值应在第3周基线测试后，由证券部根据样本类型和风险容忍度确认。小样本阶段不宜以单一百分比替代业务判断。", italic=True, color=MUTED)

    add_heading(doc, "十一、项目结论与立项建议", 1)
    add_para(doc, "本项目具备实施基础，适合采用“四周受控验证”方式推进。当前系统已经验证Windows本地运行、网页交互、公开数据同步和确定性检查；下一步技术重点是把公司资料和监管正文建设为可版本化、可检索、可引用的知识库，并用真实模型完成受来源约束的RAG审阅闭环。")
    add_callout(doc, "立项建议", "批准按本报告范围实施2周开发 + 2周Review；以P0资料到位、模型数据权限确认、证券部负责人到位作为启动门槛；四周结束后根据真实样本结果决定是否进入企业级部署和扩大使用范围。")

    add_heading(doc, "11.1 启动前需确认", 2)
    confirmation_items = [
        "项目负责人：[待确认]",
        "证券部Review负责人：[待确认]",
        "开发开始日期：[待确认]",
        "两周MVP评审日期：[待确认]",
        "四周最终验收日期：[待确认]",
        "阿里云接入方式（百炼/OpenAI兼容网关）：[待确认]",
        "材料是否允许发送至阿里云：[待确认]",
        "一次性预算上限：[待确认]",
        "月度API预算上限：[待确认]",
    ]
    for item in confirmation_items:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "附录：本报告的事实与估算边界", 1)
    add_para(doc, "事实基线来自用户提供的《上市公司证券部 AI 合规与风险管控系统》总体框架、当前项目代码与2026年8月19日已生成的数据文件。公开数据数量会随同步更新而变化。成本、工作量、Token和验收比例均为可研阶段规划参数，应在立项时由公司负责人修改确认。")
    add_para(doc, "本报告不构成法律意见、审计意见、模型厂商报价或正式采购承诺。系统输出用于辅助审查，最终合规判断由证券事务人员确认。", bold=True, color=NAVY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
