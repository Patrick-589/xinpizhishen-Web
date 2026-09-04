from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI证券合规风险管理与智能审阅平台_可行性研究报告.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "263548"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GOLD = "FFF7E6"
GOLD = "7A5A00"
BORDER = "D7DEE8"


def set_font(run, size=None, bold=None, color=None, italic=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    cell.width = Twips(width_dxa)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths):
    if sum(widths) != 9360:
        raise ValueError("Table widths must total 9360 DXA")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), BORDER)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table_text(cell, text, bold=False, color=INK, size=9.3, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    alignments = alignments or [WD_ALIGN_PARAGRAPH.LEFT] * len(headers)
    for col, header in enumerate(headers):
        shade_cell(table.rows[0].cells[col], LIGHT_GRAY)
        table_text(table.rows[0].cells[col], header, bold=True, color=NAVY, align=alignments[col])
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        current = table.add_row()
        for col, value in enumerate(row):
            if row_index % 2 == 1:
                shade_cell(current.cells[col], "FAFBFC")
            table_text(current.cells[col], value, align=alignments[col])
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_numbering_definition(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p._p.get_or_add_pPr().append(num_pr)
    run = p.add_run(text)
    set_font(run, size=11, color=INK)
    return p


def add_para(doc, text="", bold=False, color=INK, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
             before=0, after=6, line=1.10, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    label_run = p.add_run(f"{label}：")
    set_font(label_run, bold=True, color=accent)
    text_run = p.add_run(text)
    set_font(text_run, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    suffix = paragraph.add_run(" 页")
    set_font(suffix, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    heading_styles = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_styles.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.08
        style.paragraph_format.keep_with_next = True


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("AI证券合规风险管理与智能审阅平台 | 个人项目简版")
    set_font(header_run, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_run = first_footer.add_run("可编辑工作稿")
    set_font(first_run, size=8.5, color=MUTED)


def add_cover(doc):
    add_para(doc, "PERSONAL PROJECT FEASIBILITY NOTE", color=BLUE, size=10, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=108, after=20, keep=True)
    add_para(doc, "AI证券合规风险管理与智能审阅平台", color=NAVY, size=25, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.05, keep=True)
    add_para(doc, "个人项目简版可行性说明", color=NAVY, size=20, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=20, keep=True)
    add_para(doc, "用于先完成可使用的本地MVP，后续再接入公司资料和阿里云 API",
             color=DARK_BLUE, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, after=62)
    add_para(doc, "项目节奏", color=MUTED, size=9.5, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_para(doc, "第1周写框架  |  第2周实施  |  后2周 Review", color=BLUE, size=14.5, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_para(doc, "版本：V2.0 简版\n编制日期：2026年8月19日\n使用范围：个人项目和后续公司实际试用",
             color=MUTED, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.45)
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    doc.core_properties.title = "AI证券合规风险管理与智能审阅平台个人项目简版可行性说明"
    doc.core_properties.subject = "个人项目技术可行性、Token成本、公司资料与四周计划"
    doc.core_properties.author = "信披智审项目"

    bullets = add_numbering_definition(doc)
    add_cover(doc)

    add_heading(doc, "一、项目说明", 1)
    add_para(doc, "本项目先按个人使用场景开发一个可实际测试的本地应用。重点不是做普通聊天，而是让用户上传合同、公告、制度或报告后，系统能够检索公司资料和监管规则，给出带来源的风险提示，并保留人工确认。")
    add_callout(doc, "本期目标", "先把“上传材料 -> 检索依据 -> AI分析 -> 风险与来源 -> 人工确认”的完整流程跑通；公司资料到位后，再把判断改为创业黑马的个性化口径。")

    add_heading(doc, "二、技术可行性", 1)
    add_para(doc, "从实现角度看，该应用可在现有网页和 Windows 本地环境上继续完成。当前已有公告/规则索引和本地确定性检查；后续把文档解析、知识库检索和阿里云模型调用补入即可。关键不在于堆复杂技术，而在于公司制度、版本和历史案例是否真实、完整、可核验。")
    add_heading(doc, "拟实现的应用流程", 2)
    add_para(doc, "上传文件或提出问题 -> 解析文档 -> 检索公司资料、监管规则和历史案例 -> 计算金额、比例、累计口径等确定性规则 -> 调用 AI 进行带依据的分析 -> 展示风险、建议和 source_documents -> 人工确认。", color=DARK_BLUE, bold=True, line=1.25)
    add_heading(doc, "简化技术路线", 2)
    for item in [
        "前端：保留现有 HTML、CSS 和 JavaScript 工作台，增加上传、风险清单、来源跳转和 Review 操作。",
        "本地服务：Python 负责文件解析、任务保存、规则计算、知识库检索和 API 转发；模型密钥只存在服务端环境变量。",
        "知识库：公司资料、法规正文、公告和案例按版本保存；分块后做关键词和向量检索，AI 只能根据检索到的材料回答。",
        "AI：后续对接公司阿里云 API。没有 API Key 时仍可用 mock 和确定性规则开发页面与流程。",
        "结果：每个风险必须包含原文位置、法规/制度依据和来源文件；没有直接依据时提示人工进一步确认。",
    ]:
        add_list_item(doc, item, bullets)
    add_callout(doc, "可行性结论", "可按四周完成个人项目的可评审版本。真正影响效果的是公司资料和业务口径，不是是否增加登录、服务器或复杂部署。")

    doc.add_page_break()
    add_heading(doc, "三、Token 用量与费用预估", 1)
    add_para(doc, "以下为方便个人项目预算而设的计算示例，不是阿里云官方报价。示例单价按输入 2 元/百万 Token、输出 8 元/百万 Token 计算。公司确定具体模型后，只需替换单价，表中的成本即可重新计算。")
    add_callout(doc, "计算公式", "单次成本 = 输入 Token / 1,000,000 x 输入单价 + 输出 Token / 1,000,000 x 输出单价。文档越长、检索片段越多、回答越详细，Token 消耗越高。", fill=PALE_GOLD, accent=GOLD)
    add_table(
        doc,
        ["使用场景", "日均分析次数", "日均输入 / 输出 Token", "示例日成本"],
        [
            ["个人日常开发测试", "5次", "5万 / 1万", "约0.18元"],
            ["个人集中调试", "20次", "40万 / 8万", "约1.44元"],
            ["公司实际初期使用", "20次", "60万 / 10万", "约2.00元"],
            ["公司使用较频繁", "50次", "150万 / 25万", "约5.00元"],
        ],
        [2650, 1600, 3000, 2110],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
    )
    add_heading(doc, "费用怎么理解", 2)
    for item in [
        "个人开发阶段：正常测试约 0.2-1.5 元/天；按 22 个工作日计算约 4-32 元/月。为了覆盖长文重试、反复修改和临时测试，个人可预留 20-100 元/月。",
        "公司实际使用阶段：按表中 20-50 次/日的规模，生成模型约 2-5 元/天，约 44-110 元/月。建议 API 预算先按 300-1,000 元/月预留，以覆盖峰值、长文、多轮对话和模型价格变化。",
        "Embedding 在导入公司制度、法规正文和公告时会额外产生费用；OCR 只在扫描件需要识别时产生。两项不按每天固定发生，首期可另外预留 50-200 元/月。",
        "实际账单由所选阿里云模型、上下文长度、缓存和调用次数决定；系统应保留 AI 用量统计页面，按日显示分析次数、输入 Token、输出 Token 和估算成本。",
    ]:
        add_list_item(doc, item, bullets)

    doc.add_page_break()
    add_heading(doc, "四、公司后续需要提供的资料", 1)
    add_para(doc, "第一版页面和通用规则可以先开发；进入公司实际试用前，需要以下真实资料，才能做出与公司口径匹配的判断。原始 PDF、DOCX、XLSX 优先，文件需尽量标注版本、生效日期和是否仍有效。")
    add_table(
        doc,
        ["公司需要提供", "主要用途"],
        [
            ["最新公司章程及有效的历史版本", "判断董事会、股东会的权限和审批门槛。"],
            ["信息披露、关联交易、对外投资、对外担保、重大交易等制度", "判断披露、审批、关联方认定和公司内部流程。"],
            ["当前股权结构、控股股东、实控人、关联方及董监高名单", "识别关联关系、控制权变化和回避表决场景。"],
            ["3-10套完整真实或脱敏案例", "校准实际口径；每套尽量含协议、审批材料、决议、公告和最终判断。"],
            ["证券部确认的业务口径", "明确合同预审、含税口径、12个月累计、同类交易和披露标准。"],
            ["阿里云接入信息", "确认可用模型、Base URL、Embedding 模型及材料是否可发送；API Key 不提供在文档或聊天中。"],
        ],
        [4050, 5310],
    )
    add_callout(doc, "资料原则", "没有公司制度或历史审批依据时，系统只能给出通用提示，不能把通用法规误写成公司的审批结论。API Key 后续由公司通过环境变量配置，不能写进前端或文档。")

    doc.add_page_break()
    add_heading(doc, "五、四周时间表", 1)
    add_callout(doc, "排期说明", "第1周先把应用框架和资料结构写清楚；第2周开始形成可使用页面和流程；第3-4周以真实或脱敏材料进行 Review、修正和确认。本页单独作为项目排期页。")
    add_table(
        doc,
        ["时间", "工作重点", "阶段结果"],
        [
            ["第1周：写框架", "确认 MVP 页面、审阅流程、知识库分类、风险输出格式和 Token 统计口径；整理公司资料清单。", "明确可以做什么、需要哪些资料，以及每一项结果必须带来源。"],
            ["第2周：开始实施", "完成上传解析、知识库导入、规则检索、AI Provider 接口、风险清单、来源展示和任务保存。", "形成可以上传材料并输出带来源风险的本地 MVP。"],
            ["第3周：Review", "用真实或脱敏案例测试；检查风险是否命中、来源是否正确、原文定位是否可用；记录误报和漏报。", "形成 Review 问题清单和公司业务口径待确认项。"],
            ["第4周：Review与修正", "根据 Review 调整检索、提示词、规则、风险等级和页面体验；再次回归测试，确定下一阶段需求。", "交付可继续使用的个人/小团队版本，以及后续迭代清单。"],
        ],
        [1700, 4750, 2910],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    add_para(doc, "四周完成后，系统是否进入公司实际使用，以公司资料是否齐全、来源是否可追溯、证券部 Review 是否通过为准。最终合规判断始终由人工确认。", italic=True, color=MUTED, after=0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
