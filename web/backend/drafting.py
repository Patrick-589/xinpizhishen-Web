from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .rag import RAGContext
from .meeting_case_library import (
    MEETING_DOCUMENT_TYPES,
    ROLE_LABELS,
    infer_case_role,
    infer_meeting_kind,
    select_case_evidence,
    source_label,
)


MAX_DRAFT_TITLE = 120
MAX_DRAFT_SUBJECT = 4000
MAX_DRAFT_FACTS = 12000
MAX_DRAFT_REQUIREMENTS = 6000


@dataclass(frozen=True)
class DraftResult:
    title: str
    document_type: str
    content: str
    sections: list[dict[str, str]]
    sources: list[dict[str, Any]]
    warnings: list[str]
    style_sources: list[dict[str, Any]] = field(default_factory=list)


DOCUMENT_TYPES: dict[str, dict[str, Any]] = {
    "temporary_announcement": {
        "label": "临时公告",
        "heading": "关于{事项}的公告",
        "sections": ["重要提示", "事项概述", "对公司的影响", "风险提示", "备查文件"],
        "focus": "信息披露完整性、事实准确性、风险提示和备查文件",
    },
    "board_resolution": {
        "label": "董事会决议公告",
        "heading": "董事会决议公告",
        "sections": ["会议召开情况", "会议审议情况", "回避表决情况", "备查文件"],
        "focus": "董事会权限、议案逐项表决、关联董事回避和后续股东会衔接",
    },
    "shareholder_resolution": {
        "label": "股东会决议公告",
        "heading": "股东会决议公告",
        "sections": ["会议召开和出席情况", "议案审议情况", "律师见证情况", "备查文件"],
        "focus": "股东会权限、出席与表决比例、关联股东回避和律师见证",
    },
    "meeting_notice": {
        "label": "会议通知",
        "heading": "关于召开{事项}会议的通知",
        "sections": ["会议时间、地点和方式", "会议议程", "出席人员和注意事项"],
        "focus": "会议召集程序、议案清单和参会人员资格",
    },
    "meeting_sign_in": {
        "label": "会议签到表",
        "heading": "{事项}会议签到表",
        "sections": ["会议基本信息", "会议议案/事项", "签到登记"],
        "focus": "依据本次会议议案或通知形成签到表，会议名称、议案和已知信息必须与原文一致",
    },
    "meeting_agenda_summary": {
        "label": "主要议案概述",
        "heading": "{事项}主要议案概述",
        "sections": ["会议基本信息", "会议议程", "主要议案概述", "审议提示"],
        "focus": "从本次会议议案或通知中提炼议案名称、事项背景和审议要点，不补造交易事实",
    },
    "proxy_authorization": {
        "label": "授权委托书",
        "heading": "{事项}授权委托书",
        "sections": ["会议基本信息", "委托人和受托人", "授权事项", "签署确认"],
        "focus": "明确会议名称、委托权限、受托人和签署要件，未提供的信息保留待填写",
    },
    "legal_opinion": {
        "label": "法律意见书",
        "heading": "关于{事项}的法律意见书（草稿）",
        "sections": ["出具范围", "会议程序核查", "议案及文件核查", "结论与保留事项"],
        "focus": "整理会议及议案的法律核查事项，正式法律结论必须由执业律师依据完整材料确认",
    },
    "meeting_proposal": {
        "label": "会议议程/议案",
        "heading": "关于{事项}的议案",
        "sections": ["议案名称", "议案背景及审议事项", "合规审议重点", "提交审议"],
        "focus": "议案事实、审议权限、风险影响和提交层级",
    },
    "meeting_record": {
        "label": "会议记录",
        "heading": "{事项}会议记录",
        "sections": ["会议召开情况", "议案说明", "会议发言与讨论", "表决及后续事项"],
        "focus": "按具体议案记录事实确认、合规讨论、表决和后续程序",
    },
    "general_manager_meeting": {
        "label": "总经理会议",
        "heading": "{事项}总经理会议记录",
        "sections": ["会议召开情况", "议案/事项说明", "总经理审议意见", "决策及后续事项"],
        "focus": "总经理会议审议事项、经营层意见、权限衔接、风险控制和后续执行安排",
    },
    "committee_record": {
        "label": "专门委员会会议记录",
        "heading": "{事项}会议记录",
        "sections": ["会议召开情况", "议案说明", "委员发言与讨论", "表决及提交董事会"],
        "focus": "委员会职责范围、审议重点、委员独立判断和提交董事会程序",
    },
    "committee_resolution": {
        "label": "专门委员会会议决议",
        "heading": "{事项}会议决议",
        "sections": ["会议召开情况", "议案审议情况", "表决结果", "后续程序"],
        "focus": "委员会审议事项、表决结果和提交董事会的程序衔接",
    },
    "independent_director_resolution": {
        "label": "独立董事专门会议决议",
        "heading": "独立董事专门会议决议",
        "sections": ["会议召开情况", "议案审议情况", "独立董事审核意见", "提交董事会审议"],
        "focus": "独立判断、关联关系识别、回避表决和提交董事会程序",
    },
    "ballot": {
        "label": "会议表决票",
        "heading": "{事项}议案表决票",
        "sections": ["议案清单", "填表说明", "委员/董事签名"],
        "focus": "议案与表决选项完整、唯一选择和签署留痕",
    },
    "related_transaction": {
        "label": "关联交易公告",
        "heading": "关于{事项}暨关联交易的公告",
        "sections": ["关联交易概述", "关联方基本情况", "交易标的基本情况", "定价政策及定价依据", "交易协议主要内容", "关联交易目的和对公司的影响", "风险提示", "备查文件"],
        "focus": "关联方认定、交易金额和累计口径、定价公允性、审批披露和回避表决",
    },
    "investment_announcement": {
        "label": "对外投资公告",
        "heading": "关于对外投资的公告",
        "sections": ["对外投资概述", "投资标的基本情况", "投资协议主要内容", "对外投资的目的和影响", "风险分析", "备查文件"],
        "focus": "公司章程与对外投资制度的审批门槛、资金来源、投资风险和持续披露",
    },
    "periodic_report_notice": {
        "label": "定期报告风险提示",
        "heading": "定期报告披露风险提示及核对清单",
        "sections": [
            "重要提示", "公司基本情况", "主要会计数据和财务指标", "管理层讨论与分析",
            "重大事项", "股份变动及股东情况", "优先股/债券（如适用）",
            "董事、监事和高级管理人员", "公司治理", "财务报告", "备查文件",
        ],
        "focus": "深交所定期报告规则、会计类指引、历期报告对比和财务勾稽",
    },
    "szse_report_brief": {
        "label": "一页纸汇报资料",
        "heading": "关于公司业务发展及重组进展情况的汇报",
        "sections": [
            "公司业务发展历史简介",
            "公司当前业务情况",
            "重大资产重组基本情况",
            "重大资产重组最新进展",
            "后续工作安排",
        ],
        "focus": "公司历史公告事实、重组进展、监管披露规则和已确认材料",
        "one_page": True,
    },
}


def _clean(value: Any, limit: int) -> str:
    text = str(value or "").replace("\x00", "").strip()
    return text[:limit]


def _single_line(value: Any, limit: int = 160) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text if len(text) <= limit else text[:limit].rstrip() + "…"


def _evidence_excerpt(
    sources: list[dict[str, Any]],
    terms: tuple[str, ...],
    *,
    limit: int = 460,
) -> str:
    """Return short, source-labelled excerpts without inventing report facts."""
    excerpts: list[str] = []
    for source in sources:
        haystack = f"{source.get('title') or ''} {source.get('content') or ''}"
        if terms and not any(term in haystack for term in terms):
            continue
        title = _single_line(source.get("title") or "未命名公告", 80)
        content = _single_line(source.get("content"), 180)
        if content:
            excerpts.append(f"《{title}》：{content}")
        if len("；".join(excerpts)) >= limit:
            break
    return _single_line("；".join(excerpts), limit)


def _source_location(source: dict[str, Any]) -> str:
    values = []
    if source.get("page"):
        values.append(f"第{source['page']}页")
    if source.get("section"):
        values.append(str(source["section"]))
    if source.get("metadata", {}).get("csv_row_start"):
        values.append(f"CSV第{source['metadata']['csv_row_start']}行")
    if source.get("chunk_id"):
        values.append(f"Chunk {source['chunk_id']}")
    return " · ".join(values) or "位置待确认"


def _unique_sources(context: RAGContext, limit: int = 10) -> list[dict[str, Any]]:
    sources: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for source in context.citations:
        key = (
            str(source.get("document_id") or ""),
            str(source.get("chunk_id") or ""),
            str(source.get("title") or ""),
        )
        if key in seen:
            continue
        seen.add(key)
        source = dict(source)
        if source.get("document_type") == "drafting_materials":
            role = infer_case_role(str(source.get("title") or ""), str(source.get("content") or ""))
            kind = infer_meeting_kind(str(source.get("title") or ""), str(source.get("content") or ""))
            source["case_role"] = role
            source["case_role_label"] = ROLE_LABELS.get(role, role)
            metadata = dict(source.get("metadata") or {})
            metadata.update({"case_role": role, "meeting_type": kind})
            source["metadata"] = metadata
        source["citation_kind"] = "draft_basis"
        if not source.get("source_url") and source.get("document_type") in {
            "company_profile", "internal_rules", "regulations", "announcements", "cases",
            "drafting_materials",
        } and source.get("document_id"):
            source["source_url"] = f"/api/knowledge/{source['document_id']}/document"
        source["location"] = _source_location(source)
        sources.append(source)
        if len(sources) >= limit:
            break
    return sources


def _select_company_style_sources(
    sources: list[dict[str, Any]],
    limit: int = 3,
    *,
    include_all_company_announcements: bool = False,
) -> list[dict[str, Any]]:
    """Pick company reports/templates as formatting references only."""
    selected: list[dict[str, Any]] = []
    for source in sources:
        document_type = str(source.get("document_type") or "")
        if document_type not in {"announcements", "drafting_materials"}:
            continue
        title = str(source.get("title") or "")
        title_pattern = (
            r"(?:创业黑马|年度报告|半年度报告|季度报告|定期报告|重组|业务发展|重大事项|公告|汇报|模板|范本|决议)"
            if include_all_company_announcements
            else r"(?:创业黑马|年度报告|半年度报告|季度报告|定期报告|汇报|模板|范本|决议)"
        )
        if document_type != "drafting_materials" and not re.search(title_pattern, title):
            continue
        item = dict(source)
        metadata = dict(item.get("metadata") or {})
        metadata["style_reference"] = True
        item["metadata"] = metadata
        item["style_reference"] = True
        selected.append(item)
        if len(selected) >= limit:
            break
    return selected


def _format_heading(template: str, subject: str) -> str:
    clean_subject = subject.splitlines()[0].strip() if subject else "相关事项"
    if "{事项}" in template:
        return template.replace("{事项}", clean_subject[:50])
    return template


def _meeting_type_for_document(document_type: str) -> str:
    if document_type == "general_manager_meeting":
        return "general_manager"
    if document_type in {"committee_record", "committee_resolution"}:
        return "audit_committee"
    if document_type == "independent_director_resolution":
        return "independent_director"
    if document_type in {"shareholder_resolution"}:
        return "shareholders"
    if document_type == "meeting_sign_in":
        # The supplied document determines whether this is a board or
        # shareholder meeting.  A neutral value keeps both case types eligible.
        return "meeting"
    if document_type in {"meeting_agenda_summary", "legal_opinion"}:
        return "meeting"
    if document_type == "proxy_authorization":
        return "shareholders"
    if document_type in {"board_resolution", "meeting_notice", "meeting_proposal", "meeting_record"}:
        return "board"
    return "meeting"


def _meeting_source_note(source: dict[str, Any]) -> str:
    location = source.get("location") or source_label(source)
    return f"来源：{location}"


def _meeting_proposal_lines(bundle: dict[str, Any], *, limit: int = 5) -> str:
    proposals = bundle.get("proposals") or []
    if not proposals:
        return "[待补充：未从本地会议案例中检索到与本事项直接对应的议案标题和正文。]"
    lines = []
    for index, proposal in enumerate(proposals[:limit], start=1):
        title = _single_line(proposal.get("title"), 220)
        excerpt = _single_line(proposal.get("excerpt"), 480)
        source = proposal.get("source") or {}
        lines.append(f"议案{index}：《{title}》\n议案说明（本地案例摘录）：{excerpt or '[待补充]'}\n{_meeting_source_note(source)}")
    return "\n".join(lines)


def _meeting_speaking_content(bundle: dict[str, Any], document_type: str) -> str:
    proposals = bundle.get("proposals") or []
    if not proposals:
        return "[待补充：请根据本次正式议案、协议/报告书及适用制度填写发言；不得使用与本次事项无关的历史案例内容。]"
    if document_type == "general_manager_meeting":
        focus = "总经理应围绕事项背景、经营影响、资金安排、执行条件、风险控制和是否具备提交董事会/股东会审议条件发表意见"
    elif document_type == "committee_record":
        focus = "审计委员会委员应围绕议案的财务数据、评估/审计依据、风险揭示、内部控制影响及是否具备提交董事会条件发表意见"
    elif document_type == "independent_director_resolution":
        focus = "独立董事应围绕交易背景、定价/程序公允性、对中小股东影响、关联关系及是否需要回避发表独立判断"
    elif document_type in {"shareholder_resolution", "meeting_record"}:
        focus = "参会人员应围绕议案事实、审议权限、对公司及股东的影响、信息披露义务和后续审批程序进行讨论"
    else:
        focus = "参会人员应围绕议案事实、董事会审议权限、交易风险、信息披露义务及是否提交股东会进行讨论"
    lines = [f"本次发言应以本次议案为中心，重点讨论：{focus}。"]
    for index, proposal in enumerate(proposals[:5], start=1):
        title = _single_line(proposal.get("title"), 180)
        excerpt = _single_line(proposal.get("excerpt"), 360)
        source = proposal.get("source") or {}
        lines.append(
            f"议案{index}《{title}》：先确认议案背景及拟审议事项（依据：{excerpt or '待补充'}）；"
            "再核对审议权限、材料完整性、风险及披露衔接；最后记录委员/董事对提交下一层级审议的明确意见。"
            f" {_meeting_source_note(source)}"
        )
    if document_type in {"committee_record", "committee_resolution", "independent_director_resolution"}:
        lines.append("涉及关联关系的，应单独记录关联人员、回避原因及回避后的有效表决票数；本地案例未提供的姓名和票数保留待确认。")
    return "\n".join(lines)


def build_draft(
    *,
    document_type: str,
    title: str,
    subject: str,
    facts: str,
    requirements: str,
    context: RAGContext,
    structured_bundle: dict[str, Any] | None = None,
) -> DraftResult:
    template = DOCUMENT_TYPES.get(document_type) or DOCUMENT_TYPES["temporary_announcement"]
    selected_type = document_type if document_type in DOCUMENT_TYPES else "temporary_announcement"
    title = _clean(title, MAX_DRAFT_TITLE)
    subject = _clean(subject, MAX_DRAFT_SUBJECT)
    facts = _clean(facts, MAX_DRAFT_FACTS)
    requirements = _clean(requirements, MAX_DRAFT_REQUIREMENTS)
    compact_one_page = selected_type == "szse_report_brief"
    sources = _unique_sources(context, limit=4 if compact_one_page else 10)
    meeting_bundle = select_case_evidence(
        sources,
        subject=subject or facts or requirements,
        meeting_type=_meeting_type_for_document(selected_type),
        limit=8,
    ) if selected_type in MEETING_DOCUMENT_TYPES or selected_type in {"board_resolution", "shareholder_resolution"} else {}
    style_sources = _select_company_style_sources(
        sources,
        limit=2 if compact_one_page else 3,
        include_all_company_announcements=compact_one_page,
    ) if selected_type in {"periodic_report_notice", "szse_report_brief"} else []
    if style_sources:
        style_keys = {
            (str(item.get("document_id") or ""), str(item.get("chunk_id") or ""), str(item.get("title") or ""))
            for item in style_sources
        }
        for source in sources:
            key = (str(source.get("document_id") or ""), str(source.get("chunk_id") or ""), str(source.get("title") or ""))
            if key in style_keys:
                source["style_reference"] = True
    warnings: list[str] = []
    if not facts:
        if compact_one_page:
            warnings.append("未填写已确认事实，正文仅采用本地历史公告可核验摘要；具体金额、日期、主体和进展请人工确认。")
        else:
            warnings.append("未填写已确认事实，正文中的金额、日期、主体、表决结果和影响均保留待补充标记。")
    if not sources:
        warnings.append("当前知识库未检索到直接依据，请人工进一步确认。")
    else:
        warnings.append("以下内容是基于本地 RAG 检索结果生成的工作草稿，不能替代人工核验或正式签署。")
    if compact_one_page:
        warnings.append("本类型按一页纸汇报控制篇幅；超出一页的事实、附件和法规原文请另行列入附件。")

    display_title = title or _format_heading(str(template["heading"]), subject)
    lines: list[str] = [display_title, "", "【草稿状态】待人工核对，不得直接对外发布", ""]
    sections: list[dict[str, str]] = []

    def add_section(name: str, body: str) -> None:
        body = body.strip()
        lines.extend([f"{len(sections) + 1}、{name}", body, ""])
        sections.append({"title": name, "content": body})

    if selected_type in MEETING_DOCUMENT_TYPES or selected_type in {"board_resolution", "shareholder_resolution"}:
        proposal_text = _meeting_proposal_lines(meeting_bundle)
        speaking_text = _meeting_speaking_content(meeting_bundle, selected_type)
        term = meeting_bundle.get("meeting_term") or "本次会议"
        if selected_type == "meeting_sign_in":
            # A sign-in sheet is a practical working form, so keep unknown
            # fields explicit instead of borrowing people, dates or attendance
            # counts from historical cases.
            direct_meeting_source = next(
                (item for item in meeting_bundle.get("sources", []) if item.get("metadata", {}).get("direct_reference")),
                {},
            )
            direct_metadata = direct_meeting_source.get("metadata", {})
            meeting_date = direct_metadata.get("meeting_date") or "[待补充：以会议通知为准]"
            meeting_location = direct_metadata.get("meeting_location") or "[待补充：以会议通知为准]"
            add_section(
                "会议基本信息",
                f"会议名称/届次：{term}\n召开日期：{meeting_date}\n召开地点/方式：{meeting_location}",
            )
            add_section("会议议案/事项", proposal_text)
            add_section(
                "签到登记",
                "序号｜姓名/股东名称｜身份/职务｜签到时间｜签名/盖章\n"
                "1｜[待填写]｜[待填写]｜[待填写]｜________________\n"
                "2｜[待填写]｜[待填写]｜[待填写]｜________________\n"
                "3｜[待填写]｜[待填写]｜[待填写]｜________________\n"
                "4｜[待填写]｜[待填写]｜[待填写]｜________________\n"
                "5｜[待填写]｜[待填写]｜[待填写]｜________________",
            )
        elif selected_type == "meeting_agenda_summary":
            add_section(
                "会议基本信息",
                f"会议名称/届次：{term}\n召开日期：[待补充：以会议通知为准]\n召开地点/方式：[待补充：以会议通知为准]",
            )
            add_section("会议议程", proposal_text)
            add_section(
                "主要议案概述",
                "以下内容仅摘录本次导入文书中明确出现的议案和事项：\n" + proposal_text,
            )
            add_section(
                "审议提示",
                "请在会议前核对议案附件、审议权限、关联关系及回避表决要求；本稿未从参考文书识别的内容保留待补充。",
            )
        elif selected_type == "proxy_authorization":
            add_section(
                "会议基本信息",
                f"会议名称/届次：{term}\n召开日期：[待补充：以会议通知为准]\n召开地点/方式：[待补充：以会议通知为准]",
            )
            add_section(
                "委托人和受托人",
                "委托人姓名/名称：[待填写]\n委托人持有股份数及所持表决权：[待填写]\n受托人姓名/名称：[待填写]\n受托人身份证明/统一社会信用代码：[待填写]",
            )
            add_section(
                "授权事项",
                "委托人授权受托人代表本人/本单位出席上述会议，并就下列议案行使表决权：\n"
                + proposal_text
                + "\n授权期限、授权范围及是否允许转委托：[待填写；以公司提供的正式授权格式为准]",
            )
            add_section(
                "签署确认",
                "委托人签名/盖章：________________\n受托人签名：________________\n签署日期：____年__月__日\n提示：身份证明、持股证明和授权期限须在提交前核验。",
            )
        elif selected_type == "legal_opinion":
            add_section(
                "出具范围",
                f"就{term}涉及的会议文件和议案进行形式及程序核查。本稿仅根据本次导入参考文书整理，不构成执业律师正式法律意见。",
            )
            add_section(
                "会议程序核查",
                "会议召集主体、通知发出、召开时间地点、出席资格及表决程序：[待律师根据通知、签到、授权委托书和会议记录核验]。",
            )
            add_section(
                "议案及文件核查",
                "本次参考文书中识别的议案/事项：\n"
                + proposal_text
                + "\n关联关系、回避表决、审议权限、披露义务及附件完整性：[待结合公司章程、现行规则和完整底稿核验]。",
            )
            add_section(
                "结论与保留事项",
                "核查结论：[待补充：不得在缺少完整材料和律师复核的情况下填写“合法有效”等确定性结论]。\n保留事项：[待补充]。",
            )
        elif selected_type == "meeting_notice":
            add_section("会议时间、地点和方式", "会议届次、召开日期、地点、召开方式和通知送达情况：\n[待补充：请以正式会议通知和送达记录为准。]")
            add_section("会议议程", proposal_text)
            add_section("出席人员和注意事项", "出席人员、授权委托、回避事项和材料送达：\n[待补充：按公司章程、议事规则和本地案例核对。]")
        elif selected_type == "meeting_proposal":
            add_section("议案名称", proposal_text)
            add_section("议案背景及审议事项", "本议案根据本地案例中的同类议案整理，涉及事实、金额、日期、交易对方和附件的部分必须以本次底稿核对。\n" + proposal_text)
            add_section("合规审议重点", "核对审议主体和权限、议案材料完整性、关联关系/回避表决、风险影响、披露要求及是否需要提交更高层级会议。\n[待确认：本次适用制度和规则条款。]")
            add_section("提交审议", "请委员/董事对本议案逐项表决；通过后按会议层级和公司流程提交下一层级审议。")
        elif selected_type in {"meeting_record", "committee_record"}:
            add_section("会议召开情况", f"{term}的召开日期、地点、出席人员、主持人和记录人：\n[待补充：请以通知、签到及正式记录为准。]")
            add_section("议案说明", proposal_text)
            add_section("会议发言与讨论", speaking_text)
            add_section("表决及后续事项", "表决结果：同意、反对、弃权票数及回避票数[待补充]。\n后续事项：按本地案例记录是否提交董事会/股东会及后续披露安排。")
        elif selected_type == "general_manager_meeting":
            add_section("会议召开情况", f"{term}的召开日期、地点、出席人员、主持人和记录人：\n[待补充：请以会议通知、签到及正式记录为准。]")
            add_section("议案/事项说明", proposal_text)
            add_section("总经理审议意见", speaking_text + "\n\n总经理意见结论：同意/不同意/原则同意但需补充条件[待补充]。")
            add_section("决策及后续事项", "会议决定及责任部门、完成时限、资金/资源安排和风险控制措施：[待补充，以本次会议正式决定为准。]\n如事项超出总经理权限或触发董事会、股东会审议及信息披露要求，应按公司章程、内部制度和现行规则履行后续程序。")
        elif selected_type == "committee_resolution":
            add_section("会议召开情况", f"{term}的召开程序、出席人数和会议有效性：\n[待补充：以通知、签到及现行制度核对。]")
            add_section("议案审议情况", proposal_text + "\n\n委员讨论重点：\n" + speaking_text)
            add_section("表决结果", "各议案逐项记录同意、反对、弃权票数及回避情况：[待补充，不能从历史案例倒推出本次票数。]")
            add_section("后续程序", "审议通过的议案按委员会职责和公司流程提交董事会或其他有权机构审议；披露事项以正式公告和现行规则为准。")
        elif selected_type == "independent_director_resolution":
            add_section("会议召开情况", "会议通知、召开日期、出席独立董事人数及会议有效性：\n[待补充：请以正式通知和签到记录为准。]")
            add_section("议案审议情况", proposal_text)
            add_section("独立董事审核意见", speaking_text + "\n\n结论：同意/反对/弃权及是否同意提交董事会审议[待补充]。")
            add_section("提交董事会审议", "仅在独立董事专门会议形成明确意见后填写提交董事会的议案名称和日期。")
        elif selected_type == "ballot":
            add_section("议案清单", proposal_text)
            add_section("填表说明", "每一议案只能在同意、反对、弃权中选择一项；不选或多选无效。具体议案名称应与会议通知、议案正文和决议一致。")
            add_section("委员/董事签名", "签名：____________________    日期：____年__月__日")
        else:
            add_section("会议召开情况", "会议届次、召开日期、地点、方式、通知发出情况和出席人员：\n[待补充：请以会议通知、签到记录和正式会议记录为准。]")
            if selected_type == "board_resolution":
                add_section("会议审议情况", proposal_text + "\n\n议案讨论重点：\n" + speaking_text)
                add_section("回避表决情况", "[待确认：请依据关联方名录、董事任职关系和会议记录，填写应回避人员、回避原因及回避后的有效表决票数。]")
            else:
                add_section("议案审议情况", proposal_text + "\n\n议案讨论重点：\n" + speaking_text)
                add_section("律师见证情况", "[待补充：填写见证律师事务所、律师姓名及法律意见书结论。]")
            add_section("备查文件", "1. 经与会董事/股东签字确认的会议决议；\n2. [待补充：会议通知、议案、会议记录、表决票及其他应备查文件。]")
    elif selected_type == "related_transaction":
        add_section("关联交易概述", f"本事项：{subject or '[待补充]'}。\n交易金额、交易类别、连续十二个月累计金额及是否构成关联交易：[待补充]。")
        add_section("关联方基本情况", "关联方名称、统一社会信用代码/注册信息、股权及控制关系、与公司的关联关系形成原因和持续时间：[待补充]。")
        add_section("交易标的基本情况", "标的名称、权属、账面价值/评估价值、经营和财务状况、是否存在抵押质押或其他权利限制：[待补充]。")
        add_section("定价政策及定价依据", "定价方法、评估/审计依据、可比交易或市场价格、独立董事/审计委员会意见（如适用）：[待补充]。")
        add_section("交易协议主要内容", "交易主体、标的、金额、付款安排、交割条件、违约责任和生效条件：[待补充]。")
        add_section("关联交易目的和对公司的影响", "对主营业务、财务状况、经营成果和独立性的影响：[待补充，不使用无事实依据的确定性承诺。]")
        add_section("风险提示", "交易审批、定价、履约、资金、整合和市场环境风险：[待补充]。")
        add_section("备查文件", "董事会/股东会决议、关联方声明、审计或评估报告、交易协议及其他支持文件：[待补充]。")
    elif selected_type == "investment_announcement":
        add_section("对外投资概述", f"投资事项：{subject or '[待补充]'}。\n投资金额、资金来源、投资方式、持股比例和是否构成关联交易：[待补充]。")
        add_section("投资标的基本情况", "标的公司名称、注册资本、股权结构、主要业务、最近一年主要财务数据和权属情况：[待补充]。")
        add_section("投资协议主要内容", "投资主体、出资安排、治理结构、交割条件、违约责任和生效条件：[待补充]。")
        add_section("对外投资的目的和影响", "投资目的、对公司业务和财务的影响、资金来源及对公司现金流的影响：[待补充]。")
        add_section("风险分析", "政策、市场、经营、整合、投资回收和审批披露风险：[待补充]。")
        add_section("备查文件", "投资协议、董事会/股东会决议、审计或评估资料及其他支持文件：[待补充]。")
    elif selected_type == "periodic_report_notice":
        add_section("重要提示", "本报告应以经董事会审议通过的正式定期报告为准；审计意见、非标准意见说明、重大风险和前瞻性陈述：[待补充并与底稿核对]。")
        add_section("报告基本信息与公司基本情况", "公司名称、股票简称及代码、法定代表人、注册地址、联系方式和报告期：[待补充，沿用公司历期报告字段和顺序]。")
        add_section("主要财务数据和会计指标", "营业收入、归属于上市公司股东的净利润、扣除非经常性损益后的净利润、经营活动现金流量净额、基本每股收益、稀释每股收益和加权平均净资产收益率：[待补充并与 XLS/财务报表勾稽；同比百分比与百分点变化分开列示]。")
        add_section("管理层讨论与分析", "报告期经营概况、主营业务、行业情况、核心竞争力、主要经营模式、收入与成本变化、现金流和资产负债变化：[待补充，禁止照搬历史金额或结论]。")
        add_section("重大事项核对（重大事项）", "重大交易、关联交易、对外担保、对外投资、股份回购、股权激励、诉讼仲裁、承诺履行及其他应披露事项：[待补充，并根据本地 RAG 规则逐项核对]。")
        add_section("股份变动及股东情况", "股本结构、股东数量、前十名股东、控股股东和实际控制人变化：[待补充并与最新股权资料核对]。")
        add_section("优先股/债券（如适用）", "优先股、公司债、可转换债券或其他融资工具的发行、持有人、兑付和变动情况：[不适用或待补充]。")
        add_section("董事、监事和高级管理人员", "任职、持股、变动、薪酬和履职情况：[待补充并与董监高名册核对]。")
        add_section("公司治理", "股东会、董事会、监事会及专门委员会运作，内部控制和独立性情况：[待补充并与历史披露口径核对]。")
        add_section("财务报告", "审计报告、财务报表、财务报表附注及会计政策和会计估计变更：[待补充；所有金额、负号、合计和同比均须与财务底稿勾稽]。")
        add_section("备查文件", "经董事、监事和高级管理人员签署确认的定期报告、审计报告、财务报表及附注、董事会和监事会决议等：[待补充]。")
    elif selected_type == "szse_report_brief":
        structured_bundle = structured_bundle or {}
        def bundle_lines(key: str, fallback: str, limit: int = 3) -> str:
            values = structured_bundle.get(key) or []
            lines = []
            for fact in values[:limit]:
                value = _single_line(fact.get("fact_value") or fact.get("original_text"), 280)
                source = _single_line(fact.get("file_name") or "本地资料", 70)
                date_value = fact.get("publication_date") or fact.get("event_date") or "日期待确认"
                page = fact.get("page") or "页码待确认"
                if value:
                    lines.append(f"{value}（来源：{source}；披露日期：{date_value}；第{page}页）")
            return "；".join(lines) or fallback

        history_hint = _evidence_excerpt(
            style_sources,
            ("业务", "主营", "发展", "历史", "核心竞争力", "公司简介", "年度报告", "半年度报告"),
        ) or "当前知识库未检索到公司业务发展公告正文，请补充并核对公司历年公告。"
        restructuring_hint = _evidence_excerpt(
            style_sources,
            ("重组", "重大资产", "收购", "出售", "进展", "方案", "监管"),
        ) or "当前知识库未检索到重组进展公告正文，请补充并核对相关公告及监管沟通记录。"
        confirmed_facts = _single_line(facts, 900) if facts else bundle_lines(
            "business_history", history_hint,
        )
        current_business = bundle_lines(
            "business_current", "当前业务情况暂无足够结构化事实，请补充并核对最新定期报告。",
        )
        restructuring_facts = _single_line(facts, 900) if facts else bundle_lines(
            "restructuring_scheme", restructuring_hint,
        )
        latest_status = structured_bundle.get("current_status") or {}
        status_line = ""
        if latest_status:
            status_line = (
                f"最新可核验状态：{latest_status.get('status') or latest_status.get('fact_value')}（"
                f"{latest_status.get('file_name') or '本地资料'}；披露日期：{latest_status.get('publication_date') or '待确认'}；"
                f"第{latest_status.get('page') or '待确认'}页）。"
            )
        future_plan = bundle_lines(
            "future_plan", "后续工作安排暂无直接结构化事实，请以最新公告、会议材料和监管沟通记录补充。",
        )
        brief_scope = _single_line(subject or requirements, 180)
        scope_line = f"汇报范围：{brief_scope}\n" if brief_scope else ""
        add_section("公司业务发展历史简介", f"{scope_line}公司自成立/上市以来的业务发展、主营业务变化和重要阶段：{confirmed_facts}\n历史公告核验摘要：{history_hint}")
        add_section("公司当前业务情况", f"报告期/当前主要业务及经营情况：{current_business}")
        add_section("重大资产重组基本情况", f"重组事项及交易方案：{restructuring_facts}\n方案、交易对方、交易作价和募集配套资金等数字应以对应原始文件核对。")
        add_section("重大资产重组最新进展", f"{status_line or '最新可核验状态：当前知识库未检索到直接状态事实，请人工进一步确认。'}\n已完成工作、尚待完成程序及重大不确定性：{future_plan}")
        add_section("后续工作安排", "根据已检索的后续计划：" + future_plan + "；责任部门、时间表和监管沟通安排须以公司确认材料补充。")
    else:
        add_section("重要提示", "本公告所载事项和数据应以已确认底稿、正式决议及现行监管规则为依据。\n[待补充：审议主体、披露责任、风险提示和是否需要股东会审议。]")
        add_section("事项概述", f"事项名称：{subject or '[待补充]'}\n事项事实：{facts or '[待补充：金额、日期、交易对方、审批状态、资金来源等]'}")
        add_section("对公司的影响", "[待补充：对主营业务、财务状况、经营成果、股东权益和独立性的影响。不得使用未经依据支持的确定性表述。]")
        add_section("风险提示", "[待补充：根据已确认事实和本地 RAG 依据填写交易、履约、市场、财务、审批和披露风险。]")
        add_section("备查文件", "[待补充：董事会/股东会决议、协议、审计/评估报告、法律意见书和其他支持文件。]")

    if requirements:
        if compact_one_page:
            warnings.append(f"已按起草要求控制：{_single_line(requirements, 180)}")
        else:
            add_section("起草要求与待办", requirements)

    direct_reference = any(source.get("metadata", {}).get("direct_reference") for source in sources)
    lines.extend(["", "【本次导入文书依据】" if direct_reference else "【本地 RAG 起草依据】"])
    if sources:
        for index, source in enumerate(sources, 1):
            lines.extend([
                f"来源{index}：《{source.get('title') or '未命名资料'}》｜{source.get('document_type') or '资料'}｜{source.get('location')}",
                _single_line(source.get("content"), 220 if compact_one_page else 420),
                "",
            ])
    else:
        lines.extend(["当前知识库未检索到直接依据，请人工进一步确认。", ""])
    if style_sources and not compact_one_page:
        lines.extend([
            "【公司历史定期报告格式参考】",
            "以下历史公告仅用于章节顺序和披露表达参考，不得直接复制其中的金额、日期、主体或结论；正文事实必须以本次底稿核实。",
        ])
        for index, source in enumerate(style_sources, 1):
            lines.extend([
                f"格式参考{index}：《{source.get('title') or '公司历史定期报告'}》｜{source.get('location') or _source_location(source)}",
                _single_line(source.get("content"), 360),
                "",
            ])
    if compact_one_page:
        lines.extend(["【一页纸说明】", "正文按一页纸控制；引用来源仅保留精简摘要，完整法规、公告和支撑材料请在应用中打开核验。"])
    generation_note = (
        "本稿仅根据本次导入参考文书的原文片段整理；所有方括号内容都必须由人工以正式通知、签到和底稿核实后替换。"
        if direct_reference else
        f"本稿根据“{template['focus']}”的本地检索方向生成；所有方括号内容都必须由人工以原始底稿核实后替换。"
    )
    lines.extend(["【生成说明】", generation_note])

    return DraftResult(
        title=display_title,
        document_type=selected_type,
        content="\n".join(lines).strip(),
        sections=sections,
        sources=sources,
        warnings=warnings,
        style_sources=style_sources,
    )


def _set_font(run, size: float = 10.5, bold: bool = False, color: str = "222222") -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Calibri")
    rpr.rFonts.set(qn("w:hAnsi"), "Calibri")
    rpr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_docx(
    *,
    title: str,
    content: str,
    sources: list[dict[str, Any]],
    compact: bool = False,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55 if compact else 0.8)
    section.bottom_margin = Inches(0.55 if compact else 0.8)
    section.left_margin = Inches(0.65 if compact else 0.9)
    section.right_margin = Inches(0.65 if compact else 0.9)
    document.core_properties.title = title
    document.core_properties.subject = "基于本地 RAG 的工作草稿"
    document.core_properties.author = "信披智审"

    first = document.add_paragraph()
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_after = Pt(2 if compact else 4)
    run = first.add_run(title)
    _set_font(run, size=15 if compact else 18, bold=True, color="111827")
    state = document.add_paragraph()
    state.alignment = WD_ALIGN_PARAGRAPH.CENTER
    state.paragraph_format.space_after = Pt(6 if compact else 16)
    state_run = state.add_run("本地 RAG 工作草稿 · 待人工核对")
    _set_font(state_run, size=8.5 if compact else 10, color="9B1C1C")

    lines = content.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1 if compact else 2)
            index += 1
            continue
        # Convert the generated sign-in rows into a real editable Word table.
        # Other drafting content may use the full-width separator in prose, so
        # only the known sign-in header activates table mode.
        if line.startswith("序号｜姓名/股东名称｜身份/职务｜签到时间｜签名/盖章"):
            table_rows: list[list[str]] = []
            while index < len(lines) and "｜" in lines[index] and lines[index].strip():
                cells = [cell.strip() for cell in lines[index].split("｜")]
                if len(cells) == 5:
                    table_rows.append(cells)
                index += 1
            if table_rows:
                table = document.add_table(rows=0, cols=5)
                table.style = "Table Grid"
                table.autofit = True
                for row_index, cells in enumerate(table_rows):
                    row = table.add_row()
                    for cell, value in zip(row.cells, cells):
                        paragraph = cell.paragraphs[0]
                        paragraph.paragraph_format.space_after = Pt(0)
                        run = paragraph.add_run(value)
                        _set_font(run, size=8.5 if compact else 9.5, bold=row_index == 0)
                continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2 if compact else 5)
        p.paragraph_format.line_spacing = 1.0 if compact else 1.12
        if re.match(r"^\d+、", line) or line.startswith("【"):
            run = p.add_run(line)
            _set_font(run, size=9.75 if compact else 11, bold=True, color="1F4D78" if line.startswith("【") is False else "9B1C1C")
        elif line.startswith("来源"):
            run = p.add_run(line)
            _set_font(run, size=8.5 if compact else 9.5, bold=True, color="2E74B5")
        else:
            run = p.add_run(line)
            _set_font(run, size=9.25 if compact else 10.5)
        index += 1
    if sources:
        note = document.add_paragraph()
        note.paragraph_format.space_before = Pt(4 if compact else 8)
        note.paragraph_format.space_after = Pt(0)
        run = note.add_run(f"共引用 {len(sources)} 个本地 RAG 来源；请在应用中打开来源核验正文和有效版本。")
        _set_font(run, size=8 if compact else 9, color="5B6573")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


__all__ = ["DOCUMENT_TYPES", "DraftResult", "build_draft", "build_docx"]
