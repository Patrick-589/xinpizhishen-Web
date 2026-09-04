from __future__ import annotations

import html
import json
import os
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any


STATUS_LABELS = {
    "failed": "明确错误",
    "needs_review": "需要人工确认",
    "source_issue": "数据来源问题",
    "passed": "已核对无误",
}
SEVERITY_LABELS = {"high": "严重", "medium": "重要", "low": "一般", "none": "通过"}


def report_title(payload: dict[str, Any]) -> str:
    period = payload.get("periodRegistry", {})
    year = str(period.get("period_end") or "")[:4]
    report_type = str(period.get("report_type") or "定期")
    if year and report_type:
        return f"《{year}年{report_type}报告财务数据及表格勾稽审核报告》"
    return "《定期报告财务数据及表格勾稽审核报告》"


def _location_text(check: dict[str, Any]) -> str:
    values: list[str] = []
    for item in check.get("source_locations", []):
        page = f"第{item.get('page')}页" if item.get("page") else "页码待核"
        sheet = item.get("sheet_name") or item.get("sheetName")
        row_number = item.get("spreadsheet_row") or item.get("spreadsheetRow")
        table = item.get("table_name") or "正文"
        row = item.get("row_label") or ""
        column = item.get("column_label") or ""
        spreadsheet = " / ".join(
            part for part in (f"工作表 {sheet}" if sheet else "", f"第{row_number}行" if row_number else "") if part
        )
        value = " / ".join(part for part in (page, spreadsheet, table, row, column) if part)
        if value and value not in values:
            values.append(value)
    return "；".join(values) or "未形成精确位置，请人工定位原文。"


def _format_value(value: Any) -> str:
    if value is None or value == "":
        return "-"
    if isinstance(value, dict):
        return "；".join(f"{key}={_format_value(item)}" for key, item in value.items())
    return str(value)


def _source_excerpt(check: dict[str, Any]) -> str:
    values: list[str] = []
    for item in check.get("source_locations", []):
        source_text = str(item.get("source_text") or "").strip()
        if source_text and source_text not in values:
            values.append(source_text)
    return "\n".join(values) or "-"


def _case_excerpt(check: dict[str, Any]) -> str:
    values: list[str] = []
    for item in check.get("case_references", []):
        if not isinstance(item, dict):
            continue
        annotation = str(item.get("annotation") or "").strip()
        pattern = str(item.get("pattern") or "").strip()
        source = str(item.get("source") or "").strip()
        text = " / ".join(part for part in (pattern, annotation, source) if part)
        if text and text not in values:
            values.append(text)
    return "\n".join(values) or "-"


def build_html(payload: dict[str, Any]) -> str:
    title = report_title(payload)
    summary = payload.get("summary", {})
    issues = [item for item in payload.get("checks", []) if item.get("status") != "passed"]
    passed = [item for item in payload.get("checks", []) if item.get("status") == "passed"]
    category_counts = Counter(item.get("category") or "其他" for item in payload.get("checks", []))
    issue_markup = "".join(
        f"""
        <article class="issue {html.escape(str(item.get('severity') or 'medium'))}">
          <header><span>{html.escape(STATUS_LABELS.get(str(item.get('status')), str(item.get('status'))))}</span><strong>{html.escape(str(item.get('title') or ''))}</strong></header>
          <dl>
            <div><dt>原文位置</dt><dd>{html.escape(_location_text(item))}</dd></div>
             <div><dt>判断说明</dt><dd>{html.escape(str(item.get('explanation') or ''))}</dd></div>
            <div><dt>原始表格行</dt><dd>{html.escape(_source_excerpt(item))}</dd></div>
             <div><dt>计算公式</dt><dd>{html.escape(str(item.get('formula') or '-'))}</dd></div>
            <div><dt>计算输入</dt><dd>{html.escape(_format_value(item.get('inputs') or {}))}</dd></div>
            <div><dt>原始列示</dt><dd>{html.escape(_format_value(item.get('reported_value')))}</dd></div>
            <div><dt>复算结果</dt><dd>{html.escape(_format_value(item.get('calculated_value')))}</dd></div>
            <div><dt>差异</dt><dd>{html.escape(_format_value(item.get('difference')))}</dd></div>
            <div><dt>已确认案例</dt><dd>{html.escape(_case_excerpt(item))}</dd></div>
            <div><dt>建议</dt><dd>{html.escape(str(item.get('suggestion') or ''))}</dd></div>
          </dl>
        </article>"""
        for item in issues
    ) or "<p class='empty'>未发现自动检查异常。</p>"
    passed_markup = "".join(
        f"<li><strong>{html.escape(str(item.get('title') or ''))}</strong><span>{html.escape(str(item.get('formula') or item.get('explanation') or ''))}</span></li>"
        for item in passed
    )
    category_markup = "".join(
        f"<li><span>{html.escape(category)}</span><strong>{count}</strong></li>"
        for category, count in sorted(category_counts.items())
    )
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title><style>
*{{box-sizing:border-box}}body{{font-family:'Microsoft YaHei','PingFang SC',sans-serif;margin:0;color:#18181b;background:#f3f3f4;font-size:14px;line-height:1.65}}main{{max-width:1120px;margin:24px auto;background:#fff;padding:36px 44px;border-top:5px solid #b4232d}}h1{{font-size:25px;margin:0 0 8px}}h2{{font-size:18px;border-left:4px solid #b4232d;padding-left:10px;margin-top:32px}}.meta{{color:#69696f}}.kpis{{display:grid;grid-template-columns:repeat(5,1fr);gap:8px;margin:24px 0}}.kpi{{border:1px solid #dedee1;padding:12px}}.kpi span{{display:block;color:#69696f;font-size:12px}}.kpi strong{{font-size:23px}}.scope{{display:grid;grid-template-columns:repeat(3,1fr);gap:7px;padding:0;list-style:none}}.scope li{{border-bottom:1px solid #dedee1;display:flex;justify-content:space-between;padding:7px}}.issue{{border:1px solid #dedee1;border-left:4px solid #9a5a00;margin:12px 0;padding:14px}}.issue.high{{border-left-color:#b4232d}}.issue header{{display:flex;gap:10px;align-items:flex-start}}.issue header span{{background:#fff0f1;color:#b4232d;font-size:11px;padding:2px 6px;white-space:nowrap}}dl{{margin:10px 0 0}}dl div{{display:grid;grid-template-columns:84px 1fr;border-top:1px solid #ececee;padding:7px 0}}dt{{color:#69696f}}dd{{margin:0;word-break:break-word}}.passed{{padding:0;list-style:none;columns:2;column-gap:24px}}.passed li{{break-inside:avoid;border-bottom:1px solid #ececee;padding:7px 0;display:grid}}.passed span{{color:#69696f;font-size:12px}}footer{{border-top:1px solid #dedee1;color:#69696f;margin-top:30px;padding-top:12px}}@media(max-width:720px){{main{{margin:0;padding:20px 14px}}.kpis{{grid-template-columns:1fr 1fr}}.scope{{grid-template-columns:1fr}}.passed{{columns:1}}dl div{{grid-template-columns:1fr}}}}
</style></head><body><main>
<h1>{html.escape(title)}</h1><p class="meta">文件：{html.escape(str(payload.get('filename') or ''))}<br>引擎：{html.escape(str(payload.get('engineVersion') or ''))} · 全程本地确定性计算 · 未调用外部 AI</p>
<section class="kpis"><div class="kpi"><span>总检查项</span><strong>{summary.get('total', 0)}</strong></div><div class="kpi"><span>通过</span><strong>{summary.get('passed', 0)}</strong></div><div class="kpi"><span>异常</span><strong>{summary.get('exceptions', 0)}</strong></div><div class="kpi"><span>严重</span><strong>{summary.get('severe', 0)}</strong></div><div class="kpi"><span>待人工确认</span><strong>{summary.get('needs_review', 0)}</strong></div></section>
<h2>检查范围</h2><ul class="scope">{category_markup}</ul>
<h2>异常与待复核事项</h2>{issue_markup}
<h2>已核对无误清单</h2><ol class="passed">{passed_markup}</ol>
<footer>本报告由本地文档解析、结构化事实提取和 Decimal 确定性规则生成。RAG/LLM 不参与加减、百分比或勾稽计算。系统结论用于辅助复核，不替代会计师、证券事务人员和法务人员的专业判断。</footer>
</main></body></html>"""


def build_docx(payload: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    masthead = document.add_paragraph()
    masthead.paragraph_format.space_after = Pt(10)
    run = masthead.add_run("信披智审  |  财务数据与表格勾稽审计")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 35, 45)
    run.font.size = Pt(10)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    title_run = title.add_run(report_title(payload))
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.name = "Microsoft YaHei"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"文件：{payload.get('filename', '')}  |  引擎：{payload.get('engineVersion', '')}  |  全程本地确定性计算")
    meta_run.font.color.rgb = RGBColor(105, 105, 111)
    meta_run.font.size = Pt(8)

    summary = payload.get("summary", {})
    summary_table = document.add_table(rows=2, cols=5)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_table.autofit = False
    labels = ["总检查项", "通过", "异常", "严重", "待人工确认"]
    values = [summary.get("total", 0), summary.get("passed", 0), summary.get("exceptions", 0), summary.get("severe", 0), summary.get("needs_review", 0)]
    for column, (label, value) in enumerate(zip(labels, values)):
        for row in range(2):
            cell = summary_table.cell(row, column)
            cell.width = Cm(3.25)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = label if row == 0 else str(value)
            run = paragraph.add_run(text)
            run.bold = row == 1
            run.font.size = Pt(8 if row == 0 else 15)
            if row == 1 and column in {2, 3}:
                run.font.color.rgb = RGBColor(180, 35, 45)

    def heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(24, 24, 27)

    category_counts = Counter(item.get("category") or "其他" for item in payload.get("checks", []))
    heading("一、检查范围")
    document.add_paragraph("；".join(f"{category} {count} 项" for category, count in sorted(category_counts.items())))

    heading("二、异常与待复核事项")
    issues = [item for item in payload.get("checks", []) if item.get("status") != "passed"]
    if not issues:
        document.add_paragraph("未发现自动检查异常。")
    for index, item in enumerate(issues, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        run = paragraph.add_run(f"{index}. [{STATUS_LABELS.get(str(item.get('status')), item.get('status'))}] {item.get('title', '')}")
        run.bold = True
        run.font.color.rgb = RGBColor(180, 35, 45) if item.get("severity") == "high" else RGBColor(154, 90, 0)
        detail = document.add_table(rows=10, cols=2)
        detail.alignment = WD_TABLE_ALIGNMENT.CENTER
        detail.autofit = False
        rows = [
            ("原文位置", _location_text(item)),
            ("判断说明", item.get("explanation") or "-"),
            ("原始表格行", _source_excerpt(item)),
            ("计算公式", item.get("formula") or "-"),
            ("计算输入", _format_value(item.get("inputs") or {})),
            ("原始列示", _format_value(item.get("reported_value"))),
            ("复算结果", _format_value(item.get("calculated_value"))),
            ("差异", _format_value(item.get("difference"))),
            ("已确认案例", _case_excerpt(item)),
            ("修改建议", item.get("suggestion") or "-"),
        ]
        for row_index, (label, value) in enumerate(rows):
            left, right = detail.rows[row_index].cells
            left.width, right.width = Cm(2.6), Cm(13.6)
            left.text = str(label)
            right.text = str(value)
            left.paragraphs[0].runs[0].bold = True
            left.paragraphs[0].runs[0].font.size = Pt(8)
            right.paragraphs[0].runs[0].font.size = Pt(8)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F3F3F4")
            left._tc.get_or_add_tcPr().append(shading)

    heading("三、已核对无误清单")
    passed = [item for item in payload.get("checks", []) if item.get("status") == "passed"]
    for index, item in enumerate(passed, start=1):
        paragraph = document.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Cm(0.3)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.add_run(f"{index}. {item.get('title', '')}").bold = True
        paragraph.add_run(f"  {item.get('formula') or item.get('explanation') or ''}")

    heading("四、方法与责任边界")
    document.add_paragraph(
        "本报告由本地文档解析、结构化事实提取、单位标准化和 Decimal 确定性规则生成。"
        "RAG/LLM 不参与加减、百分比或勾稽计算；外部模型未被调用。系统结论用于辅助复核，"
        "不替代会计师、证券事务人员和法务人员的专业判断。"
    )
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("信披智审 · 财务勾稽审核报告 · 第 ").font.size = Pt(8)
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)
        footer.add_run(" 页").font.size = Pt(8)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def _pdf_font() -> tuple[str, str | None]:
    windows_root = Path(os.environ.get("WINDIR", r"C:\Windows"))
    candidates = [windows_root / "Fonts" / "msyh.ttc", windows_root / "Fonts" / "simsun.ttc"]
    for candidate in candidates:
        if candidate.is_file():
            return "XZS-CJK", str(candidate)
    return "Helvetica", None


def build_pdf(payload: dict[str, Any], output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name, font_path = _pdf_font()
    if font_path:
        pdfmetrics.registerFont(TTFont(font_name, font_path, subfontIndex=0))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output_path), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
        title=report_title(payload), author="信披智审",
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("XZSNormal", parent=styles["Normal"], fontName=font_name, fontSize=8.2, leading=12, textColor=colors.HexColor("#18181b"))
    title_style = ParagraphStyle("XZSTitle", parent=normal, fontSize=16, leading=22, alignment=TA_CENTER, spaceAfter=7)
    meta_style = ParagraphStyle("XZSMeta", parent=normal, fontSize=7.4, textColor=colors.HexColor("#69696f"), alignment=TA_CENTER, spaceAfter=12)
    heading_style = ParagraphStyle("XZSHeading", parent=normal, fontSize=11, leading=16, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#18181b"))
    small = ParagraphStyle("XZSSmall", parent=normal, fontSize=7.2, leading=10)
    story: list[Any] = [
        Paragraph(report_title(payload), title_style),
        Paragraph(f"文件：{html.escape(str(payload.get('filename') or ''))}<br/>引擎：{html.escape(str(payload.get('engineVersion') or ''))} · 全程本地确定性计算", meta_style),
    ]
    summary = payload.get("summary", {})
    labels = ["总检查项", "通过", "异常", "严重", "待人工确认"]
    values = [summary.get("total", 0), summary.get("passed", 0), summary.get("exceptions", 0), summary.get("severe", 0), summary.get("needs_review", 0)]
    summary_table = Table([[Paragraph(item, small) for item in labels], [Paragraph(f"<b>{item}</b>", normal) for item in values]], colWidths=[34 * mm] * 5)
    summary_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dedee1")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f3f4")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    category_counts = Counter(item.get("category") or "其他" for item in payload.get("checks", []))
    story.extend([
        summary_table,
        Paragraph("一、检查范围", heading_style),
        Paragraph("；".join(f"{category} {count} 项" for category, count in sorted(category_counts.items())), normal),
        Paragraph("二、异常与待复核事项", heading_style),
    ])
    issues = [item for item in payload.get("checks", []) if item.get("status") != "passed"]
    for index, item in enumerate(issues, start=1):
        heading = Paragraph(
            f"<b>{index}. [{html.escape(STATUS_LABELS.get(str(item.get('status')), str(item.get('status'))))}] {html.escape(str(item.get('title') or ''))}</b>",
            normal,
        )
        rows = [
            ("原文位置", _location_text(item)), ("判断说明", item.get("explanation") or "-"),
            ("原始表格行", _source_excerpt(item)), ("计算公式", item.get("formula") or "-"),
            ("计算输入", _format_value(item.get("inputs") or {})), ("原始列示", _format_value(item.get("reported_value"))),
            ("复算结果", _format_value(item.get("calculated_value"))), ("差异", _format_value(item.get("difference"))),
            ("已确认案例", _case_excerpt(item)),
            ("修改建议", item.get("suggestion") or "-"),
        ]
        data = [[Paragraph(label, small), Paragraph(html.escape(str(value)), small)] for label, value in rows]
        detail = Table(data, colWidths=[25 * mm, 144 * mm], repeatRows=0)
        detail.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#dedee1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f3f4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.extend([KeepTogether([heading, Spacer(1, 3), detail]), Spacer(1, 7)])
    story.append(Paragraph("三、已核对无误清单", heading_style))
    passed = [item for item in payload.get("checks", []) if item.get("status") == "passed"]
    for index, item in enumerate(passed, start=1):
        story.append(Paragraph(f"{index}. <b>{html.escape(str(item.get('title') or ''))}</b>　{html.escape(str(item.get('formula') or ''))}", small))
        story.append(Spacer(1, 2))
    story.extend([
        Paragraph("四、方法与责任边界", heading_style),
        Paragraph("本报告由本地文档解析、结构化事实提取、单位标准化和 Decimal 确定性规则生成。RAG/LLM 不参与加减、百分比或勾稽计算；外部模型未被调用。系统结论用于辅助复核，不替代专业判断。", normal),
    ])

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7)
        canvas.setFillColor(colors.HexColor("#69696f"))
        canvas.drawCentredString(A4[0] / 2, 8 * mm, f"信披智审 · 第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


def generate_reports(payload: dict[str, Any], output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = report_title(payload).strip("《》").replace("/", "_")
    html_path = output_dir / f"{stem}.html"
    docx_path = output_dir / f"{stem}.docx"
    pdf_path = output_dir / f"{stem}.pdf"
    html_path.write_text(build_html(payload), encoding="utf-8")
    build_docx(payload, docx_path)
    build_pdf(payload, pdf_path)
    manifest = {
        "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "engineVersion": payload.get("engineVersion"),
        "fileHash": payload.get("fileHash"),
        "reports": {"html": html_path.name, "docx": docx_path.name, "pdf": pdf_path.name},
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"html": html_path, "docx": docx_path, "pdf": pdf_path}


__all__ = ["build_docx", "build_html", "build_pdf", "generate_reports", "report_title"]
