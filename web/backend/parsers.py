from __future__ import annotations

import csv
import io
import posixpath
import re
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

from .ocr import LocalOCRError, LocalOCRService
from .windows_ifilter import WindowsIFilterError, extract_legacy_word_text


WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OFFICE_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PARSER_VERSION = 6
MAX_TEXT_CHARACTERS = 600_000
ALLOWED_EXTENSIONS = {".doc", ".docx", ".pdf", ".txt", ".md", ".json", ".csv", ".xls", ".xlsx"}
DOCX_SYMBOL_MAP = {
    ("wingdings 2", "F052"): "☑",
}


class DocumentParseError(Exception):
    pass


@dataclass(frozen=True)
class ParsedBlock:
    text: str
    page: int
    section: str
    paragraph_index: int
    block_type: str = "paragraph"
    source_start: int = 0
    source_end: int = 0
    csv_row_number: int | None = None
    source_row_numbers: tuple[int, ...] = ()
    sheet_name: str = ""


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    file_type: str
    page_count: int
    blocks: list[ParsedBlock]
    scanned_suspected: bool = False
    ocr_used: bool = False
    ocr_pages: tuple[int, ...] = ()
    ocr_engine: str = ""

    @property
    def paragraphs(self) -> list[str]:
        return [block.text for block in self.blocks]

    @property
    def character_count(self) -> int:
        return sum(len(block.text) for block in self.blocks)

    @property
    def headings(self) -> list[dict[str, object]]:
        values = []
        for block in self.blocks:
            if block.block_type == "heading" or block.paragraph_index == 0:
                values.append({
                    "text": block.text[:60],
                    "paragraphIndex": block.paragraph_index,
                    "page": block.page,
                    "section": block.section,
                })
            if len(values) >= 20:
                break
        return values


HEADING_PATTERN = re.compile(
    r"^(?:第[一二三四五六七八九十百0-9]+[编章节条]|[一二三四五六七八九十]+、|（[一二三四五六七八九十0-9]+）|[0-9]+(?:\.[0-9]+)*[、.． ])"
)


def normalize_text(value: str) -> str:
    value = re.sub(r"[\ud800-\udfff]", "\ufffd", value)
    return re.sub(r"[\t\xa0 ]+", " ", value).strip()


def looks_like_heading(text: str) -> bool:
    return len(text) <= 70 and bool(HEADING_PATTERN.search(text))


def _docx_text(element: ElementTree.Element) -> str:
    values: list[str] = []
    for node in element.iter():
        local_name = node.tag.rsplit("}", 1)[-1]
        if local_name == "t":
            values.append(node.text or "")
        elif local_name == "sym":
            font = node.get(f"{{{WORD_NAMESPACE}}}font", "").strip().casefold()
            char = node.get(f"{{{WORD_NAMESPACE}}}char", "").strip().upper()
            mapped = DOCX_SYMBOL_MAP.get((font, char))
            if mapped:
                values.append(mapped)
        elif local_name == "tab":
            values.append("\t")
        elif local_name in {"br", "cr"}:
            values.append("\n")
    return "".join(values)


def _docx_span_value(
    element: ElementTree.Element,
    path: str,
    namespace: dict[str, str],
    *,
    default: int = 1,
    minimum: int = 1,
) -> int:
    node = element.find(path, namespace)
    raw_value = node.get(f"{{{WORD_NAMESPACE}}}val", str(default)) if node is not None else str(default)
    try:
        return max(minimum, min(int(raw_value), 100))
    except ValueError:
        return default


def _docx_table_rows(
    table: ElementTree.Element,
    namespace: dict[str, str],
) -> list[str]:
    """Expand Word table grid spans so every row uses stable logical columns."""
    grid_width = len(table.findall("./w:tblGrid/w:gridCol", namespace))
    rows: list[list[str]] = []
    for row in table.findall("./w:tr", namespace):
        cells = [""] * _docx_span_value(
            row, "./w:trPr/w:gridBefore", namespace, default=0, minimum=0,
        )
        for cell in row.findall("./w:tc", namespace):
            cell_text = _docx_text(cell).strip()
            grid_span = _docx_span_value(cell, "./w:tcPr/w:gridSpan", namespace)
            cells.append(cell_text)
            cells.extend([""] * (grid_span - 1))
        cells.extend([""] * _docx_span_value(
            row, "./w:trPr/w:gridAfter", namespace, default=0, minimum=0,
        ))
        if grid_width > len(cells):
            cells.extend([""] * (grid_width - len(cells)))
        if any(cells):
            rows.append(cells)
    return [" | ".join(cells) for cells in rows]


def _iter_docx_body_blocks(
    element: ElementTree.Element,
    paragraph_tag: str,
    table_tag: str,
    alt_chunk_tag: str,
):
    """Yield visible body parts in document order, including content controls."""
    for child in list(element):
        if child.tag in {paragraph_tag, table_tag, alt_chunk_tag}:
            yield child
            continue
        yield from _iter_docx_body_blocks(child, paragraph_tag, table_tag, alt_chunk_tag)


def _normalize_html_text(value: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.replace("\r", "").split("\n")]
    return "\n".join(line for line in lines if line)


class _AltChunkHTMLParser(HTMLParser):
    """Convert an Alternative Format Import Part into ordered text/table blocks."""

    PARAGRAPH_TAGS = {"p", "li", "blockquote", "pre"}
    HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
    IGNORED_TAGS = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, str]] = []
        self._ignored_depth = 0
        self._table_depth = 0
        self._table_rows: list[list[tuple[str, int, int]]] = []
        self._current_row: list[tuple[str, int, int]] | None = None
        self._cell_parts: list[str] | None = None
        self._cell_colspan = 1
        self._cell_rowspan = 1
        self._paragraph_tag = ""
        self._paragraph_parts: list[str] = []
        self._loose_parts: list[str] = []

    @staticmethod
    def _span(attributes: list[tuple[str, str | None]], name: str) -> int:
        raw = dict(attributes).get(name) or "1"
        try:
            return max(1, min(int(raw), 100))
        except ValueError:
            return 1

    def _emit_text(self, raw_text: str, block_type: str = "paragraph") -> None:
        text = _normalize_html_text(raw_text)
        if text:
            if block_type != "table" and looks_like_heading(text):
                block_type = "heading"
            self.blocks.append((text, block_type))

    def _flush_paragraph(self) -> None:
        if self._paragraph_tag:
            block_type = "heading" if self._paragraph_tag in self.HEADING_TAGS else "paragraph"
            self._emit_text("".join(self._paragraph_parts), block_type)
        self._paragraph_tag = ""
        self._paragraph_parts = []

    def _flush_loose_text(self) -> None:
        self._emit_text("".join(self._loose_parts))
        self._loose_parts = []

    @staticmethod
    def _expand_table(rows: list[list[tuple[str, int, int]]]) -> list[list[str]]:
        expanded: list[list[str]] = []
        occupied: set[tuple[int, int]] = set()
        for row_index, cells in enumerate(rows):
            output: list[str] = []
            column = 0
            for text, colspan, rowspan in cells:
                while (row_index, column) in occupied:
                    output.append("")
                    column += 1
                while len(output) < column:
                    output.append("")
                output.append(text)
                for offset in range(1, colspan):
                    output.append("")
                    occupied.add((row_index, column + offset))
                for future_row in range(row_index + 1, row_index + rowspan):
                    for offset in range(colspan):
                        occupied.add((future_row, column + offset))
                column += colspan
            occupied_columns = [col for row, col in occupied if row == row_index]
            if occupied_columns:
                target_width = max(occupied_columns) + 1
                while len(output) < target_width:
                    output.append("")
            expanded.append(output)
        return expanded

    def _emit_table(self) -> None:
        rows = [row for row in self._expand_table(self._table_rows) if any(row)]
        if rows:
            self._emit_text("\n".join(" | ".join(row) for row in rows), "table")
        self._table_rows = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.casefold()
        if tag in self.IGNORED_TAGS:
            self._ignored_depth += 1
            return
        if self._ignored_depth:
            return
        if tag == "table":
            if self._table_depth == 0:
                self._flush_paragraph()
                self._flush_loose_text()
                self._table_rows = []
            elif self._cell_parts is not None:
                self._cell_parts.append("\n")
            self._table_depth += 1
            return
        if self._table_depth:
            if self._table_depth == 1 and tag == "tr":
                self._current_row = []
            elif self._table_depth == 1 and tag in {"td", "th"}:
                self._cell_parts = []
                self._cell_colspan = self._span(attrs, "colspan")
                self._cell_rowspan = self._span(attrs, "rowspan")
            elif tag == "br" and self._cell_parts is not None:
                self._cell_parts.append("\n")
            elif tag in self.PARAGRAPH_TAGS and self._cell_parts:
                self._cell_parts.append("\n")
            return
        if tag in self.PARAGRAPH_TAGS or tag in self.HEADING_TAGS:
            self._flush_paragraph()
            self._flush_loose_text()
            self._paragraph_tag = tag
            self._paragraph_parts = []
        elif tag == "br":
            (self._paragraph_parts if self._paragraph_tag else self._loose_parts).append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if self._ignored_depth:
            if tag in self.IGNORED_TAGS:
                self._ignored_depth -= 1
            return
        if self._table_depth:
            if self._table_depth == 1 and tag in {"td", "th"} and self._cell_parts is not None:
                text = _normalize_html_text("".join(self._cell_parts))
                if self._current_row is not None:
                    self._current_row.append((text, self._cell_colspan, self._cell_rowspan))
                self._cell_parts = None
            elif self._table_depth == 1 and tag == "tr" and self._current_row is not None:
                self._table_rows.append(self._current_row)
                self._current_row = None
            elif tag in self.PARAGRAPH_TAGS and self._cell_parts is not None:
                self._cell_parts.append("\n")
            if tag == "table":
                self._table_depth -= 1
                if self._table_depth == 0:
                    self._emit_table()
            return
        if tag == self._paragraph_tag:
            self._flush_paragraph()
        elif tag in {"body", "html"}:
            self._flush_loose_text()

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        if self._table_depth:
            if self._cell_parts is not None:
                self._cell_parts.append(data)
        elif self._paragraph_tag:
            self._paragraph_parts.append(data)
        elif data.strip():
            self._loose_parts.append(data)

    def close(self) -> None:
        super().close()
        self._flush_paragraph()
        self._flush_loose_text()
        if self._table_rows:
            self._emit_table()


def _decode_html_part(content: bytes) -> str:
    if content.startswith(b"\xef\xbb\xbf"):
        return content.decode("utf-8-sig")
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    sample = content[:4096].decode("latin-1", errors="ignore")
    match = re.search(r"charset\s*=\s*['\"]?([A-Za-z0-9._-]+)", sample, flags=re.IGNORECASE)
    encodings = [match.group(1)] if match else []
    encodings.extend(["utf-8", "gb18030"])
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except (LookupError, UnicodeDecodeError):
            continue
    return content.decode("utf-8", errors="replace")


def _parse_alt_chunk(content: bytes, target: str) -> list[tuple[str, str]]:
    if Path(target).suffix.casefold() in {".mht", ".mhtml"} or content.lstrip().startswith(b"MIME-Version:"):
        message = BytesParser(policy=policy.default).parsebytes(content)
        html_part = next(
            (part for part in message.walk() if part.get_content_type() == "text/html"),
            None,
        )
        if html_part is None:
            return []
        html = html_part.get_content()
    else:
        html = _decode_html_part(content)
    parser = _AltChunkHTMLParser()
    parser.feed(str(html))
    parser.close()
    return parser.blocks


def _load_alt_chunks(archive: zipfile.ZipFile) -> dict[str, list[tuple[str, str]]]:
    try:
        relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
    except (KeyError, ElementTree.ParseError):
        return {}
    values: dict[str, list[tuple[str, str]]] = {}
    for relationship in relationships.iter():
        if relationship.tag.rsplit("}", 1)[-1] != "Relationship":
            continue
        relationship_id = relationship.get("Id", "")
        target = relationship.get("Target", "")
        relationship_type = relationship.get("Type", "")
        if not relationship_id or not target or not relationship_type.casefold().endswith("/afchunk"):
            continue
        if relationship.get("TargetMode", "").casefold() == "external":
            continue
        archive_path = (
            posixpath.normpath(target.lstrip("/"))
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("word", target))
        )
        if archive_path.startswith("../") or archive_path not in archive.namelist():
            continue
        try:
            values[relationship_id] = _parse_alt_chunk(archive.read(archive_path), archive_path)
        except (UnicodeError, ValueError):
            continue
    return values


def _bounded_blocks(values: list[tuple]) -> list[ParsedBlock]:
    blocks: list[ParsedBlock] = []
    total = 0
    current_section = "正文"
    for value in values:
        raw_text, page, block_type, explicit_section = value[:4]
        source_row_numbers = tuple(value[4] or ()) if len(value) > 4 else ()
        sheet_name = str(value[5] or "") if len(value) > 5 else ""
        text = normalize_text(raw_text)
        if not text:
            continue
        remaining = MAX_TEXT_CHARACTERS - total
        if remaining <= 0:
            break
        text = text[:remaining]
        if explicit_section:
            current_section = explicit_section
        elif looks_like_heading(text):
            current_section = text[:100]
            block_type = "heading"
        csv_match = re.match(r"CSV行(\d+)", text)
        blocks.append(ParsedBlock(
            text,
            max(1, page),
            current_section,
            len(blocks),
            block_type,
            total,
            total + len(text),
            int(csv_match.group(1)) if csv_match else None,
            source_row_numbers,
            sheet_name,
        ))
        total += len(text)
    return blocks


def parse_docx(content: bytes, filename: str) -> ParsedDocument:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
            alt_chunks = _load_alt_chunks(archive)
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as error:
        raise DocumentParseError("DOCX 文件结构无效或已损坏。") from error

    ns = {"w": WORD_NAMESPACE}
    page = 1
    values: list[tuple[str, int, str, str]] = []
    body = root.find(".//w:body", ns)
    if body is None:
        raise DocumentParseError("DOCX 中没有可识别的正文。")

    paragraph_tag = f"{{{WORD_NAMESPACE}}}p"
    table_tag = f"{{{WORD_NAMESPACE}}}tbl"
    alt_chunk_tag = f"{{{WORD_NAMESPACE}}}altChunk"
    relationship_id_tag = f"{{{OFFICE_RELATIONSHIP_NAMESPACE}}}id"
    for element in _iter_docx_body_blocks(body, paragraph_tag, table_tag, alt_chunk_tag):
        if element.tag == paragraph_tag:
            text = _docx_text(element)
            style_node = element.find("./w:pPr/w:pStyle", ns)
            style_value = style_node.get(f"{{{WORD_NAMESPACE}}}val", "") if style_node is not None else ""
            block_type = "heading" if style_value.lower().startswith("heading") or looks_like_heading(text) else "paragraph"
            values.append((text, page, block_type, text if block_type == "heading" else ""))
            page_breaks = element.findall(".//w:br[@w:type='page']", ns)
            page += len(page_breaks)
        elif element.tag == table_tag:
            rows = _docx_table_rows(element, ns)
            if rows:
                values.append(("\n".join(rows), page, "table", ""))
            page_breaks = element.findall(".//w:br[@w:type='page']", ns)
            page += len(page_breaks)
        elif element.tag == alt_chunk_tag:
            relationship_id = element.get(relationship_id_tag, "")
            for text, block_type in alt_chunks.get(relationship_id, []):
                explicit_section = text if block_type == "heading" else ""
                values.append((text, page, block_type, explicit_section))

    blocks = _bounded_blocks(values)
    if not blocks:
        raise DocumentParseError("未从 DOCX 中提取到可审查文字。")
    return ParsedDocument(filename, "DOCX", max(page, 1), blocks)


def parse_legacy_doc(content: bytes, filename: str) -> ParsedDocument:
    if content.startswith(b"PK"):
        return parse_docx(content, f"{Path(filename).stem}.docx")
    try:
        with tempfile.TemporaryDirectory(prefix="xzs-doc-") as temporary:
            source = Path(temporary) / "source.doc"
            source.write_bytes(content)
            extracted = extract_legacy_word_text(source)
    except WindowsIFilterError as error:
        raise DocumentParseError(str(error)) from error

    values: list[tuple[str, int, str, str]] = []
    pages = extracted.replace("\r\n", "\n").replace("\r", "\n").split("\f")
    for page_number, page in enumerate(pages, start=1):
        for line in page.splitlines():
            block_type = "heading" if looks_like_heading(line) else "paragraph"
            values.append((line, page_number, block_type, line if block_type == "heading" else ""))
    blocks = _bounded_blocks(values)
    if not blocks:
        raise DocumentParseError("旧版 DOC 未提取到可索引正文。")
    return ParsedDocument(filename, "DOC", max(1, len(pages)), blocks)


def parse_pdf(content: bytes, filename: str, ocr_service: LocalOCRService | None = None) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise DocumentParseError("当前环境缺少 pypdf，无法解析 PDF。") from error
    try:
        reader = PdfReader(io.BytesIO(content))
        page_texts: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except (TypeError, ValueError, KeyError):
                text = page.extract_text() or ""
            page_texts.append(text)
    except Exception as error:
        raise DocumentParseError("PDF 无法提取文字；文件可能已加密或损坏。") from error

    ocr_targets = [
        page_number
        for page_number, text in enumerate(page_texts, start=1)
        if len(re.sub(r"\s+", "", text)) < 40
    ]
    ocr_pages: list[int] = []
    if ocr_targets and ocr_service is not None:
        try:
            recognized_pages = ocr_service.recognize_pdf_pages(content, ocr_targets)
        except LocalOCRError as error:
            raise DocumentParseError(f"扫描页本地 OCR 失败：{error}") from error
        for result in recognized_pages:
            existing = page_texts[result.page - 1]
            if len(result.text) >= len(existing.strip()):
                page_texts[result.page - 1] = result.text
            ocr_pages.append(result.page)

    values: list[tuple[str, int, str, str]] = []
    extracted_characters = 0
    for page_number, text in enumerate(page_texts, start=1):
        extracted_characters += len(text.strip())
        _append_pdf_page_values(values, text, page_number)
    blocks = _bounded_blocks(values)
    scanned = bool(ocr_targets)
    if not blocks:
        raise DocumentParseError("当前文件可能为扫描件，但本地 OCR 未能提取到可审查文字。")
    if ocr_targets and not ocr_pages and extracted_characters < max(80, len(reader.pages) * 20):
        raise DocumentParseError("当前文件为扫描件，需要启用 Windows 本地 OCR 后再解析。")
    return ParsedDocument(
        filename,
        "PDF",
        len(reader.pages),
        blocks,
        scanned,
        bool(ocr_pages),
        tuple(ocr_pages),
        ocr_service.engine_name if ocr_pages and ocr_service else "",
    )


PDF_TABLE_HEADER_TERMS = (
    "项目", "本期", "上期", "期末", "期初", "本报告期", "上年同期", "金额", "比例", "单位",
)


def _split_pdf_table_line(raw_line: str) -> list[str]:
    line = raw_line.strip()
    if not line:
        return []
    wide_cells = [normalize_text(cell) for cell in re.split(r"\t+| {2,}", line) if normalize_text(cell)]
    if len(wide_cells) >= 2:
        return wide_cells
    tokens = [normalize_text(token) for token in line.split() if normalize_text(token)]
    if len(tokens) < 2:
        return []
    numeric_count = sum(
        bool(re.fullmatch(r"[（(]?[+\-−—－]?\d[\d,，]*(?:\.\d+)?[）)]?[%％]?", token))
        for token in tokens
    )
    header_like = any(any(term in token for term in PDF_TABLE_HEADER_TERMS) for token in tokens)
    return tokens if numeric_count >= 2 or (header_like and len(tokens) >= 3) else []


def _append_pdf_page_values(
    values: list[tuple[str, int, str, str]],
    text: str,
    page_number: int,
) -> None:
    pending: list[tuple[str, list[str]]] = []

    def flush_pending() -> None:
        if not pending:
            return
        widths = [len(cells) for _raw, cells in pending]
        has_header = any(
            any(any(term in cell for term in PDF_TABLE_HEADER_TERMS) for cell in cells)
            for _raw, cells in pending
        )
        numeric_rows = sum(
            sum(parse_decimal_candidate(cell) for cell in cells) >= 1
            for _raw, cells in pending
        )
        coherent_width = max(widths, default=0) - min(widths, default=0) <= 2
        if len(pending) >= 2 and has_header and numeric_rows >= 1 and coherent_width:
            values.append(("\n".join(" | ".join(cells) for _raw, cells in pending), page_number, "table", ""))
        else:
            for raw, _cells in pending:
                block_type = "heading" if looks_like_heading(raw) else "paragraph"
                values.append((raw, page_number, block_type, raw if block_type == "heading" else ""))
        pending.clear()

    for raw_line in text.splitlines():
        cells = _split_pdf_table_line(raw_line)
        if cells:
            pending.append((raw_line, cells))
            continue
        flush_pending()
        block_type = "heading" if looks_like_heading(raw_line) else "paragraph"
        values.append((raw_line, page_number, block_type, raw_line if block_type == "heading" else ""))
    flush_pending()


def parse_decimal_candidate(value: str) -> bool:
    compact = str(value or "").strip().replace(",", "").replace("，", "")
    compact = compact.replace("−", "-").replace("—", "-").replace("－", "-").rstrip("%％")
    if compact.startswith(("(", "（")) and compact.endswith((")", "）")):
        compact = compact[1:-1]
    return bool(re.fullmatch(r"[+\-]?\d+(?:\.\d+)?", compact))


def parse_text(content: bytes, filename: str) -> ParsedDocument:
    decoded = _decode_text(content)
    values = []
    for line in re.split(r"\r?\n+", decoded):
        block_type = "heading" if looks_like_heading(line) else "paragraph"
        values.append((line, 1, block_type, line if block_type == "heading" else ""))
    blocks = _bounded_blocks(values)
    if not blocks:
        raise DocumentParseError("文本文件没有可审查内容。")
    return ParsedDocument(filename, Path(filename).suffix.removeprefix(".").upper() or "TXT", 1, blocks)


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("文本文件编码无法识别，请使用 UTF-8 或 GB18030。")


def parse_csv(content: bytes, filename: str) -> ParsedDocument:
    decoded = _decode_text(content)
    try:
        rows = list(csv.reader(io.StringIO(decoded)))
    except csv.Error as error:
        raise DocumentParseError("CSV 文件结构无效。") from error
    if not rows:
        raise DocumentParseError("CSV 文件没有可审查内容。")

    width = max(len(row) for row in rows)
    headers = [normalize_text(value) or f"列{index + 1}" for index, value in enumerate(rows[0])]
    headers.extend(f"列{index + 1}" for index in range(len(headers), width))
    current_section = Path(filename).stem
    ignored_export_section = ""
    values: list[tuple[str, int, str, str]] = []
    for row_number, raw_row in enumerate(rows[1:], start=2):
        cells = [normalize_text(value) for value in raw_row]
        cells.extend("" for _ in range(width - len(cells)))
        if not any(cells):
            continue
        first_cell = cells[0]
        if re.fullmatch(r"Sheet\d+", first_cell, flags=re.IGNORECASE):
            ignored_export_section = first_cell
            continue
        if ignored_export_section:
            if first_cell == ignored_export_section:
                continue
            ignored_export_section = ""
        if first_cell:
            current_section = first_cell
        if len(cells) > 2 and cells[1] == "序号" and cells[2]:
            headers = ["分类", *[value or f"列{index + 2}" for index, value in enumerate(cells[1:])]]
            continue
        if not cells[1] and not cells[2] and sum(bool(value) for value in cells) <= 3:
            continue
        fields = [f"{headers[index]}：{value}" for index, value in enumerate(cells) if value]
        values.append((f"CSV行{row_number}；" + "；".join(fields), 1, "table", current_section, (row_number,), ""))
    blocks = _bounded_blocks(values)
    if not blocks:
        raise DocumentParseError("CSV 文件没有可索引的数据行。")
    return ParsedDocument(filename, "CSV", 1, blocks)


def _spreadsheet_value(value: object) -> str:
    """Render spreadsheet values without losing signs, precision, or dates."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return normalize_text(str(value))


def _xlsx_sheets(content: bytes) -> list[tuple[str, list[tuple[int, list[str]]]]]:
    try:
        import openpyxl
    except ImportError as error:
        raise DocumentParseError("当前环境缺少 openpyxl，无法解析 XLSX 文件。") from error
    try:
        # Keep formulas as a fallback when a workbook has no cached result.
        formula_book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        value_book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise DocumentParseError("XLSX 文件结构无效或已损坏。") from error
    sheets: list[tuple[str, list[tuple[int, list[str]]]]] = []
    try:
        for formula_sheet, value_sheet in zip(formula_book.worksheets, value_book.worksheets, strict=True):
            rows: list[tuple[int, list[str]]] = []
            for row_number, (formula_row, value_row) in enumerate(
                zip(formula_sheet.iter_rows(), value_sheet.iter_rows(), strict=True), start=1
            ):
                cells: list[str] = []
                for formula_cell, value_cell in zip(formula_row, value_row, strict=True):
                    value = value_cell.value
                    if value is None and formula_cell.data_type == "f":
                        value = formula_cell.value
                    cells.append(_spreadsheet_value(value))
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append((row_number, cells))
            if rows:
                sheets.append((str(formula_sheet.title), rows))
    finally:
        formula_book.close()
        value_book.close()
    return sheets


def _xls_sheets(content: bytes) -> list[tuple[str, list[tuple[int, list[str]]]]]:
    try:
        import xlrd
    except ImportError as error:
        raise DocumentParseError(
            "当前环境缺少 xlrd，无法打开旧版 XLS。请运行 web\\install-local-embedding.ps1 或 pip install xlrd。"
        ) from error
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    except Exception as error:
        raise DocumentParseError("XLS 文件结构无效或已损坏。") from error
    sheets: list[tuple[str, list[tuple[int, list[str]]]]] = []
    try:
        for sheet in workbook.sheets():
            rows: list[tuple[int, list[str]]] = []
            for row_index in range(sheet.nrows):
                cells = [_spreadsheet_value(sheet.cell_value(row_index, column)) for column in range(sheet.ncols)]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append((row_index + 1, cells))
            if rows:
                sheets.append((str(sheet.name), rows))
    finally:
        workbook.release_resources()
    return sheets


def parse_spreadsheet(content: bytes, filename: str) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    sheets = _xlsx_sheets(content) if extension == ".xlsx" else _xls_sheets(content)
    if not sheets:
        raise DocumentParseError("Excel 文件没有可审查的非空工作表。")
    values: list[tuple] = []
    source_offset = 0
    paragraph_index = 0
    for sheet_name, rows in sheets:
        rendered_rows = [" | ".join(cells) for _row_number, cells in rows]
        text = "\n".join(rendered_rows)
        if not text.strip():
            continue
        values.append((text, len(values) + 1, "table", sheet_name, tuple(row_number for row_number, _cells in rows), sheet_name))
        source_offset += len(text) + 1
        paragraph_index += 1
    blocks = _bounded_blocks(values)
    if not blocks:
        raise DocumentParseError("Excel 文件没有可审查的非空工作表。")
    return ParsedDocument(filename, extension.removeprefix(".").upper(), len(sheets), blocks)


def parse_document(
    filename: str,
    content: bytes,
    ocr_service: LocalOCRService | None = None,
) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentParseError("仅支持 DOC、DOCX、PDF、XLS、XLSX、CSV、TXT、MD 或 JSON 文件。")
    if extension == ".doc":
        return parse_legacy_doc(content, filename)
    if extension == ".docx":
        return parse_docx(content, filename)
    if extension == ".pdf":
        return parse_pdf(content, filename, ocr_service)
    if extension == ".csv":
        return parse_csv(content, filename)
    if extension in {".xls", ".xlsx"}:
        return parse_spreadsheet(content, filename)
    return parse_text(content, filename)
