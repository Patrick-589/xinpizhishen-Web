from __future__ import annotations

import io
import http.client
import json
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from http.server import ThreadingHTTPServer

WEB_ROOT = Path(__file__).resolve().parents[1]
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

from docx import Document  # noqa: E402

from backend.config import Settings  # noqa: E402
from backend.database import Database, utc_now  # noqa: E402
from backend.drafting import build_docx, build_draft  # noqa: E402
from backend.meeting_case_library import extract_proposal_titles, infer_case_role  # noqa: E402
from backend.rag import RAGContext  # noqa: E402
from backend.security import AuthService, hash_password  # noqa: E402
from backend.services import PlatformService  # noqa: E402
from backend.vector_store import SearchResult  # noqa: E402
from server import ComplianceRequestHandler  # noqa: E402


class DraftingTests(unittest.TestCase):
    def _context(self, *, results: list[SearchResult] | None = None) -> RAGContext:
        return RAGContext("关联交易董事会审批", results or [])

    def _source(self) -> SearchResult:
        return SearchResult(
            chunk_id="chunk-001",
            document_id="knowledge-001",
            content="关联交易达到公司章程规定标准的，应提交董事会审议。",
            title="关联交易管理制度",
            source="证券部制度",
            date="2026-06-01",
            page=8,
            paragraph=3,
            section="第三章 审批程序",
            document_type="internal_rules",
            score=0.91,
            source_url=None,
            version="2026年修订版",
            metadata={
                "file_name": "关联交易管理制度.docx",
                "document_name": "关联交易管理制度",
                "regulation_name": "关联交易管理制度",
                "document_updated_at": "2026-06-01T08:00:00+00:00",
            },
        )

    def test_each_supported_type_has_type_specific_sections(self) -> None:
        expected = {
            "temporary_announcement": "重要提示",
            "board_resolution": "回避表决情况",
            "shareholder_resolution": "律师见证情况",
            "meeting_sign_in": "签到登记",
            "meeting_agenda_summary": "主要议案概述",
            "general_manager_meeting": "总经理审议意见",
            "proxy_authorization": "授权事项",
            "legal_opinion": "结论与保留事项",
            "related_transaction": "定价政策及定价依据",
            "investment_announcement": "风险分析",
            "periodic_report_notice": "主要财务数据",
            "szse_report_brief": "公司业务发展历史简介",
        }
        for document_type, section in expected.items():
            with self.subTest(document_type=document_type):
                result = build_draft(
                    document_type=document_type,
                    title="",
                    subject="子公司增资",
                    facts="交易金额为待确认。",
                    requirements="",
                    context=self._context(),
                )
                self.assertIn(section, result.content)
                self.assertEqual(result.document_type, document_type)

    def test_general_manager_meeting_uses_local_case_evidence(self) -> None:
        result = build_draft(
            document_type="general_manager_meeting",
            title="",
            subject="子公司增资",
            facts="公司拟向子公司增资，金额待确认。",
            requirements="",
            context=self._context(),
        )
        self.assertEqual(result.document_type, "general_manager_meeting")
        self.assertIn("总经理审议意见", result.content)
        self.assertIn("决策及后续事项", result.content)
        self.assertIn("总经理会议记录", result.title)

    def test_szse_report_brief_is_one_page_and_uses_explicit_placeholders(self) -> None:
        result = build_draft(
            document_type="szse_report_brief",
            title="",
            subject="",
            facts="",
            requirements="控制在一页纸内",
            context=self._context(),
        )
        self.assertEqual(result.document_type, "szse_report_brief")
        self.assertIn("公司业务发展历史简介", result.content)
        self.assertIn("公司当前业务情况", result.content)
        self.assertIn("重大资产重组基本情况", result.content)
        self.assertIn("重大资产重组最新进展", result.content)
        self.assertIn("后续工作安排", result.content)
        self.assertNotIn("需进一步说明或请示事项", result.content)
        self.assertIn("当前知识库未检索到", result.content)
        self.assertIn("一页纸说明", result.content)
        self.assertLess(len(result.content), 6000)
        payload = build_docx(
            title=result.title,
            content=result.content,
            sources=result.sources,
            compact=True,
        )
        document = Document(io.BytesIO(payload))
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("公司业务发展历史简介", docx_text)
        self.assertIn("重大资产重组最新进展", docx_text)

    def test_unknown_facts_remain_explicit_placeholders(self) -> None:
        result = build_draft(
            document_type="board_resolution",
            title="",
            subject="关联交易审议",
            facts="",
            requirements="",
            context=self._context(),
        )
        self.assertTrue(result.warnings)
        self.assertIn("[待补充", result.content)
        self.assertIn("[待确认", result.content)
        self.assertNotIn("本次会议于2026年", result.content)

    def test_rag_source_keeps_provenance_and_local_document_link(self) -> None:
        result = build_draft(
            document_type="related_transaction",
            title="",
            subject="子公司增资",
            facts="公司拟向子公司增资，金额待确认。",
            requirements="",
            context=self._context(results=[self._source()]),
        )
        self.assertEqual(len(result.sources), 1)
        source = result.sources[0]
        self.assertEqual(source["page"], 8)
        self.assertEqual(source["section"], "第三章 审批程序")
        self.assertEqual(source["regulation_name"], "关联交易管理制度")
        self.assertEqual(source["location"], "第8页 · 第三章 审批程序 · Chunk chunk-001")
        self.assertEqual(source["source_url"], "/api/knowledge/knowledge-001/document")
        self.assertIn("关联交易管理制度", result.content)
        self.assertEqual(source["file_name"], "关联交易管理制度.docx")

    def test_build_docx_is_editable_and_contains_sources(self) -> None:
        source = self._source().citation()
        content = "草稿正文\n\n1、事项概述\n金额：[待补充]"
        payload = build_docx(title="测试起草公告", content=content, sources=[source])
        self.assertTrue(payload.startswith(b"PK"))
        document = Document(io.BytesIO(payload))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("测试起草公告", text)
        self.assertIn("[待补充]", text)
        self.assertIn("共引用 1 个本地 RAG 来源", text)

    def test_periodic_draft_uses_company_report_style_reference_without_copying_facts(self) -> None:
        historical = self._source().__class__(
            chunk_id="hist-001",
            document_id="history-001",
            content="第一节 重要提示、目录和释义\n公司简介和主要财务指标\n管理层讨论与分析",
            title="创业黑马科技集团股份有限公司2025年半年度报告",
            source="巨潮资讯网",
            date="2025-08-28",
            page=2,
            paragraph=1,
            section="目录",
            document_type="announcements",
            score=0.94,
            source_url="https://www.cninfo.com.cn/test",
            version="正式公告",
            metadata={},
        )
        result = build_draft(
            document_type="periodic_report_notice",
            title="",
            subject="2026年半年度报告",
            facts="报告期金额待确认。",
            requirements="",
            context=self._context(results=[historical]),
        )
        self.assertTrue(result.style_sources)
        self.assertIn("公司历史定期报告格式参考", result.content)
        self.assertIn("主要财务数据和会计指标", result.content)
        self.assertNotIn("2025年半年度报告实现净利润", result.content)

    def test_drafting_material_is_used_as_style_reference_with_local_link(self) -> None:
        template_source = SearchResult(
            chunk_id="template-001",
            document_id="drafting-material-001",
            content="一、公司基本情况\n二、重大事项进展\n三、后续工作安排",
            title="公司监管汇报材料模板",
            source="证券部起草资料",
            date="2026-07-01",
            page=1,
            paragraph=1,
            section="汇报材料模板",
            document_type="drafting_materials",
            score=0.93,
            source_url=None,
            version="2026年版",
            metadata={"file_name": "监管汇报材料模板.docx"},
        )
        result = build_draft(
            document_type="periodic_report_notice",
            title="",
            subject="2026年半年度报告",
            facts="报告期金额待确认。",
            requirements="",
            context=self._context(results=[template_source]),
        )
        self.assertTrue(result.style_sources)
        self.assertEqual(result.style_sources[0]["document_type"], "drafting_materials")
        self.assertEqual(
            result.style_sources[0]["source_url"],
            "/api/knowledge/drafting-material-001/document",
        )
        self.assertIn("公司监管汇报材料模板", result.content)

    def test_meeting_record_uses_case_proposal_and_specific_compliance_speaking(self) -> None:
        source = SearchResult(
            chunk_id="meeting-case-001",
            document_id="meeting-case-doc",
            content=(
                "议案：《关于与交易对方签署《业绩承诺与补偿协议之补充协议》的议案》\n"
                "本次补充协议用于进一步明确各方权利义务，保障上市公司股东的合法权益。\n"
                "本议案审议通过后提交董事会审议。"
            ),
            title="5、2026第【】次审计委员会会议 议案.doc",
            source="公司已提供会议起草案例",
            date="2026-06-01",
            page=1,
            paragraph=3,
            section="会议议案",
            document_type="drafting_materials",
            score=0.95,
            source_url=None,
            version="案例",
            metadata={
                "case_role": "proposal",
                "meeting_type": "audit_committee",
                "file_name": "5、2026第【】次审计委员会会议 议案.doc",
                "source_priority": 68,
            },
        )
        result = build_draft(
            document_type="committee_record",
            title="",
            subject="与交易对方签署业绩承诺与补偿协议之补充协议",
            facts="",
            requirements="会议发言要围绕本次议案",
            context=self._context(results=[source]),
        )
        self.assertIn("议案说明", result.content)
        self.assertIn("业绩承诺与补偿协议之补充协议", result.content)
        self.assertIn("财务数据、评估/审计依据", result.content)
        self.assertIn("本地 RAG 起草依据", result.content)

    def test_case_role_and_proposal_title_are_deterministic(self) -> None:
        self.assertEqual(infer_case_role("审计委员会会议记录.doc", "会议发言和讨论"), "record")
        titles = extract_proposal_titles("议案一：《关于修订《公司章程》的议案》\n审议通过《关于交易的议案》")
        self.assertTrue(any("修订" in title for title in titles))

    def test_meeting_sign_in_builds_editable_word_table(self) -> None:
        source = SearchResult(
            chunk_id="sign-in-case-001",
            document_id="sign-in-doc",
            content="2025年度股东大会\n议案一：《关于修订《公司章程》的议案》",
            title="2025年度股东大会议案.docx",
            source="本次会议参考文书",
            date=None,
            page=1,
            paragraph=1,
            section="会议议案",
            document_type="drafting_materials",
            score=1.0,
            source_url=None,
            version=None,
            metadata={"direct_reference": True, "meeting_term": "2025年度股东大会"},
        )
        result = build_draft(
            document_type="meeting_sign_in",
            title="",
            subject="2025年度股东大会",
            facts="",
            requirements="",
            context=self._context(results=[source]),
        )
        self.assertIn("签到登记", result.content)
        self.assertIn("关于修订", result.content)
        payload = build_docx(title=result.title, content=result.content, sources=result.sources)
        document = Document(io.BytesIO(payload))
        self.assertEqual(len(document.tables), 1)
        self.assertIn("姓名/股东名称", document.tables[0].cell(0, 1).text)


class DraftOwnershipTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.settings = Settings(
            app_env="test",
            ai_mode="local_rag",
            database_path=root / "runtime" / "test.db",
            upload_root=root / "runtime" / "uploads",
            knowledge_root=root / "knowledge",
            max_upload_bytes=20 * 1024 * 1024,
            session_hours=12,
            bootstrap_admin_username="admin",
            bootstrap_admin_password="TestAdmin123!",
            app_secret="test-secret-not-used-by-cookie-sessions",
            aliyun_api_key="",
            aliyun_base_url="",
            aliyun_model="",
            aliyun_chat_path="/chat/completions",
            request_timeout_seconds=2,
            embedding_batch_size=8,
            vector_backend="sqlite",
            database_url="",
            vector_db_url="",
            input_cost_per_million=2,
            output_cost_per_million=8,
        )
        self.database = Database(self.settings.database_path)
        self.database.migrate()
        now = utc_now()
        for user_id, username in (("user-a", "draft-a"), ("user-b", "draft-b")):
            self.database.execute(
                "INSERT INTO users (id,username,password_hash,display_name,role,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
                (user_id, username, hash_password("DraftTest123!"), username, "user", now, now),
            )
        self.platform = PlatformService(self.settings, self.database)
        self.platform.rag.retrieve = lambda *args, **kwargs: RAGContext("", [])
        self.platform.rag.retrieve_keyword = lambda *args, **kwargs: RAGContext("", [])
        self.user_a = {"id": "user-a", "username": "draft-a", "role": "user"}
        self.user_b = {"id": "user-b", "username": "draft-b", "role": "user"}

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_drafts_are_private_between_accounts(self) -> None:
        draft = self.platform.create_draft(
            {
                "document_type": "temporary_announcement",
                "subject": "测试事项",
                "facts": "已确认事实",
            },
            self.user_a,
            "127.0.0.1",
        )
        self.assertEqual(len(self.platform.list_drafts(self.user_a)), 1)
        self.assertEqual(self.platform.list_drafts(self.user_b), [])
        with self.assertRaises(ValueError):
            self.platform.get_draft(draft["id"], self.user_b)
        with self.assertRaises(ValueError):
            self.platform.delete_draft(draft["id"], self.user_b, "127.0.0.1")
        self.assertEqual(self.platform.delete_draft(draft["id"], self.user_a, "127.0.0.1")["status"], "deleted")

    def test_szse_brief_retrieves_company_announcements_and_rules_separately(self) -> None:
        calls: list[tuple[str, ...] | None] = []

        def retrieve(*args, **kwargs):
            calls.append(tuple(kwargs.get("document_types") or []))
            return RAGContext(str(args[0]) if args else "", [])

        def retrieve_keyword(*args, **kwargs):
            calls.append(tuple(kwargs.get("document_types") or []))
            return RAGContext(str(args[0]) if args else "", [])

        self.platform.rag.retrieve = retrieve
        self.platform.rag.retrieve_keyword = retrieve_keyword
        draft = self.platform.create_draft(
            {
                "document_type": "szse_report_brief",
                "subject": "汇报业务发展及重组进展",
                "facts": "",
                "requirements": "内容包括公司业务发展历史简介和公司重组进展情况汇报，控制在一页纸内。",
            },
            self.user_a,
            "127.0.0.1",
        )
        self.assertEqual(draft["document_type"], "szse_report_brief")
        self.assertIn(("announcements",), calls)
        self.assertIn(("drafting_materials",), calls)
        self.assertIn(("regulations", "internal_rules", "company_profile", "cases"), calls)

    def test_direct_szse_request_is_inferred_without_manual_type_selection(self) -> None:
        calls: list[tuple[str, ...] | None] = []

        def retrieve(*args, **kwargs):
            calls.append(tuple(kwargs.get("document_types") or []))
            return RAGContext(str(args[0]) if args else "", [])

        self.platform.rag.retrieve = retrieve
        self.platform.rag.retrieve_keyword = retrieve
        draft = self.platform.create_draft(
            {
                "document_type": "auto",
                "requirements": "内容包括公司业务发展历史简介和公司重组进展情况汇报，控制在一页纸内。",
            },
            self.user_a,
            "127.0.0.1",
        )
        self.assertEqual(draft["document_type"], "szse_report_brief")
        self.assertIn("公司业务发展历史简介", draft["content"])

    def test_meeting_draft_retrieves_case_material_and_rules_separately(self) -> None:
        calls: list[tuple[str, ...]] = []

        def retrieve(*args, **kwargs):
            calls.append(tuple(kwargs.get("document_types") or []))
            return RAGContext(str(args[0]) if args else "", [])

        self.platform.rag.retrieve = retrieve
        self.platform.rag.retrieve_keyword = retrieve
        draft = self.platform.create_draft(
            {
                "document_type": "committee_record",
                "subject": "与交易对方签署业绩承诺与补偿协议之补充协议",
                "requirements": "会议发言围绕具体议案和审议权限",
            },
            self.user_a,
            "127.0.0.1",
        )
        self.assertEqual(draft["document_type"], "committee_record")
        self.assertIn(("drafting_materials", "announcements", "cases", "user_documents", "company_profile"), calls)
        self.assertIn(("regulations", "internal_rules"), calls)

    def test_specific_meeting_type_wins_automatic_inference(self) -> None:
        for prompt, expected in (
            ("生成会议签到表", "meeting_sign_in"),
            ("生成主要议案概述", "meeting_agenda_summary"),
            ("生成授权委托书", "proxy_authorization"),
            ("生成法律意见书", "legal_opinion"),
            ("审计委员会2026年第X次会议记录", "committee_record"),
            ("独立董事专门会议决议", "independent_director_resolution"),
            ("董事会会议决议", "board_resolution"),
            ("股东会会议决议", "shareholder_resolution"),
            ("审计委员会会议表决票", "ballot"),
            ("生成总经理会议记录", "general_manager_meeting"),
        ):
            with self.subTest(prompt=prompt):
                draft = self.platform.create_draft(
                    {"document_type": "auto", "subject": prompt},
                    self.user_a,
                    "127.0.0.1",
                )
                self.assertEqual(draft["document_type"], expected)

    def test_direct_meeting_reference_stays_private_and_is_only_source(self) -> None:
        document = Document()
        document.add_paragraph("2025年度股东大会")
        document.add_paragraph("议案一：《关于修订《公司章程》的议案》")
        output = io.BytesIO()
        document.save(output)
        draft = self.platform.create_draft_from_reference(
            "2025年度股东大会议案.docx",
            output.getvalue(),
            {"document_type": "auto"},
            self.user_a,
            "127.0.0.1",
        )
        self.assertEqual(draft["document_type"], "meeting_sign_in")
        self.assertIn("2025年度股东大会", draft["content"])
        self.assertEqual(len(draft["source_documents"]), 1)
        self.assertIn("/api/drafts/", draft["source_documents"][0]["source_url"])
        filename, original = self.platform.get_draft_reference_document(
            draft["id"], self.user_a, "127.0.0.1",
        )
        self.assertEqual(filename, "2025年度股东大会议案.docx")
        self.assertTrue(original.startswith(b"PK"))


class DraftReferenceHttpTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.settings = Settings(
            app_env="test",
            ai_mode="local_rag",
            database_path=root / "runtime" / "test.db",
            upload_root=root / "runtime" / "uploads",
            knowledge_root=root / "knowledge",
            max_upload_bytes=20 * 1024 * 1024,
            session_hours=12,
            bootstrap_admin_username="admin",
            bootstrap_admin_password="TestAdmin123!",
            app_secret="test-secret-not-used-by-cookie-sessions",
            aliyun_api_key="",
            aliyun_base_url="",
            aliyun_model="",
            aliyun_chat_path="/chat/completions",
            request_timeout_seconds=2,
            embedding_batch_size=8,
            vector_backend="sqlite",
            database_url="",
            vector_db_url="",
            input_cost_per_million=2,
            output_cost_per_million=8,
        )
        self.database = Database(self.settings.database_path)
        self.database.migrate()
        now = utc_now()
        self.database.execute(
            "INSERT INTO users (id,username,password_hash,display_name,role,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
            ("http-user", "http-draft", hash_password("DraftTest123!"), "http-draft", "user", now, now),
        )
        auth = AuthService(self.database, self.settings)
        self.platform = PlatformService(self.settings, self.database)
        self.platform.rag.retrieve = lambda *args, **kwargs: RAGContext("", [])
        self.platform.rag.retrieve_keyword = lambda *args, **kwargs: RAGContext("", [])
        ComplianceRequestHandler.settings = self.settings
        ComplianceRequestHandler.database = self.database
        ComplianceRequestHandler.auth = auth
        ComplianceRequestHandler.platform = self.platform
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), ComplianceRequestHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        self.temporary.cleanup()

    def request(self, method: str, path: str, body: bytes, headers: dict[str, str]) -> tuple[int, dict, list[tuple[str, str]]]:
        connection = http.client.HTTPConnection("127.0.0.1", self.server.server_port, timeout=10)
        connection.request(method, path, body=body, headers=headers)
        response = connection.getresponse()
        payload = json.loads(response.read().decode("utf-8"))
        response_headers = response.getheaders()
        connection.close()
        return response.status, payload, response_headers

    def test_reference_docx_multipart_route_is_not_parsed_as_json(self) -> None:
        login_body = json.dumps({"username": "http-draft", "password": "DraftTest123!"}).encode("utf-8")
        status, _payload, headers = self.request(
            "POST",
            "/api/auth/login",
            login_body,
            {"Content-Type": "application/json", "Content-Length": str(len(login_body))},
        )
        self.assertEqual(status, 200)
        cookie = next(value.split(";", 1)[0] for name, value in headers if name.casefold() == "set-cookie")

        document = Document()
        document.add_paragraph("2025年第一次临时股东会")
        document.add_paragraph("议案一：《关于修订公司章程的议案》")
        stream = io.BytesIO()
        document.save(stream)
        boundary = "----xzs-draft-reference-test"
        chunks = [
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"document_type\"\r\n\r\nauto\r\n".encode(),
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"reference_document\"; filename=\"meeting.docx\"\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n".encode(),
            stream.getvalue(),
            f"\r\n--{boundary}--\r\n".encode(),
        ]
        body = b"".join(chunks)
        status, payload, _headers = self.request(
            "POST",
            "/api/drafts/from-reference",
            body,
            {
                "Content-Type": f"multipart/form-data; boundary={boundary}",
                "Content-Length": str(len(body)),
                "Cookie": cookie,
            },
        )
        self.assertEqual(status, 201)
        self.assertEqual(payload["document_type"], "meeting_sign_in")
        self.assertIn("2025年第一次临时股东会", payload["content"])

    def test_batch_delete_route_accepts_task_id_array(self) -> None:
        content = "批量删除 HTTP 测试\n需要核对信息披露。".encode("utf-8")
        first = self.platform.create_task(
            "http-bulk-one.txt", content, "HTTP 批量任务一", "重大事项披露审查", "基础包",
            {"id": "http-user", "username": "http-draft", "role": "user"}, "127.0.0.1",
        )
        second = self.platform.create_task(
            "http-bulk-two.txt", content, "HTTP 批量任务二", "重大事项披露审查", "基础包",
            {"id": "http-user", "username": "http-draft", "role": "user"}, "127.0.0.1",
        )
        login_body = json.dumps({"username": "http-draft", "password": "DraftTest123!"}).encode("utf-8")
        _status, _payload, headers = self.request(
            "POST", "/api/auth/login", login_body,
            {"Content-Type": "application/json", "Content-Length": str(len(login_body))},
        )
        cookie = next(value.split(";", 1)[0] for name, value in headers if name.casefold() == "set-cookie")
        request_body = json.dumps({"task_ids": [first["taskId"], second["taskId"]]}).encode("utf-8")
        status, payload, _headers = self.request(
            "POST", "/api/tasks/batch-delete", request_body,
            {"Content-Type": "application/json", "Content-Length": str(len(request_body)), "Cookie": cookie},
        )
        self.assertEqual(status, 200)
        self.assertEqual(payload["taskCount"], 2)


if __name__ == "__main__":
    unittest.main()
