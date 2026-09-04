from __future__ import annotations

import io
import sys
import unittest
from decimal import Decimal
from pathlib import Path

from docx import Document
from openpyxl import Workbook

WEB_ROOT = Path(__file__).resolve().parents[1]
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

from backend.financial_audit.engine import FinancialAuditEngine
from backend.parsers import ParsedBlock, ParsedDocument, parse_document


def _parsed_table(section: str, text: str, filename: str) -> ParsedDocument:
    return ParsedDocument(
        filename,
        Path(filename).suffix.removeprefix(".").upper(),
        1,
        [ParsedBlock(text, 1, section, 0, "table", 0, len(text))],
    )


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("2026年半年度报告")
    table = document.add_table(rows=3, cols=3)
    values = [
        ("项目", "本报告期", "上年同期"),
        ("营业收入", "100", "90"),
        ("净利润", "20", "18"),
    ]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _detailed_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("2、应收账款")
    table = document.add_table(rows=4, cols=3)
    values = [
        ("项目", "2026年6月30日余额", "2025年12月31日余额"),
        ("单项计提坏账准备的应收账款", "15349846.27", "15362346.22"),
        ("按组合计提坏账准备的应收账款", "48860433.83", "49171790.25"),
        ("坏账准备", "59114841.35", "54158387.21"),
    ]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _detailed_spreadsheet_bytes(bad_debt: str = "59114841.35") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "应收账款附注"
    sheet.append(["项目", "2026年6月30日余额", "2025年12月31日余额"])
    sheet.append(["单项计提坏账准备的应收账款", "15349846.27", "15362346.22"])
    sheet.append(["按组合计提坏账准备的应收账款", "48860433.83", "49171790.25"])
    sheet.append(["坏账准备", bad_debt, "54158387.21"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _spreadsheet_bytes(net_profit: str = "20", extension: str = ".xlsx") -> bytes:
    if extension == ".xls":
        import xlwt

        output = io.BytesIO()
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("合并利润表")
        values = [
            ("项目", "本报告期", "上年同期"),
            ("营业收入", "100", "90"),
            ("净利润", net_profit, "18"),
        ]
        for row_index, row in enumerate(values):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, value)
        workbook.save(output)
        return output.getvalue()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "合并利润表"
    sheet.append(["项目", "本报告期", "上年同期"])
    sheet.append(["营业收入", "100", "90"])
    sheet.append(["净利润", net_profit, "18"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class ExcelReconciliationTests(unittest.TestCase):
    def test_current_authoritative_sheet_wins_over_archived_duplicate(self) -> None:
        """A duplicate workbook tab must not hide a current-sheet mismatch."""
        document = _parsed_table(
            "4、营业收入",
            "项目 | 2026年1-6月金额\n营业收入 | 100",
            "2026年半年度报告.docx",
        )
        workbook = Workbook()
        archived = workbook.active
        archived.title = "X0合并利润表"
        archived.append(["项目", "2026年1-6月"])
        archived.append(["营业收入", "100"])
        current = workbook.create_sheet("合并利润表")
        current.append(["项目", "2026年1-6月"])
        current.append(["营业收入", "120"])
        output = io.BytesIO()
        workbook.save(output)
        spreadsheet = parse_document("财务底稿.xlsx", output.getvalue())

        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        mismatch = next(
            item for item in result.checks
            if item.rule_id == "CROSS-DOCUMENT-FACT" and item.status == "failed"
            and "营业收入" in item.title
        )
        self.assertEqual(mismatch.calculated_value, Decimal("120"))
        self.assertEqual(mismatch.source_locations[1].sheet_name, "合并利润表")

    def test_nested_note_prefers_matching_note_sheet_over_statement_total(self) -> None:
        """A customer/detail row must use its matching note tab, not a statement aggregate."""
        document = _parsed_table(
            "（1）客户明细",
            "项目 | 2026年6月30日余额\n应收账款 | 100",
            "2026年半年度报告.docx",
        )
        workbook = Workbook()
        statement = workbook.active
        statement.title = "合并资产表"
        statement.append(["项目", "2026年6月30日"])
        statement.append(["应收账款", "999"])
        note = workbook.create_sheet("客户明细")
        note.append(["项目", "2026年6月30日"])
        note.append(["应收账款", "100"])
        output = io.BytesIO()
        workbook.save(output)
        spreadsheet = parse_document("财务底稿.xlsx", output.getvalue())

        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        check = next(
            item for item in result.checks
            if item.rule_id == "CROSS-DOCUMENT-FACT"
            and "应收账款" in item.title
        )
        self.assertEqual(check.status, "passed")
        self.assertEqual(check.source_locations[1].sheet_name, "客户明细")

    def test_nested_note_is_not_compared_to_statement_when_note_source_is_absent(self) -> None:
        """A statement aggregate alone is insufficient evidence for a detail row."""
        document = _parsed_table(
            "（1）客户明细",
            "项目 | 2026年6月30日余额\n应收账款 | 100",
            "2026年半年度报告.docx",
        )
        spreadsheet = _parsed_table(
            "合并资产表",
            "项目 | 2026年6月30日\n应收账款 | 999",
            "财务底稿.xlsx",
        )
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        self.assertFalse(any(
            item.rule_id == "CROSS-DOCUMENT-FACT"
            for item in result.checks
        ))

    def test_statement_fixed_assets_are_reconciled(self) -> None:
        """Fixed assets are a real statement metric and must not be filtered out."""
        document = _parsed_table(
            "合并资产负债表",
            "项目 | 2026年6月30日\n固定资产 | 100",
            "2026年半年度报告.docx",
        )
        spreadsheet = _parsed_table(
            "合并资产表",
            "项目 | 2026年6月30日\n固定资产 | 120",
            "财务底稿.xlsx",
        )
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        mismatch = next(
            item for item in result.checks
            if item.rule_id == "CROSS-DOCUMENT-FACT"
            and item.status == "failed"
            and "固定资产" in item.title
        )
        self.assertEqual(mismatch.reported_value, Decimal("100"))
        self.assertEqual(mismatch.calculated_value, Decimal("120"))

    def test_roe_eps_checks_do_not_cross_deducted_basis_with_prior_period(self) -> None:
        # Replace the simple test document with the actual note-table shape:
        # two profit bases in the rows, not two periods.
        document = _parsed_table(
            "2、净资产收益率及每股收益",
            "报告期利润 | 加权平均净资产收益率（%） | 每股收益 | \n"
            " |  | 基本每股收益 | 稀释每股收益\n"
            "归属于公司普通股股东的净利润 | -12.34% | -0.30 | -0.30\n"
            "扣除非经常损益后归属于普通股股东的净利润 | -14.42% | -0.36 | -0.36",
            "2026年半年度报告.docx",
        )
        spreadsheet = _parsed_table(
            "合并利润表",
            "项目 | 2026年1-6月 | 2025年1-6月\n"
            "归属于母公司股东的净利润 | -100 | -80\n"
            "基本每股收益 | -0.10 | -0.08\n"
            "稀释每股收益 | -0.10 | -0.08\n"
            "平均净资产 | 1000 | ",
            "财务底稿.xlsx",
        )
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        checks = [
            item for item in result.checks
            if item.rule_id.startswith("AUTHORITATIVE-ROE-EPS")
        ]
        self.assertEqual(len(checks), 6)
        failed = {item.inputs["metric"] for item in checks if item.status == "failed"}
        missing = {item.inputs["metric"] for item in checks if item.status == "needs_review"}
        self.assertEqual(failed, {"加权平均净资产收益率", "基本每股收益", "稀释每股收益"})
        self.assertEqual(missing, {"扣非后加权平均净资产收益率", "扣非后基本每股收益", "扣非后稀释每股收益"})
        self.assertFalse(any("扣非后" in item.title and item.status == "failed" for item in checks))

    def test_xlsx_and_xls_keep_sheet_and_row_provenance(self) -> None:
        for extension in (".xlsx", ".xls"):
            parsed = parse_document(f"财务表{extension}", _spreadsheet_bytes(extension=extension))
            self.assertEqual(parsed.file_type, extension.removeprefix(".").upper())
            self.assertEqual(parsed.blocks[0].sheet_name, "合并利润表")
            self.assertEqual(parsed.blocks[0].source_row_numbers, (1, 2, 3))

    def test_document_and_spreadsheet_mismatch_is_a_high_risk_with_two_locations(self) -> None:
        document = parse_document("2026年半年度报告.docx", _docx_bytes())
        spreadsheet = parse_document("财务表.xlsx", _spreadsheet_bytes(net_profit="19"))
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        checks = [item for item in result.checks if item.rule_id == "CROSS-DOCUMENT-FACT"]
        mismatch = next(item for item in checks if item.status == "failed" and "净利润" in item.title and "本期" in item.title)
        self.assertEqual(mismatch.severity, "high")
        self.assertEqual(len(mismatch.source_locations), 2)
        self.assertEqual(mismatch.source_locations[1].sheet_name, "合并利润表")
        self.assertEqual(mismatch.source_locations[1].spreadsheet_row, 3)

    def test_matching_values_are_marked_passed(self) -> None:
        document = parse_document("2026年半年度报告.docx", _docx_bytes())
        spreadsheet = parse_document("财务表.xls", _spreadsheet_bytes(extension=".xls"))
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        checks = [item for item in result.checks if item.rule_id == "CROSS-DOCUMENT-FACT"]
        self.assertTrue(checks)
        self.assertTrue(all(item.status == "passed" for item in checks))

    def test_all_reliably_matched_note_rows_are_reconciled(self) -> None:
        document = parse_document("财务报告附注.docx", _detailed_docx_bytes())
        spreadsheet = parse_document("财务底稿.xlsx", _detailed_spreadsheet_bytes(bad_debt="59114840.35"))
        result = FinancialAuditEngine().analyze(
            document,
            document_id="document-id",
            file_hash="document-hash",
            comparison_parsed=spreadsheet,
            comparison_document_id="spreadsheet-id",
        )
        checks = [item for item in result.checks if item.rule_id == "CROSS-DOCUMENT-FACT"]
        titles = [item.title for item in checks]
        self.assertTrue(any("单项计提坏账准备的应收账款" in title for title in titles))
        self.assertTrue(any("按组合计提坏账准备的应收账款" in title for title in titles))
        mismatch = next(item for item in checks if item.status == "failed" and "坏账准备" in item.title)
        self.assertEqual(mismatch.difference, Decimal("-1.00"))
        self.assertEqual(mismatch.source_locations[1].sheet_name, "应收账款附注")
        self.assertEqual(mismatch.source_locations[1].spreadsheet_row, 4)


if __name__ == "__main__":
    unittest.main()
