from __future__ import annotations

import io
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from docx import Document as WordDocument


WEB_ROOT = Path(__file__).resolve().parents[1]
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

from backend.config import Settings  # noqa: E402
from backend.database import Database  # noqa: E402
from backend.financial_audit.engine import ENGINE_VERSION, FinancialAuditEngine  # noqa: E402
from backend.financial_audit.extractor import FinancialExtractor  # noqa: E402
from backend.financial_audit.financial_cases import load_financial_cases  # noqa: E402
from backend.ocr import OCRPage  # noqa: E402
from backend.parsers import ParsedBlock, ParsedDocument, parse_document  # noqa: E402
from backend.security import AuthService, hash_password  # noqa: E402
from backend.services import PlatformService  # noqa: E402


def parsed_tables(*tables: tuple[str, str], filename: str = "2026年半年度报告.docx") -> ParsedDocument:
    blocks: list[ParsedBlock] = []
    offset = 0
    for index, (section, text) in enumerate(tables):
        blocks.append(ParsedBlock(text, index + 1, section, index, "table", offset, offset + len(text)))
        offset += len(text) + 1
    return ParsedDocument(filename, Path(filename).suffix.removeprefix(".").upper(), len(blocks), blocks)


def analyze(*tables: tuple[str, str], filename: str = "2026年半年度报告.docx"):
    return FinancialAuditEngine().analyze(
        parsed_tables(*tables, filename=filename), document_id="document-test", file_hash="hash-test"
    )


class FinancialFormulaTests(unittest.TestCase):
    def test_confirmed_financial_case_corpus_is_loaded(self) -> None:
        cases = load_financial_cases()
        self.assertTrue(cases)
        self.assertTrue(cases[0]["validated"])
        self.assertEqual(len(cases[0]["annotations"]), 19)

    def test_unresolved_confirmed_case_anchor_is_visible_for_manual_review(self) -> None:
        result = analyze((
            "4、其他关联方情况",
            "关联方名称 | 期末余额\n圳市达晨财智创业投资管理有限 | 0",
        ))
        checks = [
            check for check in result.checks
            if check.rule_id == "VALIDATED-CASE-ANCHOR"
        ]
        self.assertTrue(checks)
        self.assertTrue(any("不是关联方" in check.title for check in checks))
        self.assertTrue(all(check.status == "needs_review" for check in checks))
        self.assertTrue(all(check.source_locations for check in checks))

    def test_required_financial_note_amounts_are_not_silently_blank(self) -> None:
        result = analyze((
            "非经常性损益明细",
            "项目 | 本期金额 | 上期金额\n"
            "扣除所得税前非经常性损益合计 |  | \n"
            "减：所得税影响金额 |  | ",
        ))
        checks = [check for check in result.checks if check.rule_id == "REQUIRED-AMOUNT-MISSING"]
        self.assertEqual(len(checks), 2)
        self.assertTrue(all(check.status == "failed" for check in checks))
        self.assertTrue(all(check.case_references for check in checks))

    def test_repeated_specialized_fact_is_traceable_for_manual_scope_review(self) -> None:
        result = analyze(
            ("长期股权投资", "项目 | 期末余额\n长期股权投资 | 100"),
            ("长期股权投资明细", "项目 | 期末余额\n长期股权投资 | 101"),
        )
        checks = [check for check in result.checks if check.rule_id == "REPEATED-FACT-MISMATCH"]
        self.assertTrue(checks)
        self.assertEqual(checks[0].status, "needs_review")
        self.assertEqual(len(checks[0].source_locations), 2)

    def test_percentage_tolerance_and_material_difference(self) -> None:
        result = analyze((
            "主要财务指标",
            "项目 | 本报告期 | 上年同期 | 比上年同期增减\n"
            "正常四舍五入 | 120 | 100 | 20.005%\n"
            "超过容差 | 120 | 100 | 20.02%",
        ))
        checks = {check.title: check for check in result.checks if check.rule_id == "YOY-RECALCULATION"}
        self.assertEqual(checks["已核对无误：正常四舍五入"].status, "passed")
        self.assertEqual(checks["同比百分比计算不一致：超过容差"].status, "failed")
        self.assertEqual(checks["同比百分比计算不一致：超过容差"].difference, Decimal("0.0200"))

    def test_decimal_point_percentage_error_is_definite(self) -> None:
        result = analyze(
            (
                "合并利润表",
                "项目 | 本报告期 | 上年同期\n"
                "利润总额 | -13,980,078.52 | -10,000,000.00",
            ),
            (
                "非主营业务分析",
                "项目 | 金额 | 占利润总额比例\n"
                "投资收益 | -2,120,421.36 | 0.15%",
            ),
        )
        check = next(check for check in result.checks if check.rule_id == "PROFIT-TOTAL-RATIO")
        self.assertEqual(check.status, "failed")
        self.assertEqual(check.severity, "high")
        self.assertEqual(check.calculated_value, Decimal("15.17"))
        self.assertEqual(check.reported_value, Decimal("0.15"))

    def test_disclosure_row_proportions_use_total_ratio_as_denominator(self) -> None:
        result = analyze((
            "按欠款方归集的期末余额前五名的应收账款和合同资产情况",
            "单位名称 | 应收账款期末余额 | 合同资产期末余额 | 应收账款和合同资产期末余额合计数 | 占应收账款和合同资产期末余额合计数的比例 | 坏账准备和减值准备合计余额\n"
            "广州博望纵横网络科技有限公司 | 4000000.00 |  | 4000000.00 | 4.78% | 4000000.00\n"
            "深圳市金立通信设备有限公司 | 1600000.00 |  | 1600000.00 | 1.51% | 1600000.00\n"
            "凌雄科技（深圳）有限公司 | 1440000.00 |  | 1440000.00 | 1.36% | 1353600.00\n"
            "车易拍（北京）汽车技术服务有限公司 | 1200000.00 |  | 1200000.00 | 5.13% | 1200000.00\n"
            "雪豹（北京）国际传媒有限公司 | 1164471.82 |  | 1164471.82 | 1.10% | 61287.99\n"
            "合计 | 9404471.82 |  | 9404471.82 | 8.88% | 8214887.99",
        ))
        checks = [
            check for check in result.checks
            if check.rule_id == "TABLE-PROPORTION-RECALCULATION"
        ]
        by_label = {check.inputs["row_label"]: check for check in checks}
        self.assertEqual(by_label["广州博望纵横网络科技有限公司"].status, "failed")
        self.assertEqual(by_label["广州博望纵横网络科技有限公司"].calculated_value, Decimal("3.78"))
        self.assertEqual(by_label["车易拍（北京）汽车技术服务有限公司"].status, "failed")
        self.assertEqual(by_label["车易拍（北京）汽车技术服务有限公司"].calculated_value, Decimal("1.13"))
        self.assertEqual(by_label["深圳市金立通信设备有限公司"].status, "passed")
        self.assertEqual(by_label["广州博望纵横网络科技有限公司"].inputs["denominator_source"], "合计行金额与合计比例")
        self.assertEqual(len(by_label["广州博望纵横网络科技有限公司"].source_locations), 4)

    def test_disclosure_row_proportions_without_unique_denominator_need_review(self) -> None:
        result = analyze((
            "应收账款前五名",
            "单位名称 | 应收账款余额 | 占应收账款余额合计数的比例 | 坏账准备余额\n"
            "客户甲 | 100 | 10.00% | 0\n"
            "客户乙 | 200 | 20.00% | 0\n"
            "合计 | 300 |  | 0",
        ))
        checks = [
            check for check in result.checks
            if check.rule_id == "TABLE-PROPORTION-DENOMINATOR-MISSING"
        ]
        self.assertEqual(len(checks), 1)
        self.assertEqual(checks[0].status, "needs_review")
        self.assertFalse(checks[0].source_sufficient)

    def test_disclosure_row_proportions_do_not_treat_percentage_as_amount(self) -> None:
        result = analyze((
            "客户余额比例",
            "单位名称 | 应收账款余额 | 占应收账款余额合计数的比例\n"
            "客户甲 | 100 | 50.00%\n"
            "客户乙 | 100 | 50.00%\n"
            "合计 | 200 | 100.00%",
        ))
        checks = [
            check for check in result.checks
            if check.rule_id == "TABLE-PROPORTION-RECALCULATION"
        ]
        self.assertEqual(len(checks), 2)
        self.assertTrue(all(check.status == "passed" for check in checks))
        self.assertTrue(all(check.inputs["numerator"] in {Decimal("100"), Decimal("200")} for check in checks))

    def test_negative_yoy_uses_cash_flow_semantics_and_merges_repeated_disclosures(self) -> None:
        table = (
            "项目 | 本报告期 | 上年同期 | 比上年同期增减\n"
            "现金及现金等价物净增加额 | -80 | -100 | 20%"
        )
        result = analyze(("主要财务指标", table), ("现金流摘要", table))
        checks = [check for check in result.checks if check.rule_id == "YOY-NEGATIVE-SEMANTICS"]
        self.assertEqual(len(checks), 1)
        self.assertEqual(checks[0].status, "needs_review")
        self.assertEqual(checks[0].inputs["change_type"], "net_outflow_narrowed")
        self.assertIn("净流出收窄", str(checks[0].calculated_value))
        self.assertEqual(len(checks[0].source_locations), 6)

    def test_roe_uses_percentage_points_not_yoy_rate(self) -> None:
        result = analyze((
            "主要财务指标",
            "项目 | 本报告期 | 上年同期 | 比上年同期增减\n"
            "加权平均净资产收益率 | -3.14% | -2.79% | -0.35%",
        ))
        check = next(check for check in result.checks if check.rule_id == "PERCENTAGE-POINT-CHANGE")
        self.assertEqual(check.status, "passed")
        self.assertEqual(check.calculated_value, Decimal("-0.35"))
        self.assertIn("个百分点", check.explanation)

    def test_balance_profit_and_cash_flow_equations(self) -> None:
        result = analyze(
            (
                "合并资产负债表",
                "项目 | 期末余额 | 期初余额\n"
                "流动资产合计 | 60 | 50\n非流动资产合计 | 40 | 50\n资产总计 | 100 | 100\n"
                "流动负债合计 | 30 | 30\n非流动负债合计 | 20 | 20\n负债合计 | 50 | 50\n"
                "所有者权益合计 | 50 | 50\n负债和所有者权益总计 | 100 | 100",
            ),
            (
                "合并利润表",
                "项目 | 本报告期 | 上年同期\n"
                "营业利润 | 90 | 80\n营业外收入 | 20 | 10\n营业外支出 | 10 | 5\n利润总额 | 100 | 85\n"
                "所得税费用 | 20 | 17\n净利润 | 80 | 68\n归属于母公司所有者的净利润 | 70 | 60\n少数股东损益 | 10 | 8",
            ),
            (
                "合并现金流量表",
                "项目 | 本报告期 | 上年同期\n"
                "经营活动现金流入小计 | 100 | 80\n经营活动现金流出小计 | 60 | 50\n经营活动产生的现金流量净额 | 40 | 30\n"
                "投资活动现金流入小计 | 20 | 20\n投资活动现金流出小计 | 30 | 25\n投资活动产生的现金流量净额 | -10 | -5\n"
                "筹资活动现金流入小计 | 10 | 10\n筹资活动现金流出小计 | 5 | 4\n筹资活动产生的现金流量净额 | 5 | 6\n"
                "汇率变动对现金及现金等价物的影响 | 0 | 0\n现金及现金等价物净增加额 | 35 | 31\n"
                "期初现金及现金等价物余额 | 100 | 90\n期末现金及现金等价物余额 | 135 | 121",
            ),
        )
        statement_checks = [
            check for check in result.checks
            if check.category in {"资产负债表", "跨表勾稽", "利润表勾稽", "现金流勾稽"}
        ]
        self.assertGreaterEqual(len(statement_checks), 20)
        self.assertTrue(all(check.status == "passed" for check in statement_checks))

    def test_gross_margin_total_and_sign_aware_movement(self) -> None:
        result = analyze(
            (
                "主营业务分产品",
                "项目 | 营业收入 | 营业成本 | 毛利率\n"
                "企业服务 | 200 | 120 | 40%",
            ),
            (
                "分部收入",
                "项目 | 本期金额\n培训 | 60\n咨询 | 40\n合计 | 100",
            ),
            (
                "长期股权投资变动",
                "项目 | 期初余额 | 新增投资 | 减少投资 | 期末余额\n"
                "甲公司 | 100 | 20 | -30 | 90\n乙公司 | 100 | 20 | 30 | 90",
            ),
        )
        selected = [
            check for check in result.checks
            if check.rule_id in {"GROSS-MARGIN", "SIMPLE-TOTAL", "MOVEMENT-ROLLFORWARD"}
        ]
        self.assertEqual(len(selected), 4)
        self.assertTrue(all(check.status == "passed" for check in selected))

    def test_movement_ignores_paired_opening_and_closing_allowance_columns(self) -> None:
        result = analyze((
            "8、长期股权投资",
            "被投资单位 | 期初余额（账面价值） | 减值准备期初余额 | 本期增减变动 | | | | 期末余额（账面价值） | 减值准备期末余额\n"
            "| | | 追加投资 | 减少投资 | 计提减值准备 | 其他 | |\n"
            "北京黑马加速科技有限公司 | 45,608,947.07 | 20,306,570.32 | | | | | 45,608,947.07 | 20,306,570.32\n"
            "合计 | 45,608,947.07 | 20,306,570.32 | | | | | 45,608,947.07 | 20,306,570.32",
        ))
        movement_checks = [check for check in result.checks if check.rule_id == "MOVEMENT-ROLLFORWARD"]
        self.assertEqual(len(movement_checks), 2)
        self.assertTrue(all(check.status == "passed" for check in movement_checks))
        self.assertTrue(all("减值准备期初余额" not in check.inputs for check in movement_checks))

    def test_paired_balance_movement_still_detects_real_rollforward_errors(self) -> None:
        result = analyze((
            "8、长期股权投资",
            "被投资单位 | 期初余额（账面价值） | 减值准备期初余额 | 本期增减变动 | | | | 期末余额（账面价值） | 减值准备期末余额\n"
            "| | | 追加投资 | 减少投资 | 计提减值准备 | 其他 | |\n"
            "闭合项目 | 100.00 | 20.00 | 40.00 | 10.00 | 5.00 | | 125.00 | 25.00\n"
            "错误项目 | 100.00 | 20.00 | 40.00 | 10.00 | 5.00 | | 126.00 | 25.00",
        ))
        movement_checks = {
            check.title: check for check in result.checks
            if check.rule_id == "MOVEMENT-ROLLFORWARD"
        }
        self.assertEqual(movement_checks["已核对无误：闭合项目"].status, "passed")
        self.assertEqual(movement_checks["变动表不闭合：错误项目"].status, "failed")
        self.assertEqual(movement_checks["变动表不闭合：错误项目"].calculated_value, Decimal("125.00"))

    def test_unit_normalization_preserves_original_and_base_amount(self) -> None:
        text = "项目 | 期末余额\n货币资金 | 4,686.53"
        parsed = ParsedDocument(
            "附注.docx", "DOCX", 1,
            [
                ParsedBlock("单位：万元", 1, "货币资金", 0, "paragraph", 0, 5),
                ParsedBlock(text, 1, "货币资金", 1, "table", 6, 6 + len(text)),
            ],
        )
        extracted = FinancialExtractor().extract(parsed, "document-unit")
        cell = next(cell for cell in extracted.tables[0].cells if cell.row_label == "货币资金" and cell.normalized_value)
        self.assertEqual(cell.original_unit, "万元")
        self.assertEqual(cell.normalized_value, Decimal("4686.53"))
        self.assertEqual(cell.base_amount, Decimal("46865300.00"))

    def test_source_registry_does_not_invent_missing_source(self) -> None:
        text = "国家统计局数据显示，行业规模继续增长。"
        parsed = ParsedDocument(
            "2026年半年度报告.docx", "DOCX", 1,
            [ParsedBlock(text, 1, "行业情况", 0, "paragraph", 0, len(text))],
        )
        result = FinancialAuditEngine().analyze(parsed, document_id="source-document", file_hash="source-hash")
        check = next(check for check in result.checks if check.rule_id == "SOURCE-SUFFICIENCY")
        self.assertEqual(check.status, "source_issue")
        self.assertFalse(check.source_sufficient)


class FinancialDocumentParsingTests(unittest.TestCase):
    def test_docx_table_reaches_financial_engine_with_cell_provenance(self) -> None:
        document = WordDocument()
        document.add_heading("合并资产负债表", level=1)
        table = document.add_table(rows=4, cols=3)
        values = [
            ("项目", "期末余额", "期初余额"),
            ("流动资产合计", "60", "50"),
            ("非流动资产合计", "40", "50"),
            ("资产总计", "100", "100"),
        ]
        for row_index, row in enumerate(values):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
        buffer = io.BytesIO()
        document.save(buffer)
        parsed = parse_document("2026年半年度报告.docx", buffer.getvalue())
        extracted = FinancialExtractor().extract(parsed, "docx-source")
        self.assertEqual(len(extracted.tables), 1)
        amount = next(cell for cell in extracted.tables[0].cells if cell.row_label == "资产总计" and cell.column_index == 1)
        self.assertEqual(amount.normalized_value, Decimal("100"))
        self.assertEqual(amount.location.table_id, extracted.tables[0].table_id)
        self.assertEqual(amount.location.source_text, "资产总计 | 100 | 100")

    def test_scanned_pdf_ocr_table_is_reconstructed_and_audited(self) -> None:
        try:
            from pypdf import PdfWriter
        except ImportError:
            self.skipTest("pypdf is not installed")
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        class FakeOCRService:
            engine_name = "fake-local-ocr"

            @staticmethod
            def recognize_pdf_pages(_content: bytes, pages: list[int]) -> list[OCRPage]:
                self.assertEqual(pages, [1])
                return [OCRPage(1, (
                    "项目    本报告期    上年同期    比上年同期增减\n"
                    "营业收入    120    100    20%"
                ))]

        parsed = parse_document("2026年半年度报告.pdf", buffer.getvalue(), FakeOCRService())
        self.assertTrue(parsed.ocr_used)
        self.assertEqual([block.block_type for block in parsed.blocks], ["table"])
        result = FinancialAuditEngine().analyze(parsed, document_id="pdf-source", file_hash="pdf-hash")
        yoy = next(check for check in result.checks if check.rule_id == "YOY-RECALCULATION")
        self.assertEqual(yoy.status, "passed")
        self.assertEqual(yoy.source_locations[0].page, 1)


class FinancialPersistenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.settings = Settings(
            app_env="test",
            ai_mode="local_rag",
            database_path=root / "runtime" / "test.db",
            upload_root=root / "runtime" / "uploads",
            knowledge_root=root / "knowledge",
            max_upload_bytes=20 * 1024 * 1024,
            session_hours=12,
            bootstrap_admin_username="admin",
            bootstrap_admin_password="TestAdmin123!",
            app_secret="test-secret-not-used-by-cookie-sessions",
            aliyun_api_key="",
            aliyun_base_url="",
            aliyun_model="",
            aliyun_chat_path="/chat/completions",
            request_timeout_seconds=2,
            embedding_batch_size=8,
            vector_backend="sqlite",
            database_url="",
            vector_db_url="",
            input_cost_per_million=2,
            output_cost_per_million=8,
        )
        self.database = Database(self.settings.database_path)
        self.database.migrate()
        auth = AuthService(self.database, self.settings)
        auth.bootstrap()
        login = auth.login("admin", "TestAdmin123!", "127.0.0.1")
        assert login is not None
        self.admin = login[1]
        self.platform = PlatformService(self.settings, self.database)
        self.platform.initialize()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_result_owner_isolation_hash_cache_and_report_formats(self) -> None:
        content = (
            "2026年半年度报告\n"
            "国家统计局数据显示，行业规模继续增长。"
        ).encode("utf-8")
        first = self.platform.create_task(
            "2026年半年度报告.txt", content, "半年报财务测试", "定期报告一致性审查",
            "创业板定期报告规则包", self.admin, "127.0.0.1",
        )
        payload = self.platform.get_financial_audit(first["taskId"], self.admin)
        self.assertEqual(payload["engineVersion"], ENGINE_VERSION)
        self.assertFalse(payload["externalAiUsed"])

        other_id = "financial-other-user"
        now = "2026-08-27T00:00:00+00:00"
        self.database.execute(
            "INSERT INTO users (id,username,password_hash,display_name,role,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
            (other_id, "financial-other", hash_password("OtherUser123!"), "其他账号", "user", now, now),
        )
        other = {"id": other_id, "username": "financial-other", "display_name": "其他账号", "role": "user"}
        with self.assertRaises(PermissionError):
            self.platform.get_financial_audit(first["taskId"], other)

        second = self.platform.create_task(
            "2026年半年度报告.txt", content, "半年报财务复核", "定期报告一致性审查",
            "创业板定期报告规则包", self.admin, "127.0.0.1",
        )
        second_task = self.platform.get_task_result(second["taskId"], user=self.admin)
        self.assertTrue(second_task["financialAudit"]["cacheReused"])
        first_financial_risks = [risk for risk in first["risks"] if risk["type"] == "财务数据勾稽审计"]
        second_financial_risks = [risk for risk in second_task["risks"] if risk["type"] == "财务数据勾稽审计"]
        self.assertTrue(first_financial_risks)
        self.assertEqual(
            {risk["calculation"]["check_id"] for risk in first_financial_risks},
            {risk["calculation"]["check_id"] for risk in second_financial_risks},
        )
        self.assertTrue(
            {risk["id"] for risk in first_financial_risks}.isdisjoint(
                {risk["id"] for risk in second_financial_risks}
            )
        )

        for report_format, signature in (("html", b"<!doctype html>"), ("docx", b"PK"), ("pdf", b"%PDF")):
            filename, report = self.platform.get_financial_audit_report(
                first["taskId"], report_format, self.admin, "127.0.0.1"
            )
            self.assertTrue(filename.lower().endswith(f".{report_format}"))
            self.assertTrue(report.startswith(signature))

    def test_stale_financial_result_refreshes_task_risks_without_losing_review_identity(self) -> None:
        content = (
            "2026年半年度报告\n"
            "国家统计局数据显示，行业规模继续增长。"
        ).encode("utf-8")
        created = self.platform.create_task(
            "2026年半年度报告.txt", content, "旧引擎刷新测试", "定期报告一致性审查",
            "创业板定期报告规则包", self.admin, "127.0.0.1",
        )
        old_risk = next(risk for risk in created["risks"] if risk["type"] == "财务数据勾稽审计")
        self.database.execute(
            "UPDATE financial_audit_results SET engine_version=? WHERE task_id=?",
            ("1.0.3", created["taskId"]),
        )

        refreshed = self.platform.get_task_result(created["taskId"], user=self.admin)
        self.assertEqual(refreshed["financialAudit"]["engineVersion"], ENGINE_VERSION)
        current_risk = next(risk for risk in refreshed["risks"] if risk["type"] == "财务数据勾稽审计")
        self.assertEqual(current_risk["id"], old_risk["id"])
        self.assertEqual(current_risk["calculation"]["engine_version"], ENGINE_VERSION)


if __name__ == "__main__":
    unittest.main()
