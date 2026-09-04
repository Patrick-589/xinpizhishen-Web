import io
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from docx import Document as WordDocument


WEB_ROOT = Path(__file__).resolve().parents[1]
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

from backend.config import Settings  # noqa: E402
from backend.database import Database  # noqa: E402
from backend.security import AuthService  # noqa: E402
from backend.services import PlatformService  # noqa: E402


def xlsx_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def docx_bytes(paragraphs: list[str]) -> bytes:
    document = WordDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class MultipleComparisonTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        settings = Settings(
            app_env="test",
            ai_mode="local_rag",
            database_path=root / "runtime" / "test.db",
            upload_root=root / "runtime" / "uploads",
            knowledge_root=root / "knowledge",
            max_upload_bytes=20 * 1024 * 1024,
            session_hours=12,
            bootstrap_admin_username="admin",
            bootstrap_admin_password="TestAdmin123!",
            app_secret="test-secret",
            aliyun_api_key="",
            aliyun_base_url="",
            aliyun_model="",
            aliyun_chat_path="/chat/completions",
            request_timeout_seconds=2,
            embedding_batch_size=8,
            vector_backend="sqlite",
            database_url="",
            vector_db_url="",
            input_cost_per_million=2,
            output_cost_per_million=8,
        )
        self.database = Database(settings.database_path)
        self.database.migrate()
        self.auth = AuthService(self.database, settings)
        self.auth.bootstrap()
        self.user = self.auth.login("admin", "TestAdmin123!", "127.0.0.1")[1]
        self.platform = PlatformService(settings, self.database)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_two_spreadsheets_are_saved_read_and_checked_independently(self) -> None:
        primary = xlsx_bytes([
            ["项目", "本报告期", "上年同期"],
            ["营业收入", 120, 100],
            ["资产总计", 300, 280],
        ])
        first = xlsx_bytes([
            ["项目", "本报告期", "上年同期"],
            ["营业收入", 120, 100],
            ["资产总计", 300, 280],
        ])
        second = xlsx_bytes([
            ["项目", "本报告期", "上年同期"],
            ["营业收入", 121, 100],
            ["资产总计", 301, 280],
        ])
        result = self.platform.create_task(
            "主报告.xlsx", primary, "多表格审查", "定期报告一致性审查", "自动匹配",
            self.user, "127.0.0.1",
            comparison_documents=[("财务底稿一.xlsx", first), ("财务底稿二.xlsx", second)],
        )
        comparisons = result["comparisonDocuments"]
        self.assertEqual([item["filename"] for item in comparisons], ["财务底稿一.xlsx", "财务底稿二.xlsx"])
        self.assertTrue(all(item.get("documentId") for item in comparisons))
        rows = self.database.fetch_all(
            "SELECT id,filename,metadata_json FROM documents WHERE task_id=? ORDER BY created_at,id",
            (result["taskId"],),
        )
        self.assertEqual(
            {row["filename"] for row in rows},
            {"主报告.xlsx", "财务底稿一.xlsx", "财务底稿二.xlsx"},
        )
        checks = self.database.fetch_one(
            "SELECT checks_json FROM financial_audit_results WHERE task_id=?", (result["taskId"],)
        )
        check_text = checks["checks_json"] if checks else ""
        self.assertIn("财务底稿一.xlsx", check_text)
        self.assertIn("财务底稿二.xlsx", check_text)
        for comparison in comparisons:
            preview = self.platform.get_task_related_document_preview(
                result["taskId"], comparison["documentId"], self.user, "127.0.0.1",
            )
            self.assertEqual(preview["filename"], comparison["filename"])
            self.assertGreater(preview["paragraphCount"], 0)

    def test_multiple_word_documents_are_each_available_and_summary_results_are_separate(self) -> None:
        primary = docx_bytes(["报告摘要", "报告期：2026年6月30日", "营业收入：100"])
        first = docx_bytes(["完整报告一", "报告期：2026年6月30日", "营业收入：100"])
        second = docx_bytes(["完整报告二", "报告期：2026年6月30日", "营业收入：90"])
        result = self.platform.create_task(
            "报告摘要.docx", primary, "多文档审查", "定期报告一致性审查", "自动匹配",
            self.user, "127.0.0.1",
            comparison_documents=[("完整报告一.docx", first), ("完整报告二.docx", second)],
        )
        self.assertEqual(len(result["comparisonDocuments"]), 2)
        self.assertEqual(len(result["summaryComparisons"]), 2)
        self.assertTrue(all(item["status"] == "completed" for item in result["summaryComparisons"]))
        self.assertTrue(result["summaryComparison"].get("multi"))
        for comparison in result["comparisonDocuments"]:
            filename, content = self.platform.get_task_related_document(
                result["taskId"], comparison["documentId"], self.user, "127.0.0.1",
            )
            self.assertEqual(filename, comparison["filename"])
            self.assertEqual(content, {"完整报告一.docx": first, "完整报告二.docx": second}[filename])


if __name__ == "__main__":
    unittest.main()
