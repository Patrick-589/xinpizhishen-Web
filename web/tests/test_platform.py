from __future__ import annotations

import sys
import tempfile
import unittest
import json
import io
import hashlib
import time
import zipfile
from decimal import Decimal
from xml.etree import ElementTree
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from docx import Document as WordDocument


WEB_ROOT = Path(__file__).resolve().parents[1]
if str(WEB_ROOT) not in sys.path:
    sys.path.insert(0, str(WEB_ROOT))

from backend.config import Settings  # noqa: E402
from backend.ai import AIServiceError, AliyunProvider  # noqa: E402
from backend.casting import review_casting  # noqa: E402
from backend.chunking import SecuritiesChunker  # noqa: E402
from backend.database import Database  # noqa: E402
from backend.docx_comments import create_annotated_docx  # noqa: E402
from backend.embeddings import MockEmbeddingProvider, create_embedding_provider  # noqa: E402
from backend.financial_change import calculate_financial_change  # noqa: E402
from backend.parsers import PARSER_VERSION, ParsedBlock, ParsedDocument, normalize_text, parse_document  # noqa: E402
from backend.ocr import LocalOCRService, OCRPage, normalize_ocr_text  # noqa: E402
from backend.periodic_compare import (  # noqa: E402
    PeriodicReportIdentity,
    compare_periodic_report,
    extract_narrative_sections,
    select_baseline_rows,
)
from backend.rag import RAGContext  # noqa: E402
from backend.rules import DeterministicRuleEngine  # noqa: E402
from backend.security import ROLE_PERMISSIONS, AuthService, hash_password  # noqa: E402
from backend.services import PlatformService, detect_review_configuration  # noqa: E402
from backend.vector_store import FaissVectorStore, SearchResult  # noqa: E402


class FinancialChangeTests(unittest.TestCase):
    def test_required_financial_change_boundaries(self) -> None:
        cases = [
            ("normal growth", "100", "120", "profit_loss", "normal_growth", Decimal("20.00"), "同比增长"),
            ("normal decline", "100", "80", "profit_loss", "normal_decline", Decimal("-20.00"), "同比下降"),
            ("loss narrowed", "-100", "-80", "profit_loss", "loss_narrowed", None, "亏损收窄"),
            ("loss expanded", "-100", "-120", "profit_loss", "loss_expanded", None, "亏损扩大"),
            ("negative to positive", "-100", "50", "profit_loss", "negative_to_positive", None, "由负转正"),
            ("positive to negative", "100", "-50", "profit_loss", "positive_to_negative", None, "由正转负"),
            ("zero to positive", "0", "100", "profit_loss", "previous_zero", None, "上期为0，本期增加"),
            ("zero to negative", "0", "-100", "profit_loss", "previous_zero", None, "上期为0，本期减少"),
            ("zero unchanged", "0", "0", "profit_loss", "unchanged", None, "无变化"),
            (
                "operating cash outflow narrowed", "-42,492,502.59", "-24,379,248.64", "cash_flow",
                "net_outflow_narrowed", None, "净流出收窄",
            ),
            (
                "investing cash outflow narrowed", "-98,847,347.67", "-22,237,942.79", "cash_flow",
                "net_outflow_narrowed", None, "净流出收窄",
            ),
            (
                "financing cash outflow expanded", "-4,214,096.62", "-10,971,363.00", "cash_flow",
                "net_outflow_expanded", None, "净流出扩大",
            ),
        ]
        for name, previous, current, semantic, change_type, display_rate, display_text in cases:
            with self.subTest(name=name):
                result = calculate_financial_change(
                    current,
                    previous,
                    semantic=semantic,
                    label="测试项目",
                )
                self.assertEqual(result.current, Decimal(current.replace(",", "")))
                self.assertEqual(result.previous, Decimal(previous.replace(",", "")))
                self.assertEqual(result.absolute_change, result.current - result.previous)
                self.assertEqual(result.change_type, change_type)
                self.assertEqual(result.yoy_rate, display_rate)
                self.assertEqual(result.display_yoy_rate, display_rate)
                self.assertIn(display_text, result.display_change_text)

    def test_raw_rate_is_separate_from_display_rate(self) -> None:
        result = calculate_financial_change("1,827,920.96", "-1,204,119.50", label="财务费用")
        self.assertEqual(result.change_type, "negative_to_positive")
        self.assertEqual(result.raw_yoy_rate, Decimal("251.81"))
        self.assertIsNone(result.yoy_rate)
        self.assertIsNone(result.display_yoy_rate)
        self.assertEqual(result.display_change_text, "财务费用由负转正")


class PlatformIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        root = Path(self.temporary.name)
        self.settings = Settings(
            app_env="test",
            ai_mode="local_rag",
            database_path=root / "runtime" / "test.db",
            upload_root=root / "runtime" / "uploads",
            knowledge_root=root / "knowledge",
            max_upload_bytes=20 * 1024 * 1024,
            session_hours=12,
            bootstrap_admin_username="admin",
            bootstrap_admin_password="TestAdmin123!",
            app_secret="test-secret-not-used-by-cookie-sessions",
            aliyun_api_key="",
            aliyun_base_url="",
            aliyun_model="",
            aliyun_chat_path="/chat/completions",
            request_timeout_seconds=2,
            embedding_batch_size=8,
            vector_backend="sqlite",
            database_url="",
            vector_db_url="",
            input_cost_per_million=2,
            output_cost_per_million=8,
        )
        self.database = Database(self.settings.database_path)
        self.database.migrate()
        self.auth = AuthService(self.database, self.settings)
        self.auth.bootstrap()
        login = self.auth.login("admin", "TestAdmin123!", "127.0.0.1")
        assert login is not None
        self.admin = login[1]
        self.platform = PlatformService(self.settings, self.database)
        self.platform.initialize()

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_parser_replaces_invalid_unicode_surrogates(self) -> None:
        normalized = normalize_text("法规正文\ud835损坏字形")
        self.assertEqual(normalized, "法规正文\ufffd损坏字形")
        normalized.encode("utf-8")

    def test_primary_regulation_basis_is_one_precise_current_article(self) -> None:
        selected = PlatformService._select_primary_regulation_basis([
            {
                "title": "创业板上市公司持续监管办法",
                "section": "总则",
                "content": "适用于创业板上市公司。",
                "citation_kind": "local_rag",
            },
            {
                "title": "深圳证券交易所创业板股票上市规则",
                "article": "第8.7.5条",
                "detail": "达到标准的重大交易应履行审议及披露义务。",
                "citation_kind": "deterministic_rule",
                "source_url": "https://www.szse.cn/rule/test",
            },
            {
                "title": "深圳证券交易所创业板股票上市规则",
                "article": "第8.7.6条",
                "detail": "其他事项。",
                "citation_kind": "deterministic_rule",
            },
        ])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected[0]["article"], "第8.7.5条")

    def test_end_to_end_local_rag_review_memory_and_audit(self) -> None:
        policy = (
            "创业黑马关联交易管理制度\n"
            "第一条 关联交易提交证券部预审。\n"
            "第二条 达到公司章程规定标准的事项应履行董事会审议程序。\n"
            "第三条 具体金额门槛应以当前有效公司章程及监管规则为准。"
        ).encode("utf-8")
        metadata = {
            "document_type": "internal_rules",
            "company": "创业黑马科技集团股份有限公司",
            "title": "创业黑马关联交易管理制度",
            "source": "测试资料",
            "version": "测试版",
        }
        imported = self.platform.import_knowledge("关联交易制度.txt", policy, metadata, self.admin, "127.0.0.1")
        self.assertEqual(imported["status"], "success")
        self.assertGreater(imported["document"]["chunks"], 0)
        duplicate = self.platform.import_knowledge("关联交易制度.txt", policy, metadata, self.admin, "127.0.0.1")
        self.assertEqual(duplicate["status"], "skipped")

        draft = (
            "关于对外投资暨关联交易的公告\n"
            "一、交易概述\n"
            "本次交易总规模为10000万元，公司出资3500万元，占总规模的30%。\n"
            "本事项构成关联交易，确保公司未来持续稳定增长。\n"
        ).encode("utf-8")
        result = self.platform.create_task(
            "关联交易公告.txt", draft, "关联交易公告审查", "关联交易专项审查",
            "关联交易专项规则包", self.admin, "127.0.0.1",
        )
        self.assertGreaterEqual(result["summary"]["total"], 3)
        self.assertEqual(result["analysisMode"], "local_rag")
        self.assertFalse(result["aiUsed"])
        self.assertTrue(any(risk["source_documents"] for risk in result["risks"]))
        wording_risk = next(risk for risk in result["risks"] if risk["type"] == "表述风险")
        self.assertIn("1. 核对", wording_risk["suggestion"])
        self.assertIn("2. 修改", wording_risk["suggestion"])
        self.assertIn("完成标准", wording_risk["suggestion"])
        self.assertIn("本地 RAG", wording_risk["suggestion"])
        self.assertIn("标记误报", wording_risk["suggestion"])
        task_citations = [
            source
            for risk in result["risks"]
            for source in risk["source_documents"]
            if source["document_id"] == result["documentId"]
        ]
        self.assertFalse(task_citations)
        self.assertTrue(any(
            source.get("document_type") == "internal_rules"
            for risk in result["risks"] for source in risk["source_documents"]
        ))

        task_id = result["taskId"]
        downloaded_name, downloaded_content = self.platform.get_task_document(task_id, self.admin, "127.0.0.1")
        self.assertEqual(downloaded_name, "关联交易公告.txt")
        self.assertEqual(downloaded_content, draft)
        original_preview = self.platform.get_task_document_preview(task_id, self.admin, "127.0.0.1")
        self.assertEqual(original_preview["filename"], "关联交易公告.txt")
        self.assertIn("本事项构成关联交易", "\n".join(original_preview["paragraphs"]))
        self.assertGreaterEqual(original_preview["parserVersion"], 2)
        response = self.platform.chat(task_id, "公司关联交易预审要求是什么？", None, self.admin, "127.0.0.1")
        self.assertEqual(response["mode"], "local_rag")
        self.assertIn("本地知识检索模式", response["answer"])
        self.assertTrue(response["source_documents"])
        conversations = self.platform.list_conversations(self.admin, task_id)
        self.assertEqual(len(conversations), 1)
        messages = self.platform.get_messages(response["conversation_id"], self.admin)
        self.assertEqual([message["role"] for message in messages], ["user", "assistant"])
        memory = self.database.fetch_one("SELECT summary FROM memory_summaries WHERE task_id=?", (task_id,))
        self.assertIsNotNone(memory)
        continued = self.platform.chat(task_id, "继续说明刚才的风险。", None, self.admin, "127.0.0.1")
        self.assertNotEqual(continued["conversation_id"], response["conversation_id"])
        latest_chat_audit = self.database.fetch_one(
            "SELECT after_json FROM audit_logs WHERE action='ai_chat' ORDER BY created_at DESC LIMIT 1"
        )
        self.assertTrue(json.loads(latest_chat_audit["after_json"])["task_memory_loaded"])

        risk_id = result["risks"][0]["id"]
        reviewed = self.platform.review_risk(risk_id, "confirm", "测试确认", self.admin, "127.0.0.1")
        self.assertEqual(reviewed["status"], "confirmed")
        restored = self.platform.get_task_result(task_id, user=self.admin)
        self.assertEqual(next(risk for risk in restored["risks"] if risk["id"] == risk_id)["status"], "confirmed")
        report = self.platform.generate_report(task_id, self.admin, "127.0.0.1")
        self.assertEqual(report["reviewSummary"]["confirmed"], 1)
        self.assertTrue(report["generatedAt"])
        reopened = self.platform.review_risk(risk_id, "reopen", "撤销测试确认", self.admin, "127.0.0.1")
        self.assertEqual(reopened["status"], "pending")
        restored = self.platform.get_task_result(task_id, user=self.admin)
        self.assertEqual(next(risk for risk in restored["risks"] if risk["id"] == risk_id)["status"], "pending")
        review_actions = self.database.fetch_all(
            "SELECT action,before_json,after_json FROM reviews WHERE risk_id=? ORDER BY created_at", (risk_id,)
        )
        self.assertEqual([item["action"] for item in review_actions], ["confirm", "reopen"])
        self.assertEqual(json.loads(review_actions[-1]["before_json"])["status"], "confirmed")
        self.assertEqual(json.loads(review_actions[-1]["after_json"])["status"], "pending")
        self.assertGreaterEqual(len(self.platform.audit_logs(self.admin)), 5)

    def test_local_rag_recommendation_is_detailed_and_does_not_auto_select_template_option(self) -> None:
        recommendation = self.platform._local_rag_recommendation(
            {
                "type": "模板选项完整性",
                "title": "必选项未勾选：无形资产减值测试情况",
                "description": "适用和不适用两个选项均未勾选。",
                "source_text": "□适用 □不适用",
                "recommendation": "根据实际情况仅勾选一项。",
            },
            [{
                "document_type": "regulations",
                "title": "半年度报告内容与格式准则",
                "page": 8,
                "section": "第十条",
                "version": "现行版本",
                "content": "上市公司应当按照规定的内容和格式编制半年度报告。",
            }],
        )
        self.assertIn("1. 核对", recommendation)
        self.assertIn("2. 修改", recommendation)
        self.assertIn("系统不得代替作出选择", recommendation)
        self.assertIn("☑适用 □不适用", recommendation)
        self.assertIn("□适用 ☑不适用", recommendation)
        self.assertIn("完成标准", recommendation)
        self.assertIn("《半年度报告内容与格式准则》（第十条 · 现行版本）", recommendation)
        self.assertIn("上市公司应当按照规定的内容和格式编制半年度报告", recommendation)

    def test_review_configuration_is_detected_from_filename_and_content(self) -> None:
        cases = (
            ("2026年半年度报告.docx", "第一节 重要提示\n公司简介和主要财务指标", "定期报告一致性审查", "创业板定期报告规则包"),
            ("关联交易公告.docx", "本事项构成关联交易，关联董事回避表决。", "关联交易专项审查", "关联交易专项规则包"),
            ("子公司增资协议.docx", "公司拟对全资子公司增资，交易金额待确认。", "重大事项披露审查", "创业板持续监管规则包"),
            ("采购合同.docx", "本合同金额及合同标的以附件为准。", "重大事项披露审查", "创业板持续监管规则包"),
        )
        for filename, content, expected_task_type, expected_rule_pack in cases:
            with self.subTest(filename=filename):
                parsed = parse_document(filename.replace(".docx", ".txt"), content.encode("utf-8"))
                detected = detect_review_configuration(filename, parsed)
                self.assertEqual(detected["taskType"], expected_task_type)
                self.assertEqual(detected["rulePack"], expected_rule_pack)
                self.assertFalse(detected["needsReview"])

        unknown = parse_document("普通材料.txt", "关于日常经营情况的说明。".encode("utf-8"))
        fallback = detect_review_configuration("普通材料.txt", unknown)
        self.assertEqual(fallback["taskType"], "临时公告通用审查")
        self.assertEqual(fallback["rulePack"], "上市公司信披基础包")
        self.assertTrue(fallback["needsReview"])

    def test_periodic_baseline_selection_excludes_summaries_and_cancelled_reports(self) -> None:
        rows = [
            {"id": "half", "title": "2025年半年度报告", "publication_date": "2025-08-26"},
            {"id": "half-summary", "title": "2025年半年度报告摘要", "publication_date": "2025-08-26"},
            {"id": "annual", "title": "2025年年度报告", "publication_date": "2026-04-08"},
            {"id": "annual-cancelled", "title": "2025年年度报告（已取消）", "publication_date": "2026-04-09"},
        ]
        selected = select_baseline_rows(rows, PeriodicReportIdentity(2026, "half-year"))
        self.assertEqual([item["id"] for item in selected], ["half", "annual"])

    def test_historical_baseline_parser_cache_survives_memory_cache_clear(self) -> None:
        content = "创业黑马科技集团股份有限公司2025年半年度报告\n主要会计数据和财务指标。".encode("utf-8")
        digest = hashlib.sha256(content).hexdigest()
        source = self.settings.knowledge_root / "announcements" / f"{digest[:16]}_2025年半年度报告.txt"
        source.parent.mkdir(parents=True, exist_ok=True)
        source.write_bytes(content)
        row = {"file_path": str(source), "file_hash": digest}
        with patch("backend.services.parse_document", wraps=parse_document) as mocked_parser:
            first = self.platform._load_cached_baseline_document(row)
            self.platform._parsed_baseline_cache.clear()
            second = self.platform._load_cached_baseline_document(row)
        self.assertEqual(mocked_parser.call_count, 1)
        self.assertEqual(first.paragraphs, second.paragraphs)
        cache_file = self.settings.upload_root.parent / "parsed-cache" / f"{digest}-v{PARSER_VERSION}.json"
        self.assertTrue(cache_file.is_file())

    def test_periodic_comparison_checks_same_period_year_end_and_statement_opening_values(self) -> None:
        current_text = (
            "2026年半年度报告\n"
            "四、主要会计数据和财务指标\n"
            "营业收入（元） | 80.00 | 70.00 | 14.29%\n"
            "总资产（元） | 120.00 | 100.00 | 20.00%\n"
            "1、合并资产负债表\n"
            "货币资金 | 90.00 | 80.00\n"
            "资产总计 | 120.00 | 100.00\n"
            "2、母公司资产负债表\n"
            "1、货币资金\n"
            "合计 | 90.00 | 80.00\n"
        )
        half_text = (
            "2025年半年度报告\n"
            "四、主要会计数据和财务指标\n"
            "营业收入（元） | 71.00 | 60.00 | 18.33%\n"
            "总资产（元） | 110.00 | 90.00 | 22.22%\n"
        )
        annual_text = (
            "2025年年度报告\n"
            "五、主要会计数据和财务指标\n"
            "总资产（元） | 101.00 | 90.00 | 12.22%\n"
            "1、合并资产负债表\n"
            "货币资金 | 81.00 | 75.00\n"
            "资产总计 | 101.00 | 90.00\n"
            "2、母公司资产负债表\n"
        )
        root = Path(self.temporary.name)
        half_path = root / "2025年半年度报告.txt"
        annual_path = root / "2025年年度报告.txt"
        half_path.write_text(half_text, encoding="utf-8")
        annual_path.write_text(annual_text, encoding="utf-8")
        current = parse_document("2026年半年度报告.txt", current_text.encode("utf-8"))
        rows = [
            {
                "id": "half", "title": "2025年半年度报告", "publication_date": "2025-08-26",
                "source": "巨潮资讯网", "source_url": "https://example.test/half.pdf",
                "file_path": str(half_path),
            },
            {
                "id": "annual", "title": "2025年年度报告", "publication_date": "2026-04-08",
                "source": "巨潮资讯网", "source_url": "https://example.test/annual.pdf",
                "file_path": str(annual_path),
            },
        ]
        risks, summary = compare_periodic_report(current, "2026年半年度报告.txt", rows)
        titles = [risk["title"] for risk in risks]
        self.assertEqual(summary["status"], "completed")
        self.assertEqual([item["title"] for item in summary["baselines"]], ["2025年半年度报告", "2025年年度报告"])
        self.assertEqual(summary["historicalMismatches"], 4)
        self.assertEqual(summary["internalMismatches"], 0)
        self.assertIn("跨期比较数不一致：营业收入（元）", titles)
        self.assertEqual(titles.count("跨期比较数不一致：总资产（元）"), 1)
        self.assertIn("跨期比较数不一致：货币资金", titles)
        self.assertIn("跨期比较数不一致：资产总计", titles)
        revenue = next(risk for risk in risks if risk["title"].endswith("营业收入（元）"))
        self.assertEqual(revenue["source_documents"][0]["source_url"], "https://example.test/half.pdf")
        self.assertEqual(revenue["source_documents"][0]["citation_kind"], "period_comparison")
        self.assertEqual(
            [item["article"] for item in revenue["regulation_basis"]],
            ["第5.2.4条", "第五十三条"],
        )

    def test_periodic_financial_checks_recalculate_equations_percentages_and_levels(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "二、主营业务分析\n"
                "销售费用 | 50.00 | 100.00 | -40.00%\n"
                "销售费用为50.00元，同比下降45.00%\n"
                "3、合并利润表\n"
                "营业利润 | 100.00 | 90.00\n"
                "营业外收入 | 10.00 | 5.00\n"
                "营业外支出 | 5.00 | 1.00\n"
                "利润总额 | 120.00 | 94.00\n"
                "所得税费用 | 20.00 | 10.00\n"
                "净利润 | 100.00 | 84.00\n"
                "4、母公司利润表\n"
                "5、合并现金流量表\n"
                "现金及现金等价物净增加额 | 10.00 | 8.00\n"
                "期初现金及现金等价物余额 | 100.00 | 90.00\n"
                "期末现金及现金等价物余额 | 115.00 | 98.00\n"
                "6、母公司现金流量表"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        by_title = {risk["title"]: risk for risk in risks}
        self.assertEqual(by_title["利润总额勾稽不一致"]["level"], "high")
        self.assertEqual(by_title["期末现金及现金等价物余额勾稽不一致"]["level"], "high")
        self.assertEqual(by_title["同比比例计算不一致：销售费用"]["level"], "medium")
        self.assertEqual(by_title["文字分析与表格同比比例不一致：销售费用"]["level"], "medium")
        narrative_risk = by_title["文字分析与表格同比比例不一致：销售费用"]
        self.assertEqual(len(narrative_risk["location"]["items"]), 2)
        self.assertEqual(
            [item["label"] for item in narrative_risk["location"]["items"]],
            ["主营业务分析表：同比比例", "主营业务文字分析：同比说明"],
        )
        self.assertTrue(all(
            any(item["article"] == "第5.2.4条" for item in risk["regulation_basis"])
            for risk in risks
        ))

    def test_negative_prior_percentage_uses_semantic_change_types(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "二、主营业务分析\n"
                "财务费用 | 1,827,920.96 | -1,204,119.50 | 251.81%\n"
                "所得税费用 | 116,285.86 | 539,413.43 | -78.44%\n"
                "经营活动产生的现金流量净额 | -24,379,428.64 | -42,492,502.59 | 42.63%\n"
                "投资活动产生的现金流量净额 | -22,237,942.79 | -98,847,347.67 | 77.50%\n"
                "筹资活动产生的现金流量净额 | -10,971,363.00 | -4,214,096.62 | -160.35%\n"
                "现金及现金等价物净增加额 | -57,709,522.81 | -145,580,288.70 | 60.36%"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        by_title = {risk["title"]: risk for risk in risks}
        self.assertNotIn("同比比例计算不一致：所得税费用", by_title)
        for label in (
            "财务费用", "经营活动产生的现金流量净额", "投资活动产生的现金流量净额",
            "筹资活动产生的现金流量净额", "现金及现金等价物净增加额",
        ):
            self.assertIn(f"同比变动展示语义不当：{label}", by_title)
        self.assertEqual(
            by_title["同比变动展示语义不当：财务费用"]["calculation"]["change_type"],
            "negative_to_positive",
        )
        self.assertEqual(
            by_title["同比变动展示语义不当：经营活动产生的现金流量净额"]["calculation"]["change_type"],
            "net_outflow_narrowed",
        )
        self.assertEqual(
            by_title["同比变动展示语义不当：筹资活动产生的现金流量净额"]["calculation"]["change_type"],
            "net_outflow_expanded",
        )

    def test_operating_analysis_is_checked_against_cash_flow_statement_signs(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "二、主营业务分析\n"
                "筹资活动产生的现金流量净额 | 10,971,363.00 | -4,214,096.62 | -160.35%\n"
                "5、合并现金流量表\n"
                "经营活动产生的现金流量净额 | -24,379,428.64 | -42,492,502.59\n"
                "投资活动产生的现金流量净额 | -22,237,942.79 | -98,847,347.67\n"
                "筹资活动产生的现金流量净额 | -10,971,363.00 | -4,214,096.62\n"
                "汇率变动对现金及现金等价物的影响 | -120,788.38 | -26,341.82\n"
                "现金及现金等价物净增加额 | -57,709,522.81 | -145,580,288.70\n"
                "6、母公司现金流量表"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        risk = next(
            item for item in risks
            if item["title"] == "主营业务分析表与合并现金流量表不一致：筹资活动产生的现金流量净额"
        )
        self.assertEqual(risk["level"], "high")
        self.assertIn("本期差额 21,942,726.00", risk["description"])
        self.assertEqual(
            [item["label"] for item in risk["location"]["items"]],
            ["主营业务分析表", "合并现金流量表"],
        )

    def test_cash_flow_components_must_equal_net_cash_change(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "5、合并现金流量表\n"
                "经营活动产生的现金流量净额 | -24,379,428.64 | -42,492,502.59\n"
                "投资活动产生的现金流量净额 | -22,237,942.79 | -98,847,347.67\n"
                "筹资活动产生的现金流量净额 | 10,971,363.00 | -4,214,096.62\n"
                "汇率变动对现金及现金等价物的影响 | -120,788.38 | -26,341.82\n"
                "现金及现金等价物净增加额 | -57,709,522.81 | -145,580,288.70\n"
                "6、母公司现金流量表"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        risk = next(item for item in risks if item["title"] == "现金及现金等价物净增加额勾稽不一致")
        self.assertIn("经营活动净额 + 投资活动净额 + 筹资活动净额", risk["description"])
        self.assertEqual(len(risk["location"]["items"]), 5)

    def test_negative_prior_percentage_is_never_presented_as_normal_yoy(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "二、主营业务分析\n"
                "财务费用 | 1,827,920.96 | -1,204,119.50 | -251.81%\n"
                "经营活动产生的现金流量净额 | -24,379,428.64 | -42,492,502.59 | -42.63%"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        by_title = {risk["title"]: risk for risk in risks}
        expense = by_title["同比变动展示语义不当：财务费用"]
        operating = by_title["同比变动展示语义不当：经营活动产生的现金流量净额"]
        self.assertIn("财务费用由负转正", expense["description"])
        self.assertIn("经营活动产生的现金流量净额净流出收窄", operating["description"])
        self.assertIsNone(expense["calculation"]["display_yoy_rate"])
        self.assertIsNone(operating["calculation"]["display_yoy_rate"])
        self.assertIn("不得使用内部 raw_yoy_rate", expense["recommendation"])

    @staticmethod
    def _narrative_report_text(year: int, risk_text: str, business_text: str | None = None) -> str:
        business = business_text or "公司重点发展人工智能业务，天启产品已实现商业化并产生收入。"
        return (
            f"创业黑马科技集团股份有限公司{year}年半年度报告\n"
            "一、报告期内公司从事的主要业务\n"
            "（一）主要业务概述\n"
            f"{business}\n"
            "二、核心竞争力分析\n"
            "公司以天启大模型和企业服务网络形成差异化人工智能业务竞争力。\n"
            "三、社会责任情况\n"
            "公司通过人工智能服务支持中小企业数字化转型。\n"
            "四、公司面临的风险和应对措施\n"
            f"{risk_text}\n"
            "五、其他事项\n"
        )

    def test_narrative_sections_detect_cross_section_opposite_facts(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            self._narrative_report_text(
                2026,
                "天启产品尚未商业化，尚未产生收入。",
            ).encode("utf-8"),
        )
        sections = extract_narrative_sections(parsed)
        self.assertEqual(
            set(sections),
            {"main_business", "business_overview", "core_competitiveness", "risk_disclosure", "social_responsibility"},
        )
        risks, summary = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        titles = [risk["title"] for risk in risks]
        self.assertIn("文字章节存在相反事实表述：人工智能业务", titles)
        contradiction = next(risk for risk in risks if risk["title"].startswith("文字章节存在相反事实表述"))
        self.assertEqual(contradiction["type"], "定期报告文字一致性")
        self.assertEqual(len(contradiction["location"]["items"]), 2)
        self.assertEqual(summary["narrativeSectionsFound"], 5)
        self.assertGreaterEqual(summary["narrativeInternalConflicts"], 1)

    def test_narrative_review_ignores_explicit_prior_period_contrast(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            self._narrative_report_text(
                2026,
                "人工智能业务商业化仍受客户预算及技术迭代影响，未来存在不确定性。",
                "天启产品本期已实现商业化并产生收入，上年同期尚未商业化且未产生收入。",
            ).encode("utf-8"),
        )
        risks, summary = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        self.assertFalse(any(risk["title"].startswith("文字章节存在相反事实表述") for risk in risks))
        self.assertEqual(summary["narrativeInternalConflicts"], 0)

    def test_social_responsibility_keeps_nested_numbered_subsections_until_next_report_section(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "第四节 环境和社会责任\n"
                "五、社会责任情况\n"
                "一、公司概况和主要举措\n"
                "公司通过人工智能服务支持中小企业数字化转型。\n"
                "二、员工权益保护\n"
                "公司持续完善员工培训和权益保障。\n"
                "第五节 重要事项\n"
                "一、承诺事项\n"
                "本节文字不得进入社会责任章节。\n"
            ).encode("utf-8"),
        )
        social = extract_narrative_sections(parsed)["social_responsibility"]
        self.assertIn("中小企业数字化转型", social.text)
        self.assertIn("员工权益保护", social.text)
        self.assertNotIn("本节文字不得进入社会责任章节", social.text)

    def test_unstructured_annual_risk_prompt_keeps_nested_risks_until_next_report_section(self) -> None:
        parsed = parse_document(
            "2025年年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2025年年度报告\n"
                "第一节 重要提示、目录和释义\n"
                "风险因素提示如下：\n"
                "（1）市场竞争加剧的风险\n"
                "企业服务市场同质化竞争加剧。\n"
                "（2）人工智能服务业务开拓风险\n"
                "人工智能技术和商业化存在不确定性。\n"
                "第二节 公司简介和主要财务指标\n"
                "本节人工智能文字不得进入风险提示章节。\n"
            ).encode("utf-8"),
        )
        risk_section = extract_narrative_sections(parsed)["risk_disclosure"]
        self.assertIn("人工智能服务业务开拓风险", risk_section.text)
        self.assertIn("商业化存在不确定性", risk_section.text)
        self.assertNotIn("本节人工智能文字不得进入风险提示章节", risk_section.text)

    def test_research_investment_conflict_uses_precise_risk_title(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            self._narrative_report_text(
                2026,
                "公司持续加大人工智能研发投入力度。",
                "公司相应减少了人工智能基础性研发投入。",
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        titles = [risk["title"] for risk in risks]
        risk = next(item for item in risks if item["title"] == "文字章节存在相反事实表述：研发投入")
        self.assertIn("文字章节存在相反事实表述：研发投入", titles)
        self.assertNotIn("文字章节存在相反事实表述：人工智能业务", titles)
        self.assertIn("研发费用同比下降", risk["recommendation"])
        self.assertIn("局部增加与研发费用总额下降的关系", risk["recommendation"])

    def test_narrative_history_flags_unexplained_disappearing_risk_with_sources(self) -> None:
        current = parse_document(
            "2026年半年度报告.txt",
            self._narrative_report_text(
                2026,
                "企业服务市场竞争加剧，公司将持续提升服务能力。",
            ).encode("utf-8"),
        )
        root = Path(self.temporary.name)
        rows = []
        for title, date in (("2025年半年度报告", "2025-08-26"), ("2025年年度报告", "2026-04-08")):
            baseline_text = self._narrative_report_text(
                2025,
                "人工智能服务业务开拓存在技术迭代和商业化不确定性风险。",
            ).replace("2025年半年度报告", title)
            path = root / f"{title}.txt"
            path.write_text(baseline_text, encoding="utf-8")
            rows.append({
                "id": title,
                "title": title,
                "publication_date": date,
                "source": "巨潮资讯网",
                "source_url": f"https://example.test/{title}.pdf",
                "file_path": str(path),
            })
        risks, summary = compare_periodic_report(current, "2026年半年度报告.txt", rows)
        risk = next(item for item in risks if item["title"] == "既往重大风险本期未延续披露：人工智能业务")
        self.assertEqual(risk["level"], "needs_review")
        self.assertEqual(len(risk["source_documents"]), 2)
        self.assertTrue(all(source["citation_kind"] == "period_narrative_comparison" for source in risk["source_documents"]))
        self.assertEqual(summary["narrativeHistoricalChanges"], 1)
        self.assertEqual(summary["narrativeSectionsChecked"], 5)

    def test_social_responsibility_zero_debt_claim_is_checked_against_balance_sheet(self) -> None:
        text = self._narrative_report_text(
            2026,
            "企业服务市场竞争加剧，公司将持续提升服务能力。",
        ).replace(
            "公司通过人工智能服务支持中小企业数字化转型。",
            "公司坚持零有息负债和轻资产运营。",
        ) + (
            "1、合并资产负债表\n"
            "短期借款 | 100.00 | 0.00\n"
            "长期借款 | 20.00 | 0.00\n"
            "2、母公司资产负债表\n"
        )
        parsed = parse_document("2026年半年度报告.txt", text.encode("utf-8"))
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        risk = next(item for item in risks if item["title"] == "社会责任表述与财务报表有息借款不一致")
        self.assertEqual(risk["level"], "high")
        self.assertEqual(len(risk["location"]["items"]), 3)

    def test_annotated_docx_comments_every_location_and_materializes_target_altchunk(self) -> None:
        source_document = WordDocument()
        source_document.add_paragraph("2026年半年度报告")
        source_buffer = io.BytesIO()
        source_document.save(source_buffer)
        source = source_buffer.getvalue()

        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        content_type_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(source)) as archive, zipfile.ZipFile(output, "w") as destination:
            document_root = ElementTree.fromstring(archive.read("word/document.xml"))
            body = document_root.find(f".//{{{word_ns}}}body")
            self.assertIsNotNone(body)
            alt_chunk = ElementTree.Element(
                f"{{{word_ns}}}altChunk",
                {f"{{{office_rel_ns}}}id": "rIdXzsAltChunk"},
            )
            body.insert(max(0, len(body) - 1), alt_chunk)
            document_xml = ElementTree.tostring(document_root, encoding="utf-8", xml_declaration=True)

            relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
            ElementTree.SubElement(relationships, f"{{{package_rel_ns}}}Relationship", {
                "Id": "rIdXzsAltChunk",
                "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk",
                "Target": "xzs-altchunk.htm",
            })
            relationship_xml = ElementTree.tostring(relationships, encoding="utf-8", xml_declaration=True)

            content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
            if not any(node.get("Extension") == "htm" for node in content_types):
                ElementTree.SubElement(content_types, f"{{{content_type_ns}}}Default", {
                    "Extension": "htm", "ContentType": "text/html",
                })
            content_type_xml = ElementTree.tostring(content_types, encoding="utf-8", xml_declaration=True)
            for item in archive.infolist():
                replacement = {
                    "word/document.xml": document_xml,
                    "word/_rels/document.xml.rels": relationship_xml,
                    "[Content_Types].xml": content_type_xml,
                }.get(item.filename)
                destination.writestr(item, replacement if replacement is not None else archive.read(item.filename))
            destination.writestr(
                "word/xzs-altchunk.htm",
                "<html><body><p>销售费用文字分析同比下降44.29%</p>"
                "<table><tr><td>项目</td><td>本期</td><td>上期</td><td>同比</td></tr>"
                "<tr><td>销售费用</td><td>50.00</td><td>100.00</td><td>-40.00%</td></tr>"
                "</table></body></html>".encode("utf-8"),
            )
        source_with_altchunk = output.getvalue()
        source_hash = hashlib.sha256(source_with_altchunk).hexdigest()
        risks = [{
            "id": "risk-active",
            "title": "销售费用披露不一致",
            "level": "medium",
            "status": "pending",
            "description": "文字分析与表格列示比例不一致。",
            "recommendation": "核对底稿后统一修改两处披露。",
            "regulation_basis": [{"title": "测试规则", "article": "第一条"}],
            "originalLocations": [
                {"label": "文字分析", "paragraphIndex": 1, "text": "销售费用文字分析同比下降44.29%"},
                {"label": "费用表格", "paragraphIndex": 2, "text": "销售费用 | 50.00 | 100.00 | -40.00%"},
            ],
        }, {
            "id": "risk-false-positive",
            "title": "已标记误报",
            "level": "low",
            "status": "false_positive",
            "description": "不应写入批注。",
            "recommendation": "无。",
            "originalLocations": [{"paragraphIndex": 0, "text": "2026年半年度报告"}],
        }]
        annotated = create_annotated_docx(source_with_altchunk, "半年报.docx", risks)
        self.assertEqual(hashlib.sha256(source_with_altchunk).hexdigest(), source_hash)
        with zipfile.ZipFile(io.BytesIO(annotated)) as archive:
            comments_xml = archive.read("word/comments.xml").decode("utf-8")
            document_xml = archive.read("word/document.xml").decode("utf-8")
            settings_xml = archive.read("word/settings.xml").decode("utf-8")
        annotated_root = ElementTree.fromstring(document_xml)
        self.assertEqual(len(annotated_root.findall(f".//{{{word_ns}}}commentRangeStart")), 2)
        self.assertEqual(len(annotated_root.findall(f".//{{{word_ns}}}commentRangeEnd")), 2)
        self.assertEqual(len(annotated_root.findall(f".//{{{word_ns}}}commentReference")), 2)
        self.assertIn("销售费用披露不一致", comments_xml)
        self.assertIn("核对底稿后统一修改两处披露", comments_xml)
        self.assertNotIn("已标记误报", comments_xml)
        self.assertFalse(annotated_root.findall(f".//{{{word_ns}}}altChunk"))
        self.assertIn("销售费用文字分析同比下降44.29%", document_xml)
        self.assertIn("-40.00%", document_xml)
        self.assertEqual(len(annotated_root.findall(f".//{{{word_ns}}}ins")), 2)
        self.assertIn("trackRevisions", settings_xml)

    def test_annotated_task_document_preserves_original_enforces_owner_and_audits(self) -> None:
        source_document = WordDocument()
        source_document.add_paragraph("创业黑马科技集团股份有限公司2026年半年度报告")
        source_document.add_paragraph("第一节 重要提示、目录和释义")
        source_document.add_paragraph("公司董事会、监事会及董事、监事、高级管理人员保证报告内容真实。")
        source_buffer = io.BytesIO()
        source_document.save(source_buffer)
        source = source_buffer.getvalue()
        task = self.platform.create_task(
            "2026年半年度报告.docx",
            source,
            "批注版权限测试",
            "定期报告一致性审查",
            "创业板定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        self.assertTrue(task["risks"])
        document_row = self.database.fetch_one(
            "SELECT file_path,file_hash FROM documents WHERE id=?", (task["documentId"],),
        )
        source_path = Path(document_row["file_path"])
        before_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
        output_name, annotated = self.platform.get_annotated_task_document(
            task["taskId"], self.admin, "127.0.0.1",
        )
        self.assertTrue(output_name.endswith("_AI合规审阅批注版.docx"))
        self.assertTrue(annotated.startswith(b"PK"))
        self.assertEqual(hashlib.sha256(source_path.read_bytes()).hexdigest(), before_hash)
        self.assertEqual(before_hash, document_row["file_hash"])
        audit = self.database.fetch_one(
            "SELECT action,after_json FROM audit_logs WHERE resource_id=? AND action='annotated_document_generate'",
            (task["taskId"],),
        )
        self.assertIsNotNone(audit)
        self.assertGreater(json.loads(audit["after_json"])["annotated_locations"], 0)

        _token, other_user = self.auth.register(
            "other_user", "OtherUser123!", "其他用户", "127.0.0.2",
        )
        with self.assertRaises(PermissionError):
            self.platform.get_annotated_task_document(task["taskId"], other_user, "127.0.0.2")

    def test_monetary_funds_difference_preserves_statement_and_note_locations(self) -> None:
        blocks = [
            ParsedBlock("2026年半年度报告", 1, "正文", 0, "heading"),
            ParsedBlock("1、合并资产负债表", 10, "1、合并资产负债表", 1, "heading"),
            ParsedBlock("货币资金 | 195,670,788.01 | 254,033,402.63", 10, "1、合并资产负债表", 2, "table"),
            ParsedBlock("2、母公司资产负债表", 12, "2、母公司资产负债表", 3, "heading"),
            ParsedBlock("1、货币资金", 15, "1、货币资金", 4, "heading"),
            ParsedBlock(
                "项目 | 期末余额 | 期初余额\n"
                "库存现金 | 10,000.00 | 0.00\n"
                "银行存款 | 220,679,372.37 | 243,907,385.08\n"
                "其他货币资金 | 6,981,415.64 | 10,088,683.58\n"
                "合计 | 227,670,788.01 | 254,033,402.63",
                15, "1、货币资金", 5, "table",
            ),
            ParsedBlock("2、应收票据", 16, "2、应收票据", 6, "heading"),
        ]
        parsed = ParsedDocument(
            filename="2026年半年度报告.txt",
            file_type="txt",
            page_count=16,
            blocks=blocks,
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        risk = next(item for item in risks if item["title"] == "货币资金附注合计与资产负债表不一致")
        self.assertEqual(len(risk["location"]["items"]), 2)
        self.assertEqual(
            [item["label"] for item in risk["location"]["items"]],
            ["合并资产负债表：货币资金期末余额", "货币资金附注：合计"],
        )

    def test_task_result_serializes_every_original_location(self) -> None:
        result = self.platform.create_task(
            "2026年半年度报告.txt",
            (
                "2026年半年度报告\n"
                "二、主营业务分析\n"
                "销售费用 | 50.00 | 100.00 | -40.00%\n"
                "销售费用为50.00元，同比下降45.00%"
            ).encode("utf-8"),
            "多原文位置序列化测试",
            "定期报告一致性审查",
            "定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        risk = next(
            item for item in result["risks"]
            if item["title"] == "文字分析与表格同比比例不一致：销售费用"
        )
        self.assertEqual(len(risk["originalLocations"]), 2)
        self.assertEqual(risk["originalLocations"], risk["original_locations"])
        self.assertIn("涉及 2 处原文", risk["location"])

    def test_financial_change_is_persisted_but_raw_rate_is_not_exposed_to_frontend(self) -> None:
        result = self.platform.create_task(
            "2026年半年度报告.txt",
            (
                "2026年半年度报告\n"
                "二、主营业务分析\n"
                "财务费用 | 1,827,920.96 | -1,204,119.50 | -251.81%"
            ).encode("utf-8"),
            "财务同比语义持久化测试",
            "定期报告一致性审查",
            "定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        risk = next(
            item for item in result["risks"]
            if item["title"] == "同比变动展示语义不当：财务费用"
        )
        self.assertEqual(risk["calculation"]["display_change_text"], "财务费用由负转正")
        self.assertNotIn("raw_yoy_rate", risk["calculation"])

        row = self.database.fetch_one("SELECT calculation_json FROM risks WHERE id=?", (risk["id"],))
        stored = json.loads(row["calculation_json"])
        self.assertEqual(stored["raw_yoy_rate"], "251.81")
        self.assertIsNone(stored["display_yoy_rate"])

    def test_periodic_comparison_separates_historical_and_internal_differences(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "2026年半年度报告\n"
                "二、主营业务分析\n"
                "销售费用 | 50.00 | 100.00 | -50.00%\n"
                "销售费用为50.00元，同比下降45.00%\n"
            ).encode("utf-8"),
        )
        risks, summary = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        self.assertTrue(risks)
        self.assertEqual(summary["historicalMismatches"], 0)
        self.assertEqual(summary["internalMismatches"], len(risks))

    def test_local_casting_checks_simple_totals_without_external_ai(self) -> None:
        correct_table = (
            "项目 | 期末余额 | 期初余额\n"
            "库存现金 | 10.00 | 20.00\n"
            "银行存款 | 90.00 | 180.00\n"
            "合计 | 100.00 | 200.00"
        )
        parsed = ParsedDocument(
            "财务附注.docx", "DOCX", 8,
            [ParsedBlock(correct_table, 8, "货币资金", 0, "table", 0, len(correct_table))],
        )
        risks, coverage = review_casting(parsed)
        self.assertEqual(risks, [])
        self.assertEqual(coverage["checksPerformed"], 2)
        self.assertFalse(coverage["externalAiUsed"])

        incorrect_table = correct_table.replace("合计 | 100.00 | 200.00", "合计 | 101.00 | 200.00")
        parsed.blocks[0] = ParsedBlock(
            incorrect_table, 8, "货币资金", 0, "table", 0, len(incorrect_table),
        )
        risks, coverage = review_casting(parsed)
        self.assertEqual(len(risks), 1)
        self.assertEqual(risks[0]["title"], "附注明细合计不一致：货币资金")
        self.assertIn("差额 1.00", risks[0]["description"])
        self.assertEqual(risks[0]["location"]["page"], 8)
        self.assertFalse(coverage["externalAiUsed"])

    def test_local_casting_checks_carrying_value_and_movement_closure(self) -> None:
        carrying_table = (
            "项目 | 账面余额 | 坏账准备 | 账面价值\n"
            "应收账款 | 1,000.00 | 100.00 | 900.00\n"
            "其他应收款 | 500.00 | 50.00 | 460.00"
        )
        movement_table = (
            "项目 | 期初余额 | 本期增加 | | 本期减少 | | 期末余额\n"
            "| | 现金变动 | 非现金变动 | 现金变动 | 非现金变动 |\n"
            "坏账准备 | 100.00 | 5.00 | 20.00 | 10.00 | 3.00 | 112.00\n"
            "累计摊销 | 200.00 | 10.00 | 20.00 | 5.00 | 3.00 | 220.00"
        )
        parsed = ParsedDocument(
            "财务附注.docx", "DOCX", 20,
            [
                ParsedBlock(carrying_table, 15, "应收款项", 0, "table", 0, len(carrying_table)),
                ParsedBlock(movement_table, 20, "准备变动", 1, "table", 0, len(movement_table)),
            ],
        )
        risks, coverage = review_casting(parsed)
        self.assertEqual(coverage["checksPerformed"], 4)
        self.assertEqual(len(risks), 2)
        self.assertEqual(
            {risk["title"] for risk in risks},
            {"账面价值计算不一致：其他应收款", "变动表不闭合：累计摊销"},
        )
        self.assertTrue(all(risk["type"] == "财务附注 Casting" for risk in risks))
        self.assertTrue(all(risk["location"]["items"] for risk in risks))
        self.assertFalse(coverage["externalAiUsed"])

    def test_local_casting_recovers_legacy_merged_movement_header_without_false_positive(self) -> None:
        movement_table = (
            "项目 | 期初余额 | 本期增加 | 本期减少 | 期末余额\n"
            "| | 现金变动 | 非现金变动 | 现金变动 | 非现金变动 |\n"
            "租赁负债（一年内到期的非流动负债） | 6,543,922.25 | | 103,160.45 | "
            "1,867,353.26 | | 4,779,729.44\n"
            "合计 | 6,543,922.25 | | 103,160.45 | 1,867,353.26 | | 4,779,729.44"
        )
        parsed = ParsedDocument(
            "2026年半年度报告.docx", "DOCX", 14,
            [ParsedBlock(movement_table, 14, "分类列示筹资活动产生的各项负债变动情况", 0, "table")],
        )
        risks, coverage = review_casting(parsed)
        self.assertFalse(any(risk["title"].startswith("变动表不闭合") for risk in risks))
        self.assertGreaterEqual(coverage["checksPerformed"], 2)

        incorrect_table = movement_table.replace(
            "租赁负债（一年内到期的非流动负债） | 6,543,922.25 | | 103,160.45 | "
            "1,867,353.26 | | 4,779,729.44",
            "租赁负债（一年内到期的非流动负债） | 6,543,922.25 | | 103,160.45 | "
            "1,867,353.26 | | 4,700,000.00",
        )
        parsed.blocks[0] = ParsedBlock(
            incorrect_table, 14, "分类列示筹资活动产生的各项负债变动情况", 0, "table",
        )
        risks, _coverage = review_casting(parsed)
        self.assertIn(
            "变动表不闭合：租赁负债（一年内到期的非流动负债）",
            {risk["title"] for risk in risks},
        )

    def test_local_casting_applies_only_explicit_subtractive_total_semantics(self) -> None:
        finance_costs = (
            "项目 | 本期发生额 | 上期发生额\n"
            "利息支出 | 2,310,329.59 | 208,218.75\n"
            "减：利息收入 | 596,780.85 | 1,573,427.20\n"
            "汇兑损益 | -9,570.21 | 10,402.02\n"
            "其他 | 123,942.43 | 150,686.93\n"
            "合计 | 1,827,920.96 | -1,204,119.50"
        )
        non_recurring = (
            "项目 | 金额 | 说明\n"
            "非流动性资产处置损益 | 3,664.22 |\n"
            "政府补助 | 2,105,806.00 |\n"
            "委托投资损益 | 388,769.69 |\n"
            "其他营业外收入和支出 | -22,173.92 |\n"
            "减：所得税影响额 | 160,588.98 |\n"
            "少数股东权益影响额（税后） | -1,320.14 |\n"
            "合计 | 2,316,797.15 |"
        )
        parsed = ParsedDocument(
            "2026年半年度报告.docx", "DOCX", 20,
            [
                ParsedBlock(finance_costs, 15, "38、财务费用", 0, "table", 0, len(finance_costs)),
                ParsedBlock(non_recurring, 20, "非经常性损益", 1, "table", 0, len(non_recurring)),
            ],
        )
        risks, coverage = review_casting(parsed)
        self.assertEqual(risks, [])
        self.assertEqual(coverage["checksPerformed"], 3)

    def test_periodic_task_returns_local_casting_coverage(self) -> None:
        document = WordDocument()
        document.add_heading("2026年半年度报告", level=1)
        document.add_heading("七、合并财务报表项目注释", level=2)
        table = document.add_table(rows=4, cols=3)
        for row_index, values in enumerate((
            ("项目", "期末余额", "期初余额"),
            ("库存现金", "10.00", "20.00"),
            ("银行存款", "90.00", "180.00"),
            ("合计", "101.00", "200.00"),
        )):
            for column_index, value in enumerate(values):
                table.cell(row_index, column_index).text = value
        buffer = io.BytesIO()
        document.save(buffer)
        task = self.platform.create_task(
            "2026年半年度报告.docx",
            buffer.getvalue(),
            "Casting任务接口测试",
            "定期报告一致性审查",
            "创业板定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        self.assertFalse(task["castingReview"]["externalAiUsed"])
        self.assertGreater(task["castingReview"]["checksPerformed"], 0)
        self.assertGreater(task["castingReview"]["mismatches"], 0)
        self.assertIn("财务附注 Casting", {risk["type"] for risk in task["risks"]})

    def test_periodic_comparison_uses_populated_duplicate_statement_row(self) -> None:
        current = parse_document(
            "2026年半年度报告.txt",
            (
                "2026年半年度报告\n"
                "3、合并利润表\n"
                "利息收入 | 0.00 | 0.00\n"
                "财务费用 | 100.00 | 200.00\n"
                "利息收入 | 596,780.85 | 1,573,427.20\n"
                "4、母公司利润表\n"
            ).encode("utf-8"),
        )
        baseline_path = Path(self.temporary.name) / "2025年半年度报告.txt"
        baseline_path.write_text(
            "2025年半年度报告\n3、合并利润表\n利息收入 | 1,573,427.20 | 3,950,450.73\n4、母公司利润表\n",
            encoding="utf-8",
        )
        risks, _ = compare_periodic_report(current, "2026年半年度报告.txt", [{
            "id": "half", "title": "2025年半年度报告", "publication_date": "2025-08-26",
            "source": "巨潮资讯网", "source_url": "https://example.test/half.pdf",
            "file_path": str(baseline_path),
        }])
        self.assertNotIn("跨期比较数不一致：利息收入", [risk["title"] for risk in risks])

    def test_missing_monetary_note_table_does_not_reuse_receivables_total(self) -> None:
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "2026年半年度报告\n"
                "1、合并资产负债表\n"
                "货币资金 | 195,670,788.01 | 254,033,402.63\n"
                "2、母公司资产负债表\n"
                "七、合并财务报表项目注释\n"
                "1、货币资金\n"
                "2、应收票据\n"
                "3、应收账款\n"
                "（2）按坏账计提方法分类披露\n"
                "合计 | 64,210,280.10 | 59,114,841.35\n"
            ).encode("utf-8"),
        )
        risks, _ = compare_periodic_report(parsed, "2026年半年度报告.txt", [])
        self.assertNotIn("货币资金附注合计与资产负债表不一致", [risk["title"] for risk in risks])

    def test_auto_review_configuration_is_saved_and_manual_selection_overrides_it(self) -> None:
        draft = "关于关联交易的公告\n本事项构成关联交易，关联董事已回避表决。".encode("utf-8")
        automatic = self.platform.create_task(
            "关联交易公告.txt", draft, "自动分类测试", "auto", "auto", self.admin, "127.0.0.1",
        )
        self.assertEqual(automatic["taskType"], "关联交易专项审查")
        self.assertEqual(automatic["rulePack"], "关联交易专项规则包")
        self.assertEqual(automatic["classification"]["method"], "automatic")
        self.assertEqual(automatic["classification"]["taskTypeSource"], "automatic")

        manual = self.platform.create_task(
            "关联交易公告.txt", draft, "人工覆盖测试", "重大事项披露审查",
            "上市公司信披基础包", self.admin, "127.0.0.1",
        )
        self.assertEqual(manual["taskType"], "重大事项披露审查")
        self.assertEqual(manual["rulePack"], "上市公司信披基础包")
        self.assertEqual(manual["classification"]["method"], "manual")
        self.assertEqual(manual["classification"]["detectedTaskType"], "关联交易专项审查")

    def test_large_table_chunking_preserves_text_without_duplicate_tail_chunk(self) -> None:
        rows = [f"第{index}行 | 交易金额{index * 100}万元 | 审批状态待确认" for index in range(1, 181)]
        table_text = "\n".join(rows)
        parsed = ParsedDocument(
            "大表格.docx",
            "DOCX",
            1,
            [ParsedBlock(table_text, 1, "交易明细", 0, "table", 0, len(table_text))],
        )
        chunks = SecuritiesChunker(max_chars=900, overlap_chars=120).chunk(
            "large-table", parsed, "user_documents", {"file_name": "大表格.docx"}
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk.content) <= 900 for chunk in chunks))
        self.assertEqual(len({chunk.content_hash for chunk in chunks}), len(chunks))
        self.assertTrue(all(chunk.metadata["contains_table"] for chunk in chunks))
        self.assertIn(rows[0], chunks[0].content)
        self.assertIn(rows[-1], chunks[-1].content)

    def test_embedding_uses_body_only_and_keeps_operational_fields_as_metadata(self) -> None:
        body = "第一条 公司应当按照现行规则履行审议和信息披露义务。"
        excluded_markers = {
            "stock_code": "STOCK-CODE-METADATA-ONLY",
            "announcement_date": "ANNOUNCEMENT-DATE-METADATA-ONLY",
            "user_id": "USER-ID-METADATA-ONLY",
            "file_status": "FILE-STATUS-METADATA-ONLY",
            "risk_level": "RISK-LEVEL-METADATA-ONLY",
            "transaction_amount": "TRANSACTION-AMOUNT-METADATA-ONLY",
            "page_label": "PAGE-METADATA-ONLY",
        }
        parsed = ParsedDocument(
            "规则正文.txt", "TXT", 7,
            [ParsedBlock(body, 7, "第一条", 0, "paragraph", 0, len(body))],
        )
        chunks = SecuritiesChunker().chunk(
            "metadata-boundary", parsed, "regulations", excluded_markers,
        )
        anchor = self.platform.import_knowledge(
            "边界测试锚点.txt", "锚点正文".encode("utf-8"),
            {"document_type": "regulations", "title": "边界测试锚点", "source": "测试"},
            self.admin, "127.0.0.1",
        )
        captured: list[str] = []
        original_embed_batch = self.platform.embedding_provider.embed_batch

        def capture_embedding(texts: object) -> list[list[float]]:
            values = [str(value) for value in texts]  # type: ignore[union-attr]
            captured.extend(values)
            return original_embed_batch(values)

        self.platform.embedding_provider.embed_batch = capture_embedding  # type: ignore[method-assign]
        try:
            self.platform._index_chunks(
                chunks, knowledge_document_id=anchor["document"]["id"],
            )
        finally:
            self.platform.embedding_provider.embed_batch = original_embed_batch  # type: ignore[method-assign]

        self.assertEqual(captured, [body])
        self.assertEqual(chunks[0].embedding_text, body)
        for value in excluded_markers.values():
            self.assertNotIn(value, "\n".join(captured))
            self.assertIn(value, chunks[0].metadata.values())

    def test_all_supported_knowledge_types_produce_body_embeddings(self) -> None:
        supported = (
            "company_profile", "regulations", "internal_rules",
            "announcements", "cases", "drafting_materials", "user_documents",
        )
        for document_type in supported:
            with self.subTest(document_type=document_type):
                body = f"{document_type} 正文检索内容"
                parsed = ParsedDocument(
                    f"{document_type}.txt", "TXT", 1,
                    [ParsedBlock(body, 1, "正文", 0, "paragraph", 0, len(body))],
                )
                chunks = SecuritiesChunker().chunk(
                    f"document-{document_type}", parsed, document_type,
                    {"document_name": f"{document_type} 元数据标题"},
                )
                self.assertEqual([chunk.embedding_text for chunk in chunks], [body])

    def test_docx_parser_preserves_wingdings_checked_boxes_in_paragraphs_and_tables(self) -> None:
        document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body>
            <w:p>
              <w:r><w:t xml:space="preserve">□适用 </w:t></w:r>
              <w:r><w:sym w:font="Wingdings 2" w:char="F052"/></w:r>
              <w:r><w:t>不适用</w:t></w:r>
            </w:p>
            <w:tbl><w:tr><w:tc><w:p>
              <w:r><w:sym w:font="Wingdings 2" w:char="F052"/></w:r>
              <w:r><w:t>是 □否</w:t></w:r>
            </w:p></w:tc></w:tr></w:tbl>
          </w:body>
        </w:document>"""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", document_xml)
        parsed = parse_document("复选框测试.docx", buffer.getvalue())
        self.assertEqual(parsed.paragraphs, ["□适用 ☑不适用", "☑是 □否"])
        self.assertEqual(parsed.blocks[1].block_type, "table")

    def test_docx_parser_expands_grid_span_in_multirow_financial_header(self) -> None:
        document = WordDocument()
        table = document.add_table(rows=3, cols=7)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "期初余额"
        table.cell(0, 2).merge(table.cell(0, 3)).text = "本期增加"
        table.cell(0, 4).merge(table.cell(0, 5)).text = "本期减少"
        table.cell(0, 6).text = "期末余额"
        for column_index, value in enumerate(("", "", "现金变动", "非现金变动", "现金变动", "非现金变动", "")):
            table.cell(1, column_index).text = value
        for column_index, value in enumerate((
            "租赁负债", "6,543,922.25", "", "103,160.45", "1,867,353.26", "", "4,779,729.44",
        )):
            table.cell(2, column_index).text = value
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse_document("合并表头.docx", buffer.getvalue())
        table_block = next(block for block in parsed.blocks if block.block_type == "table")
        header_cells = [cell.strip() for cell in table_block.text.splitlines()[0].split("|")]
        self.assertEqual(
            header_cells,
            ["项目", "期初余额", "本期增加", "", "本期减少", "", "期末余额"],
        )
        risks, coverage = review_casting(parsed)
        self.assertFalse(any(risk["title"].startswith("变动表不闭合") for risk in risks))
        self.assertEqual(coverage["checksPerformed"], 1)

    def test_docx_parser_reads_tables_nested_in_content_controls(self) -> None:
        document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body>
            <w:p><w:r><w:t>七、合并财务报表项目注释</w:t></w:r></w:p>
            <w:sdt><w:sdtContent>
              <w:p><w:r><w:t>1、货币资金</w:t></w:r></w:p>
              <w:tbl>
                <w:tr><w:tc><w:p><w:r><w:t>项目</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>期末余额</w:t></w:r></w:p></w:tc></w:tr>
                <w:tr><w:tc><w:p><w:r><w:t>银行存款</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>220,679,372.37</w:t></w:r></w:p></w:tc></w:tr>
              </w:tbl>
            </w:sdtContent></w:sdt>
          </w:body>
        </w:document>"""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", document_xml)
        parsed = parse_document("嵌套表格.docx", buffer.getvalue())
        self.assertEqual([block.block_type for block in parsed.blocks], ["heading", "heading", "table"])
        self.assertIn("银行存款 | 220,679,372.37", parsed.blocks[-1].text)

    def test_docx_parser_reads_alt_chunk_html_in_document_order(self) -> None:
        document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
        <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                    xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
          <w:body>
            <w:p><w:r><w:t>1、货币资金</w:t></w:r></w:p>
            <w:altChunk r:id="money-table" />
            <w:p><w:r><w:t>2、应收票据</w:t></w:r></w:p>
          </w:body>
        </w:document>"""
        relationships_xml = """<?xml version="1.0" encoding="UTF-8"?>
        <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
          <Relationship Id="money-table"
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/aFChunk"
            Target="afchunk3a.htm" />
        </Relationships>"""
        html = """<html><body><p>单位：元</p><table>
          <tr><td>项目</td><td>期末账面余额</td><td>期初账面余额</td></tr>
          <tr><td>银行存款</td><td>220,679,372.37</td><td>243,907,385.08</td></tr>
          <tr><td colspan="1">合计</td><td>227,670,788.01</td><td>254,033,402.63</td></tr>
        </table></body></html>"""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("word/document.xml", document_xml)
            archive.writestr("word/_rels/document.xml.rels", relationships_xml)
            archive.writestr("word/afchunk3a.htm", html.encode("utf-8"))
        parsed = parse_document("内嵌HTML表格.docx", buffer.getvalue())
        self.assertEqual([block.block_type for block in parsed.blocks], ["heading", "paragraph", "table", "heading"])
        self.assertIn("银行存款 | 220,679,372.37 | 243,907,385.08", parsed.blocks[2].text)
        self.assertIn("合计 | 227,670,788.01 | 254,033,402.63", parsed.blocks[2].text)
        self.assertEqual(parsed.blocks[3].text, "2、应收票据")

    def test_legacy_doc_parser_uses_local_ifilter_and_preserves_pages(self) -> None:
        extracted = "第一条 上市公司应当依法披露。\f第二条 本制度自发布之日起施行。"
        with patch("backend.parsers.extract_legacy_word_text", return_value=extracted):
            parsed = parse_document("历史制度.doc", b"\xd0\xcf\x11\xe0legacy")
        self.assertEqual(parsed.file_type, "DOC")
        self.assertEqual(parsed.page_count, 2)
        self.assertEqual([block.page for block in parsed.blocks], [1, 2])
        self.assertIn("第一条", parsed.blocks[0].text)

    def test_background_analysis_job_reports_progress_and_is_account_scoped(self) -> None:
        _, other_user = self.auth.register("job_viewer", "JobViewer123!", "其他用户", "127.0.0.1")
        content = (
            "关于对外投资的公告\n一、交易概述\n本次投资金额为1000万元。\n"
            "相关审批与披露义务应以当前有效规则为准。"
        ).encode("utf-8")
        created = self.platform.create_analysis_job(
            "后台审查.txt", content, "后台审查", "重大事项披露审查", "基础包", self.admin, "127.0.0.1"
        )
        self.assertIn(created["status"], {"queued", "processing"})
        with self.assertRaises(PermissionError):
            self.platform.get_analysis_job(created["jobId"], other_user)

        deadline = time.monotonic() + 10
        job = created
        while job["status"] not in {"completed", "failed"} and time.monotonic() < deadline:
            time.sleep(0.05)
            job = self.platform.get_analysis_job(created["jobId"], self.admin)
        self.assertEqual(job["status"], "completed", job.get("error"))
        self.assertEqual(job["stage"], "completed")
        self.assertEqual(job["progress"], 100)
        self.assertTrue(job["taskId"])
        self.assertEqual(self.platform.get_task_result(job["taskId"], user=self.admin)["taskId"], job["taskId"])

    def test_failed_background_analysis_removes_partial_task_and_hides_internal_error(self) -> None:
        content = "测试公告\n本次合作确保公司未来实现预期收益。".encode("utf-8")
        with patch.object(self.platform, "_insert_risk", side_effect=RuntimeError("UNIQUE constraint failed: risks.id")):
            created = self.platform.create_analysis_job(
                "失败清理测试.txt", content, "失败清理测试", "临时公告通用审查", "基础包",
                self.admin, "127.0.0.1",
            )
            deadline = time.monotonic() + 10
            job = created
            while job["status"] not in {"completed", "failed"} and time.monotonic() < deadline:
                time.sleep(0.05)
                job = self.platform.get_analysis_job(created["jobId"], self.admin)
        self.assertEqual(job["status"], "failed")
        self.assertNotIn("constraint", job["error"].lower())
        self.assertIsNone(job["taskId"])
        self.assertFalse(any(task["name"] == "失败清理测试" for task in self.platform.list_tasks(self.admin)))
        self.assertIsNone(self.database.fetch_one(
            "SELECT id FROM analysis_tasks WHERE name='失败清理测试' AND owner_user_id=?",
            (self.admin["id"],),
        ))

    def test_wording_rule_ignores_governance_boilerplate_but_flags_future_guarantee(self) -> None:
        engine = DeterministicRuleEngine()
        ordinary = parse_document(
            "ordinary.txt",
            "董事积极履行职责，确保决策科学、及时、高效，有效维护公司和股东利益。".encode("utf-8"),
        )
        self.assertFalse(any(risk["type"] == "表述风险" for risk in engine.analyze(ordinary, "年度报告审查")))
        promise = parse_document(
            "promise.txt",
            "本次合作确保公司未来持续稳定增长，并保证实现预期收益。".encode("utf-8"),
        )
        self.assertTrue(any(risk["type"] == "表述风险" for risk in engine.analyze(promise, "临时公告审查")))

    def test_periodic_report_rule_pack_checks_approval_audit_statements_and_deadline(self) -> None:
        engine = DeterministicRuleEngine()
        incomplete = parse_document(
            "2025年年度报告.txt",
            "创业黑马科技集团股份有限公司2025年年度报告\n报告期内公司经营情况如下。".encode("utf-8"),
        )
        risks = engine.analyze(incomplete, "定期报告一致性审查")
        titles = {risk["title"] for risk in risks}
        self.assertIn("定期报告审批程序待核验", titles)
        approval = next(risk for risk in risks if risk["title"] == "定期报告审批程序待核验")
        self.assertIn("审计委员会审核", approval["description"])
        self.assertIn("董事会审议通过", approval["description"])
        self.assertEqual(len(approval["regulation_basis"]), 1)
        self.assertEqual(len(approval["location"]["items"]), 1)
        self.assertEqual(approval["location"]["items"][0]["label"], "审批程序共同对应位置")
        self.assertIn("董高书面确认意见待核验", titles)
        self.assertIn("年度报告审计文件缺失", titles)
        self.assertIn("年度财务报表完整性待核验", titles)
        self.assertIn("定期报告披露时限与预约日期待核验", titles)
        deadline = next(risk for risk in risks if risk["title"] == "定期报告披露时限与预约日期待核验")
        self.assertEqual(deadline["regulation_basis"][0]["article"], "第6.1.2条、第6.1.3条")
        self.assertTrue(deadline["regulation_basis"][0]["source_url"].startswith("https://www.szse.cn/"))

        complete = parse_document(
            "2025年年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2025年年度报告\n"
                "审计委员会审核年度报告财务信息并经全体成员过半数同意。\n"
                "本年度报告已经董事会审议通过。\n"
                "全体董事及高级管理人员签署书面确认意见，保证报告真实、准确、完整。\n"
                "审计报告 会计师事务所 注册会计师 审计意见\n"
                "资产负债表\n利润表\n现金流量表\n所有者权益变动表"
            ).encode("utf-8"),
        )
        complete_titles = {risk["title"] for risk in engine.analyze(complete, "定期报告一致性审查")}
        self.assertNotIn("定期报告审批程序待核验", complete_titles)
        self.assertNotIn("董高书面确认意见待核验", complete_titles)
        self.assertNotIn("年度报告审计文件缺失", complete_titles)
        self.assertNotIn("年度财务报表完整性待核验", complete_titles)
        self.assertIn("定期报告披露时限与预约日期待核验", complete_titles)

    def test_periodic_report_rag_cites_clickable_szse_rule(self) -> None:
        regulation = (
            "深圳证券交易所创业板股票上市规则（2026年修订）\n"
            "第六章 定期报告\n"
            "第6.1.4条 定期报告中的财务信息应当经审计委员会审核，由审计委员会全体成员过半数同意后提交董事会审议。"
        ).encode("utf-8")
        source_url = "https://www.szse.cn/lawrules/rule/stock/supervision/chinext/t20260424_620189.html"
        imported = self.platform.import_knowledge(
            "创业板股票上市规则2026.txt",
            regulation,
            {
                "document_type": "regulations",
                "title": "深圳证券交易所创业板股票上市规则（2026年修订）",
                "source": "深圳证券交易所",
                "source_url": source_url,
                "version": "2026年修订",
            },
            self.admin,
            "127.0.0.1",
        )
        self.assertEqual(imported["status"], "success")
        result = self.platform.create_task(
            "2025年年度报告.txt",
            "创业黑马科技集团股份有限公司2025年年度报告\n报告期经营情况。".encode("utf-8"),
            "2025年年度报告审查",
            "定期报告一致性审查",
            "创业板定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        approval_risk = next(risk for risk in result["risks"] if risk["title"] == "定期报告审批程序待核验")
        official_sources = [
            source for source in approval_risk["source_documents"]
            if source.get("source_url") == source_url
        ]
        self.assertTrue(official_sources)
        self.assertTrue(any(
            "6.1.4" in source.get("content", "") or "6.1.4" in source.get("section", "")
            for source in official_sources
        ))
        self.assertGreater(result["ragCoverage"]["ruleDocumentsScanned"], 0)
        self.assertGreater(result["ragCoverage"]["ruleChunksScanned"], 0)
        self.assertFalse(any(
            source.get("document_id") == result["documentId"]
            for risk in result["risks"] for source in risk["source_documents"]
        ))

    def test_periodic_report_aggregates_unselected_template_choices(self) -> None:
        engine = DeterministicRuleEngine()
        parsed = parse_document(
            "2026年半年度报告.txt",
            (
                "创业黑马科技集团股份有限公司2026年半年度报告\n"
                "优先股相关情况\n□适用 □不适用\n"
                "半年度报告是否经过审计\n□是 □否"
            ).encode("utf-8"),
        )
        risks = [
            risk for risk in engine.analyze(parsed, "定期报告一致性审查")
            if risk["type"] == "模板选项完整性"
        ]
        self.assertEqual(len(risks), 1)
        self.assertEqual(risks[0]["level"], "needs_review")
        self.assertIn("共2处", risks[0]["title"])
        self.assertIn("优先股相关情况", risks[0]["description"])
        self.assertIn("半年度报告是否经过审计", risks[0]["description"])
        self.assertEqual(len(risks[0]["location"]["items"]), 2)
        self.assertTrue(all(risk["regulation_basis"] for risk in risks))
        self.assertTrue(all("仅选一项" in risk["recommendation"] for risk in risks))

    def test_periodic_report_does_not_show_page_only_accounting_guidance(self) -> None:
        source_url = "https://www.szse.cn/lawrules/csrcrules/guide/P020240802361881526199.pdf"
        imported = self.platform.import_knowledge(
            "监管规则适用指引——会计类第4号.txt",
            (
                "监管规则适用指引——会计类第4号\n"
                "上市公司应当按照企业会计准则和证券监管要求，核对确认、计量、列报和披露事项。"
            ).encode("utf-8"),
            {
                "document_type": "regulations",
                "title": "监管规则适用指引——会计类第4号",
                "source": "中国证券监督管理委员会",
                "source_url": source_url,
                "version": "官方现行收录版本",
            },
            self.admin,
            "127.0.0.1",
        )
        self.assertEqual(imported["status"], "success")
        result = self.platform.create_task(
            "2025年年度报告.txt",
            "创业黑马科技集团股份有限公司2025年年度报告\n报告期经营情况。".encode("utf-8"),
            "2025年年度报告审查",
            "定期报告一致性审查",
            "创业板定期报告规则包",
            self.admin,
            "127.0.0.1",
        )
        financial_risk = next(risk for risk in result["risks"] if risk["title"] == "年度财务报表完整性待核验")
        accounting_sources = [
            source for source in financial_risk["source_documents"]
            if source.get("source_url") == source_url
        ]
        self.assertFalse(accounting_sources)
        self.assertTrue(any(
            source.get("citation_kind") == "deterministic_rule"
            and source.get("article")
            for source in financial_risk["source_documents"]
        ))

    def test_registration_validation_and_duplicate_username(self) -> None:
        token, user = self.auth.register("reviewer01", "PrivatePass123!", "审阅用户", "127.0.0.1")
        self.assertTrue(token)
        self.assertEqual(user["role"], "user")
        self.assertEqual(self.auth.authenticate(token)["id"], user["id"])
        with self.assertRaisesRegex(ValueError, "已被注册"):
            self.auth.register("REVIEWER01", "PrivatePass123!", "重复用户", "127.0.0.1")
        with self.assertRaisesRegex(ValueError, "8-128"):
            self.auth.register("shortpass", "1234567", "短密码", "127.0.0.1")

    def test_only_admin_can_manage_shared_knowledge(self) -> None:
        self.assertIn("knowledge", ROLE_PERMISSIONS["admin"])
        self.assertIn("users", ROLE_PERMISSIONS["admin"])
        self.assertNotIn("knowledge", ROLE_PERMISSIONS["user"])
        self.assertNotIn("users", ROLE_PERMISSIONS["user"])
        self.assertNotIn("knowledge", ROLE_PERMISSIONS["securities"])
        _, regular = self.auth.register("knowledge_reader", "ReaderPass123!", "资料使用者", "127.0.0.1")
        imported = self.platform.import_knowledge(
            "共享制度.txt",
            "共享制度\n第一条 所有审查任务均应核对有效规则。".encode("utf-8"),
            {"document_type": "internal_rules", "title": "共享制度", "source": "管理员测试"},
            self.admin,
            "127.0.0.1",
        )
        self.assertEqual(imported["document"]["scope"], "shared")
        regular_documents = self.platform.list_knowledge(regular)
        self.assertEqual([item["title"] for item in regular_documents], ["共享制度"])
        self.assertFalse(regular_documents[0]["can_manage"])
        opened_name, opened_content = self.platform.get_knowledge_document(
            imported["document"]["id"], self.admin, "127.0.0.1",
        )
        self.assertEqual(opened_name, "共享制度.txt")
        self.assertEqual(opened_content, "共享制度\n第一条 所有审查任务均应核对有效规则。".encode("utf-8"))
        reader_name, reader_content = self.platform.get_knowledge_document(
            imported["document"]["id"], regular, "127.0.0.1",
        )
        self.assertEqual(reader_name, "共享制度.txt")
        self.assertEqual(reader_content, opened_content)
        self.database.execute(
            "UPDATE knowledge_documents SET owner_user_id=? WHERE id=?",
            (self.admin["id"], imported["document"]["id"]),
        )
        with self.assertRaisesRegex(PermissionError, "其他账号的私有资料"):
            self.platform.get_knowledge_document(
                imported["document"]["id"], regular, "127.0.0.1",
            )
        self.database.execute(
            "UPDATE knowledge_documents SET owner_user_id=NULL,status='inactive' WHERE id=?",
            (imported["document"]["id"],),
        )
        with self.assertRaisesRegex(PermissionError, "未启用"):
            self.platform.get_knowledge_document(
                imported["document"]["id"], regular, "127.0.0.1",
            )
        self.database.execute(
            "UPDATE knowledge_documents SET status='active' WHERE id=?",
            (imported["document"]["id"],),
        )
        view_audit = self.database.fetch_one(
            "SELECT action FROM audit_logs WHERE action='knowledge_document_view' ORDER BY created_at DESC LIMIT 1"
        )
        self.assertIsNotNone(view_audit)
        with self.assertRaisesRegex(PermissionError, "只有管理员"):
            self.platform.import_knowledge(
                "越权资料.txt",
                "普通账号不能导入资料。".encode("utf-8"),
                {"document_type": "internal_rules", "title": "越权资料", "source": "普通账号"},
                regular,
                "127.0.0.1",
            )
        with self.assertRaisesRegex(PermissionError, "只有管理员"):
            self.platform.rebuild_knowledge(imported["document"]["id"], regular, "127.0.0.1")
        with self.assertRaisesRegex(PermissionError, "只有管理员"):
            self.platform.delete_knowledge(imported["document"]["id"], regular, "127.0.0.1")
        self.database.migrate()
        self.assertEqual([item["title"] for item in self.platform.list_knowledge(regular)], ["共享制度"])

    def test_admin_custom_knowledge_category_is_shared_but_read_only_for_regular_users(self) -> None:
        _, regular = self.auth.register("category_reader", "ReaderPass123!", "分类阅读者", "127.0.0.1")
        category = self.platform.create_knowledge_category("法务资料", "合同和法律意见", self.admin, "127.0.0.1")
        self.assertEqual(category["name"], "法务资料")
        imported = self.platform.import_knowledge(
            "法律意见.txt",
            "法律意见书\n第一条 合同审查应保留书面依据。".encode("utf-8"),
            {"document_type": "user_documents", "title": "法律意见书", "source": "法务部", "category_id": category["id"]},
            self.admin,
            "127.0.0.1",
        )
        visible = self.platform.list_knowledge(regular, category_id=category["id"])
        self.assertEqual([item["title"] for item in visible], ["法律意见书"])
        self.assertEqual(visible[0]["category_name"], "法务资料")
        self.assertFalse(visible[0]["can_manage"])
        with self.assertRaisesRegex(PermissionError, "只有管理员"):
            self.platform.create_knowledge_category("诉讼资料", "", regular, "127.0.0.1")
        self.platform.delete_knowledge_category(category["id"], self.admin, "127.0.0.1")
        self.assertEqual(self.platform.list_knowledge_categories(regular), [])

    def test_workspace_chat_is_local_and_account_isolated(self) -> None:
        _, regular = self.auth.register("chat_reader", "ReaderPass123!", "独立对话用户", "127.0.0.1")
        self.platform.import_knowledge(
            "聊天依据.txt",
            "关联交易制度\n第一条 关联交易应当履行审议和披露程序。".encode("utf-8"),
            {"document_type": "internal_rules", "title": "关联交易制度", "source": "证券部"},
            self.admin,
            "127.0.0.1",
        )
        response = self.platform.workspace_chat("关联交易审议和披露程序", None, self.admin, "127.0.0.1")
        self.assertEqual(response["mode"], "local_rag")
        self.assertEqual(response["provider"], "local-rag-retrieval")
        self.assertTrue(response["source_documents"])
        self.assertEqual(len(self.platform.list_workspace_conversations(self.admin)), 1)
        self.assertEqual(self.platform.list_workspace_conversations(regular), [])
        with self.assertRaisesRegex(ValueError, "独立对话不存在"):
            self.platform.get_workspace_messages(response["conversation_id"], regular)

    def test_current_review_retrieval_excludes_repealed_regulations(self) -> None:
        current = self.platform.import_knowledge(
            "现行董事会规则.txt",
            "董事会决议公告应披露会议通知、召开方式和逐项表决结果。".encode("utf-8"),
            {
                "document_type": "regulations",
                "title": "现行董事会规则",
                "source": "深圳证券交易所",
                "version": "现行有效",
            },
            self.admin,
            "127.0.0.1",
        )
        repealed = self.platform.import_knowledge(
            "已废止董事会规则.txt",
            "董事会决议公告应披露会议通知、召开方式和逐项表决结果。旧版本。".encode("utf-8"),
            {
                "document_type": "regulations",
                "title": "已废止董事会规则",
                "source": "深圳证券交易所",
                "version": "repealed-list",
            },
            self.admin,
            "127.0.0.1",
        )
        context = self.platform.rag.retrieve_keyword(
            "董事会决议公告 会议通知 表决结果",
            top_k=10,
            document_types=["regulations"],
            current_regulations_only=True,
            user_id=str(self.admin["id"]),
        )
        document_ids = {result.document_id for result in context.results}
        self.assertIn(current["document"]["id"], document_ids)
        self.assertNotIn(repealed["document"]["id"], document_ids)

    def test_admin_can_import_shared_drafting_materials_for_all_users(self) -> None:
        _, regular = self.auth.register(
            "drafting_reader", "DraftingPass123!", "起草资料使用者", "127.0.0.1",
        )
        imported = self.platform.import_knowledge(
            "汇报模板.txt",
            "一、公司基本情况\n二、事项进展\n三、后续安排".encode("utf-8"),
            {
                "document_type": "drafting_materials",
                "title": "公司汇报材料模板",
                "source": "证券部起草资料",
            },
            self.admin,
            "127.0.0.1",
        )
        self.assertEqual(imported["status"], "success")
        self.assertEqual(imported["document"]["document_type"], "drafting_materials")
        visible = self.platform.list_knowledge(regular)
        self.assertEqual(visible[0]["document_type"], "drafting_materials")
        opened_name, opened_content = self.platform.get_knowledge_document(
            imported["document"]["id"], regular, "127.0.0.1",
        )
        self.assertEqual(opened_name, "汇报模板.txt")
        self.assertIn("事项进展", opened_content.decode("utf-8"))

    def test_folder_import_metadata_preserves_relative_path(self) -> None:
        imported = self.platform.import_knowledge(
            "制度说明.txt",
            "文件夹导入测试\n第一条 文件内容。".encode("utf-8"),
            {
                "document_type": "internal_rules",
                "title": "制度目录/制度说明",
                "source": "本地文件夹导入",
                "relative_path": "制度目录\\子目录\\制度说明.txt",
            },
            self.admin,
            "127.0.0.1",
        )
        self.assertEqual(imported["status"], "success")
        self.assertEqual(imported["document"]["relativePath"], "制度目录/子目录/制度说明.txt")
        stored = self.database.fetch_one(
            "SELECT source_name FROM knowledge_documents WHERE id=?",
            (imported["document"]["id"],),
        )
        self.assertEqual(stored["source_name"], "制度目录/子目录/制度说明.txt")
        chunk = self.database.fetch_one(
            "SELECT metadata_json FROM chunks WHERE knowledge_document_id=?",
            (imported["document"]["id"],),
        )
        self.assertIn("制度目录/子目录/制度说明.txt", chunk["metadata_json"])

    def test_interrupted_knowledge_embedding_never_remains_active(self) -> None:
        original_embed_batch = self.platform.embedding_provider.embed_batch

        def interrupt_embedding(texts: object) -> list[list[float]]:
            raise KeyboardInterrupt("test interruption")

        self.platform.embedding_provider.embed_batch = interrupt_embedding  # type: ignore[method-assign]
        content = "中断测试公告\n本公告包含需要执行本地向量计算的独立内容。".encode("utf-8")
        with self.assertRaises(KeyboardInterrupt):
            self.platform.import_knowledge(
                "中断测试公告.txt",
                content,
                {"document_type": "announcements", "title": "中断测试公告", "source": "巨潮资讯网"},
                self.admin,
                "127.0.0.1",
            )
        interrupted = self.database.fetch_one(
            "SELECT id,status FROM knowledge_documents WHERE title='中断测试公告'"
        )
        self.assertEqual(interrupted["status"], "failed")
        self.platform.embedding_provider.embed_batch = original_embed_batch  # type: ignore[method-assign]
        rebuilt = self.platform.rebuild_knowledge(interrupted["id"], self.admin, "127.0.0.1")
        self.assertGreater(rebuilt["chunks"], 0)
        restored = self.database.fetch_one("SELECT status FROM knowledge_documents WHERE id=?", (interrupted["id"],))
        self.assertEqual(restored["status"], "active")

    def test_password_change_invalidates_existing_sessions(self) -> None:
        first_token, user = self.auth.register("password_user", "OriginalPass123!", "密码测试", "127.0.0.1")
        second_login = self.auth.login("password_user", "OriginalPass123!", "127.0.0.2")
        self.assertIsNotNone(second_login)
        second_token = second_login[0]

        new_token, refreshed = self.auth.change_password(
            user["id"], "OriginalPass123!", "UpdatedPass456!", "127.0.0.1"
        )
        self.assertEqual(refreshed["id"], user["id"])
        self.assertIsNone(self.auth.authenticate(first_token))
        self.assertIsNone(self.auth.authenticate(second_token))
        self.assertIsNotNone(self.auth.authenticate(new_token))
        self.assertIsNone(self.auth.login("password_user", "OriginalPass123!", "127.0.0.1"))
        self.assertIsNotNone(self.auth.login("password_user", "UpdatedPass456!", "127.0.0.1"))
        with self.assertRaisesRegex(ValueError, "当前密码错误"):
            self.auth.change_password(user["id"], "wrong", "AnotherPass789!", "127.0.0.1")

    def test_account_deletion_clears_private_data_anonymizes_audit_identity_and_allows_reregistration(self) -> None:
        token, user = self.auth.register("delete_me", "DeleteMe123!", "待注销用户", "127.0.0.1")
        self.assertTrue(self.auth.verify_current_password(user["id"], "DeleteMe123!"))
        self.assertFalse(self.auth.verify_current_password(user["id"], "wrong"))
        task = self.platform.create_task(
            "注销资料.txt",
            "注销账号测试资料\n本事项金额1000万元。".encode("utf-8"),
            "注销清理测试", "重大事项披露审查", "基础包", user, "127.0.0.1",
        )
        document = self.database.fetch_one(
            "SELECT file_path FROM documents WHERE id=?", (task["documentId"],)
        )
        assert document is not None
        file_path = Path(document["file_path"])
        self.platform.chat(task["taskId"], "查询当前任务依据。", None, user, "127.0.0.1")
        self.assertTrue(file_path.is_file())

        deleted = self.platform.delete_account(user["id"], self.admin, "127.0.0.1")
        self.assertEqual(deleted["status"], "deleted")
        self.assertEqual(deleted["deleted"]["tasks"], 1)
        self.assertFalse(file_path.exists())
        self.assertIsNone(self.auth.authenticate(token))
        deleted_user = self.database.fetch_one(
            "SELECT username,display_name,active,deleted_at,password_hash FROM users WHERE id=?", (user["id"],)
        )
        self.assertTrue(deleted_user["username"].startswith("deleted_"))
        self.assertEqual(deleted_user["display_name"], "已注销账号")
        self.assertEqual(deleted_user["active"], 0)
        self.assertTrue(deleted_user["deleted_at"])
        self.assertEqual(deleted_user["password_hash"], "account_deleted")
        self.assertIsNone(self.database.fetch_one("SELECT id FROM analysis_tasks WHERE owner_user_id=?", (user["id"],)))
        self.assertIsNone(self.database.fetch_one("SELECT id FROM documents WHERE owner_user_id=?", (user["id"],)))
        self.assertIsNone(self.database.fetch_one("SELECT id FROM conversations WHERE user_id=?", (user["id"],)))
        self.assertIsNone(self.database.fetch_one("SELECT id FROM memory_summaries WHERE user_id=?", (user["id"],)))
        self.assertIsNotNone(self.database.fetch_one(
            "SELECT id FROM audit_logs WHERE action='account_delete' AND resource_id=?", (user["id"],)
        ))

        _, replacement = self.auth.register("delete_me", "Replacement123!", "重新注册用户", "127.0.0.1")
        self.assertNotEqual(replacement["id"], user["id"])

    def test_last_active_admin_cannot_be_deleted(self) -> None:
        with self.assertRaisesRegex(ValueError, "至少保留一个"):
            self.platform.delete_account(
                self.admin["id"], self.admin, "127.0.0.1", self_service=True,
            )

    def test_task_delete_removes_document_vectors_memory_and_file(self) -> None:
        content = "删除测试公告\n关联交易金额1000万元，应核对审批和披露要求。".encode("utf-8")
        first = self.platform.create_task(
            "删除测试.txt", content, "待删除任务", "关联交易审查", "基础包", self.admin, "127.0.0.1"
        )
        second = self.platform.create_task(
            "删除测试.txt", content, "保留任务", "关联交易审查", "基础包", self.admin, "127.0.0.1"
        )
        task_id = first["taskId"]
        document_id = first["documentId"]
        document = self.database.fetch_one("SELECT file_path FROM documents WHERE id=?", (document_id,))
        assert document is not None
        shared_path = Path(document["file_path"])
        self.assertTrue(shared_path.is_file())
        response = self.platform.chat(task_id, "这个事项需要什么审批？", None, self.admin, "127.0.0.1")
        self.assertTrue(response["conversation_id"])
        self.assertIsNotNone(self.database.fetch_one("SELECT id FROM memory_summaries WHERE task_id=?", (task_id,)))

        deleted = self.platform.delete_task(task_id, self.admin, "127.0.0.1")
        self.assertEqual(deleted["status"], "deleted")
        self.assertGreater(deleted["deleted"]["chunks"], 0)
        self.assertGreater(deleted["deleted"]["messages"], 0)
        self.assertGreater(deleted["deleted"]["memories"], 0)
        for table, column in (
            ("analysis_tasks", "id"), ("documents", "id"), ("risks", "task_id"),
            ("conversations", "task_id"), ("memory_summaries", "task_id"), ("usage_records", "task_id"),
        ):
            self.assertIsNone(self.database.fetch_one(f"SELECT 1 FROM {table} WHERE {column}=?", (task_id if column != "id" or table == "analysis_tasks" else document_id,)))
        self.assertIsNone(self.database.fetch_one("SELECT id FROM chunks WHERE document_id=?", (document_id,)))
        self.assertTrue(shared_path.is_file(), "另一任务仍引用相同文件时不能删除磁盘文件")

        self.platform.delete_task(second["taskId"], self.admin, "127.0.0.1")
        self.assertFalse(shared_path.exists())
        audit = self.database.fetch_one(
            "SELECT after_json FROM audit_logs WHERE action='analysis_task_delete' AND resource_id=?", (task_id,)
        )
        self.assertIsNotNone(audit)

    def test_batch_task_delete_validates_all_ids_before_cascading_cleanup(self) -> None:
        content = "批量删除测试公告\n需要核对交易审批和信息披露。".encode("utf-8")
        first = self.platform.create_task(
            "批量删除一.txt", content, "批量删除任务一", "重大事项披露审查", "基础包", self.admin, "127.0.0.1"
        )
        second = self.platform.create_task(
            "批量删除二.txt", content, "批量删除任务二", "重大事项披露审查", "基础包", self.admin, "127.0.0.1"
        )
        ids = [first["taskId"], second["taskId"]]
        with self.assertRaisesRegex(ValueError, "部分文档审查已不存在"):
            self.platform.delete_tasks([ids[0], "missing-task-id"], self.admin, "127.0.0.1")
        self.assertEqual(len(self.platform.list_tasks(self.admin)), 2)

        deleted = self.platform.delete_tasks(ids, self.admin, "127.0.0.1")
        self.assertEqual(deleted["status"], "deleted")
        self.assertEqual(deleted["taskCount"], 2)
        self.assertEqual(set(deleted["taskIds"]), set(ids))
        self.assertEqual(self.platform.list_tasks(self.admin), [])

    def test_scanned_pdf_uses_page_level_local_ocr(self) -> None:
        try:
            from pypdf import PdfWriter
        except ImportError:
            self.skipTest("pypdf is not installed")

        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.write(output)

        class FakeOCRService:
            engine_name = "fake-local-ocr"

            @staticmethod
            def recognize_pdf_pages(content: bytes, pages: list[int]) -> list[OCRPage]:
                self.assertTrue(content)
                self.assertEqual(pages, [1])
                return [OCRPage(1, "创业黑马董事会决议\n第一条 本事项提交审议。")]

        parsed = parse_document("扫描决议.pdf", output.getvalue(), FakeOCRService())
        self.assertTrue(parsed.ocr_used)
        self.assertEqual(parsed.ocr_pages, (1,))
        self.assertEqual(parsed.ocr_engine, "fake-local-ocr")
        self.assertIn("创业黑马董事会决议", parsed.paragraphs)
        self.assertEqual(normalize_ocr_text("创 业 黑 马 2 0 2 6 年"), "创业黑马2026年")

    def test_local_ocr_allows_an_intentionally_blank_page(self) -> None:
        service = object.__new__(LocalOCRService)
        service.max_pages = 2
        service.status = lambda: {"available": True}  # type: ignore[method-assign]
        service._render_page = lambda pdf, root, page: root / f"page-{page}.png"  # type: ignore[method-assign]
        service._recognize_image = lambda image: ""  # type: ignore[method-assign]
        results = service.recognize_pdf_pages(b"%PDF test", [1])
        self.assertEqual(results, [OCRPage(1, "")])

    def test_private_workspace_isolation_for_tasks_usage_and_audit(self) -> None:
        _, user_a = self.auth.register("account_a", "AccountA123!", "账号甲", "127.0.0.1")
        _, user_b = self.auth.register("account_b", "AccountB123!", "账号乙", "127.0.0.1")

        task_a = self.platform.create_task(
            "甲公告.txt", "甲账户公告\n合同金额1000万元。".encode("utf-8"),
            "甲任务", "重大事项披露审查", "基础包", user_a, "127.0.0.1",
        )
        task_b = self.platform.create_task(
            "乙公告.txt", "乙账户公告\n合同金额2000万元。".encode("utf-8"),
            "乙任务", "重大事项披露审查", "基础包", user_b, "127.0.0.1",
        )
        self.assertEqual([item["id"] for item in self.platform.list_tasks(user_a)], [task_a["taskId"]])
        self.assertEqual([item["id"] for item in self.platform.list_tasks(user_b)], [task_b["taskId"]])
        for viewer, hidden_task in ((user_a, task_b), (user_b, task_a), (self.admin, task_a)):
            with self.assertRaises(PermissionError):
                self.platform.get_task_result(hidden_task["taskId"], user=viewer)

        self.platform._record_usage(user_a["id"], task_a["taskId"], "mock", "mock", "review", 100, 20)
        self.platform._record_usage(user_b["id"], task_b["taskId"], "mock", "mock", "review", 300, 60)
        self.database.audit(user_a["id"], "private_a", "task", task_a["taskId"], "127.0.0.1")
        self.database.audit(user_b["id"], "private_b", "task", task_b["taskId"], "127.0.0.1")
        self.assertEqual(self.platform.usage_summary(user_a)["records"][0]["input_tokens"], 100)
        self.assertEqual(self.platform.usage_summary(user_b)["records"][0]["input_tokens"], 300)
        self.assertFalse(any(log["action"] == "private_b" for log in self.platform.audit_logs(user_a)))
        self.assertFalse(any(log["action"] == "private_a" for log in self.platform.audit_logs(user_b)))

    def test_company_profile_uniqueness_is_scoped_to_account(self) -> None:
        _, user_a = self.auth.register("profile_a", "ProfileA123!", "画像甲", "127.0.0.1")
        _, user_b = self.auth.register("profile_b", "ProfileB123!", "画像乙", "127.0.0.1")
        now = "2026-08-20T00:00:00+00:00"
        values = []
        for owner, value in ((user_a["id"], "甲口径"), (user_b["id"], "乙口径")):
            values.append((
                f"fact-{owner}", owner, "创业黑马科技集团股份有限公司", "300688", "approval",
                "董事会门槛", value, "测试制度", "2026-01-01", "v1", now, 1, now, now,
            ))
        self.database.executemany(
            """INSERT INTO company_profiles (
                   id,owner_user_id,company,stock_code,fact_type,fact_key,fact_value,source,
                   effective_date,version,last_verified_at,confirmed,created_at,updated_at
               ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            values,
        )
        rows = self.database.fetch_all("SELECT owner_user_id,fact_value FROM company_profiles ORDER BY fact_value")
        self.assertEqual(len(rows), 2)

    def test_legacy_company_profile_schema_migrates_without_losing_facts(self) -> None:
        legacy_path = Path(self.temporary.name) / "runtime" / "legacy.db"
        legacy = Database(legacy_path)
        with legacy.connection() as connection:
            connection.executescript(
                """
                CREATE TABLE users (
                    id TEXT PRIMARY KEY, username TEXT NOT NULL UNIQUE, password_hash TEXT NOT NULL,
                    display_name TEXT NOT NULL, role TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE company_profiles (
                    id TEXT PRIMARY KEY, owner_user_id TEXT REFERENCES users(id), company TEXT NOT NULL,
                    stock_code TEXT, fact_type TEXT NOT NULL, fact_key TEXT NOT NULL, fact_value TEXT NOT NULL,
                    source TEXT NOT NULL, effective_date TEXT, version TEXT, last_verified_at TEXT,
                    confirmed INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
                    UNIQUE(company, fact_type, fact_key, version)
                );
                INSERT INTO users (id,username,password_hash,display_name,role,created_at,updated_at)
                VALUES ('legacy-user','legacy','hash','旧账号','user','2026-08-01','2026-08-01');
                INSERT INTO company_profiles (
                    id,owner_user_id,company,fact_type,fact_key,fact_value,source,version,created_at,updated_at
                ) VALUES ('legacy-fact','legacy-user','创业黑马','approval','门槛','旧口径','旧制度','v1','2026-08-01','2026-08-01');
                """
            )
        legacy.migrate()
        migrated = legacy.fetch_one("SELECT owner_user_id,fact_value FROM company_profiles WHERE id='legacy-fact'")
        self.assertEqual(migrated, {"owner_user_id": "legacy-user", "fact_value": "旧口径"})
        schema = legacy.fetch_one("SELECT sql FROM sqlite_master WHERE type='table' AND name='company_profiles'")
        self.assertNotIn("UNIQUE(company, fact_type, fact_key, version)", schema["sql"])

    def test_regular_user_cannot_open_another_users_task(self) -> None:
        now = "2026-08-19T00:00:00+00:00"
        other_id = "user-other"
        self.database.execute(
            "INSERT INTO users (id,username,password_hash,display_name,role,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
            (other_id, "other", hash_password("OtherUser123!"), "普通用户", "user", now, now),
        )
        draft = "公告草稿\n本次合同存在风险提示。".encode("utf-8")
        result = self.platform.create_task(
            "公告.txt", draft, "公告审查", "临时公告通用审查", "基础包", self.admin, "127.0.0.1"
        )
        with self.assertRaises(PermissionError):
            self.platform.get_task_result(
                result["taskId"],
                user={"id": other_id, "username": "other", "display_name": "普通用户", "role": "user"},
            )

    def test_cloud_mode_without_key_still_runs_deterministic_review(self) -> None:
        cloud_platform = PlatformService(replace(self.settings, ai_mode="cloud"), self.database)
        cloud_platform.initialize()
        self.assertFalse(cloud_platform.status()["aiConfigured"])
        result = cloud_platform.create_task(
            "合同.txt", "重大合同\n交易金额1000万元，需进一步确认审批和披露标准。".encode("utf-8"),
            "合同审查", "重大事项披露审查", "基础包", self.admin, "127.0.0.1",
        )
        self.assertFalse(result["aiUsed"])
        self.assertIn("尚未配置", result["aiMessage"])
        with self.assertRaises(AIServiceError):
            cloud_platform.chat(result["taskId"], "判断该合同风险", None, self.admin, "127.0.0.1")

    def test_csv_parser_preserves_rows_and_skips_exported_sheet_noise(self) -> None:
        content = (
            "分类,序号,名称,关联关系说明\n"
            "法务集团内公司,1,北京测试公司,全资子公司\n"
            "Sheet1,制度,部门,\n"
            "Sheet1,财务制度,财务部,\n"
            "年报其他关联方,序号,名称,关联关系说明\n"
            "年报其他关联方,1,张三,董高\n"
        ).encode("utf-8-sig")
        parsed = parse_document("关联方.csv", content)
        self.assertEqual(parsed.file_type, "CSV")
        self.assertEqual(len(parsed.blocks), 2)
        self.assertIn("北京测试公司", parsed.blocks[0].text)
        self.assertIn("CSV行2", parsed.blocks[0].text)
        self.assertIn("张三", parsed.blocks[1].text)
        self.assertFalse(any("财务制度" in block.text for block in parsed.blocks))
        chunks = SecuritiesChunker().chunk(
            "csv-document", parsed, "company_profile",
            {"file_name": "关联方.csv", "document_name": "关联方汇总表"},
        )
        self.assertEqual(chunks[0].metadata["csv_row_start"], 2)
        self.assertEqual(chunks[-1].metadata["csv_row_end"], 6)
        self.assertEqual(chunks[0].metadata["file_name"], "关联方.csv")
        self.assertIn("start_offset", chunks[0].metadata["original_text_position"])

    def test_embedding_selection_is_local_even_in_cloud_ai_mode(self) -> None:
        settings = replace(
            self.settings,
            ai_mode="cloud",
            aliyun_api_key="configured-only-for-chat",
            aliyun_base_url="https://example.invalid/v1",
            aliyun_model="chat-model",
            local_embedding_provider="mock",
        )
        provider = create_embedding_provider(settings)
        self.assertIsInstance(provider, MockEmbeddingProvider)
        self.assertEqual(provider.model_name(), "local-hash-ngram-v1")

    def test_cloud_analysis_sends_only_retrieved_context_not_full_document(self) -> None:
        settings = replace(
            self.settings,
            ai_mode="cloud",
            aliyun_api_key="configured-chat-key",
            aliyun_base_url="https://example.invalid/v1",
            aliyun_model="chat-model",
        )
        provider = AliyunProvider(settings)
        context = RAGContext("测试", [SearchResult(
            chunk_id="chunk-1", document_id="document-1", content="必要的本地检索片段",
            title="测试规则", source="本地", date=None, page=1, paragraph=0, section="第一条",
            document_type="regulations", score=0.9, source_url=None, version="v1", metadata={},
        )])
        captured: list[list[dict[str, str]]] = []

        def fake_request(messages: list[dict[str, str]], *, json_mode: bool = False) -> tuple[str, dict[str, int]]:
            captured.append(messages)
            return (
                '{"summary":"无新增风险","risk_score":0,"risk_level":"low","risks":[]}',
                {"input": 10, "output": 5},
            )

        provider._request = fake_request  # type: ignore[method-assign]
        secret_document = "完整文件不得外发-UNIQUE-MARKER"
        provider.analyze(secret_document, context)
        payload = json.dumps(captured, ensure_ascii=False)
        self.assertNotIn(secret_document, payload)
        self.assertIn("必要的本地检索片段", payload)

    def test_faiss_index_and_metadata_persist_and_reload(self) -> None:
        try:
            import faiss  # noqa: F401
        except ImportError:
            self.skipTest("faiss-cpu is not installed")
        faiss_root = Path(self.temporary.name) / "证券合规向量" / "faiss"
        settings = replace(
            self.settings,
            vector_backend="faiss",
            faiss_index_root=faiss_root,
            local_embedding_provider="mock",
            retrieval_mode="hybrid",
        )
        platform = PlatformService(settings, self.database)
        platform.initialize()
        peer_platform = PlatformService(settings, self.database)
        peer_platform.initialize()
        result = platform.import_knowledge(
            "本地向量制度.txt",
            "本地向量制度\n第一条 关联交易应当核对审议和披露标准。".encode("utf-8"),
            {"document_type": "regulations", "title": "本地向量制度", "source": "测试"},
            self.admin,
            "127.0.0.1",
        )
        self.assertGreater(result["document"]["chunks"], 0)
        self.assertIsInstance(platform.vector_store, FaissVectorStore)
        root = platform.vector_store.root
        self.assertTrue((root / "index.faiss").is_file())
        self.assertTrue((root / "metadata.json").is_file())
        self.assertTrue((root / "manifest.json").is_file())
        metadata = json.loads((root / "metadata.json").read_text(encoding="utf-8"))["records"]
        first = next(iter(metadata.values()))["metadata"]
        for key in ("file_name", "document_name", "original_text_position", "regulation_name", "document_updated_at"):
            self.assertIn(key, first)

        reloaded = FaissVectorStore(self.database, faiss_root, platform.embedding_provider.model_name())
        results = reloaded.search(
            platform.embedding_provider.embed_text("关联交易披露标准"),
            top_k=3,
            threshold=-1.0,
            filters={"user_id": self.admin["id"]},
        )
        self.assertEqual(results[0].title, "本地向量制度")
        self.assertEqual(results[0].metadata["file_name"], "本地向量制度.txt")

        second = peer_platform.import_knowledge(
            "外部同步公告.txt",
            "外部同步公告\n新增公告正文应当立即进入运行中服务的向量检索。".encode("utf-8"),
            {"document_type": "announcements", "title": "外部同步公告", "source": "巨潮资讯网"},
            self.admin,
            "127.0.0.1",
        )
        self.assertGreater(second["document"]["chunks"], 0)
        hot_results = platform.vector_store.search(
            platform.embedding_provider.embed_text("新增公告正文向量检索"),
            top_k=10,
            threshold=-1.0,
            filters={"user_id": self.admin["id"]},
        )
        self.assertIn("外部同步公告", [item.title for item in hot_results])
        final_metadata = json.loads((root / "metadata.json").read_text(encoding="utf-8"))["records"]
        self.assertEqual(
            {item["title"] for item in final_metadata.values()},
            {"本地向量制度", "外部同步公告"},
        )

        embed_batch_calls = 0
        original_embed_batch = platform.embedding_provider.embed_batch

        def count_embed_batch(texts: object) -> list[list[float]]:
            nonlocal embed_batch_calls
            embed_batch_calls += 1
            return original_embed_batch(texts)  # type: ignore[arg-type]

        platform.embedding_provider.embed_batch = count_embed_batch  # type: ignore[method-assign]
        queries = ["关联交易披露标准", "新增公告正文向量检索"]
        query_vectors = platform.rag.embed_queries(queries)
        batch_contexts = platform.rag.retrieve_many(
            queries,
            top_k=3,
            threshold=-1.0,
            document_types=["regulations", "announcements"],
            user_id=self.admin["id"],
            query_vectors=query_vectors,
        )
        self.assertEqual(embed_batch_calls, 1)
        self.assertEqual(set(batch_contexts), set(queries))
        self.assertTrue(all(context.results for context in batch_contexts.values()))

        def unexpected_embedding(texts: object) -> list[list[float]]:
            raise AssertionError("content_hash cache should avoid repeated embedding")

        platform.embedding_provider.embed_batch = unexpected_embedding  # type: ignore[method-assign]
        rebuilt = platform.rebuild_knowledge(result["document"]["id"], self.admin, "127.0.0.1")
        self.assertGreater(rebuilt["chunks"], 0)

    def test_faiss_recovers_corrupt_files_from_sqlite_vectors(self) -> None:
        try:
            import faiss  # noqa: F401
        except ImportError:
            self.skipTest("faiss-cpu is not installed")
        faiss_root = Path(self.temporary.name) / "faiss-recovery"
        settings = replace(
            self.settings,
            vector_backend="faiss",
            faiss_index_root=faiss_root,
            local_embedding_provider="mock",
        )
        platform = PlatformService(settings, self.database)
        platform.initialize()
        imported = platform.import_knowledge(
            "恢复测试规则.txt",
            "第一条 FAISS 损坏后应从 SQLite ready 向量自动恢复。".encode("utf-8"),
            {"document_type": "regulations", "title": "恢复测试规则", "source": "测试"},
            self.admin,
            "127.0.0.1",
        )
        self.assertGreater(imported["document"]["chunks"], 0)
        store = platform.vector_store
        assert isinstance(store, FaissVectorStore)
        store.index_path.write_bytes(b"\x00" * 64)
        store.metadata_path.write_bytes(b"\x00" * 64)
        store.manifest_path.write_bytes(b"\x00" * 64)

        recovered = FaissVectorStore(self.database, faiss_root, platform.embedding_provider.model_name())
        self.assertGreater(recovered.index_count, 0)
        results = recovered.search(
            platform.embedding_provider.embed_text("SQLite 向量恢复"),
            top_k=3,
            threshold=-1.0,
            filters={"user_id": self.admin["id"]},
        )
        self.assertEqual(results[0].title, "恢复测试规则")
        self.assertNotEqual(recovered.manifest_path.read_bytes()[:1], b"\x00")

    def test_faiss_recovers_when_sqlite_has_newer_ready_vectors(self) -> None:
        try:
            import faiss  # noqa: F401
        except ImportError:
            self.skipTest("faiss-cpu is not installed")
        faiss_root = Path(self.temporary.name) / "faiss-count-recovery"
        settings = replace(
            self.settings,
            vector_backend="faiss",
            faiss_index_root=faiss_root,
            local_embedding_provider="mock",
        )
        platform = PlatformService(settings, self.database)
        platform.initialize()
        platform.import_knowledge(
            "数量恢复规则.txt",
            "第一条 批量中断后应补齐 SQLite 已保存的向量。".encode("utf-8"),
            {"document_type": "regulations", "title": "数量恢复规则", "source": "测试"},
            self.admin,
            "127.0.0.1",
        )
        store = platform.vector_store
        assert isinstance(store, FaissVectorStore)
        persisted_count = store.index_count
        self.database.execute(
            """INSERT INTO chunks (
                   id,document_id,knowledge_document_id,content,page,section,paragraph,chunk_index,
                   content_hash,embedding_status,embedding_model,vector_json,metadata_json,created_at,updated_at
               ) SELECT ?,document_id,knowledge_document_id,content || ?,page,section,paragraph,chunk_index + 1,
                        content_hash || ?,embedding_status,embedding_model,vector_json,metadata_json,created_at,updated_at
                   FROM chunks WHERE embedding_model=? LIMIT 1""",
            ("sqlite-only-ready-vector", " SQLite 追加向量", "-sqlite-only", platform.embedding_provider.model_name()),
        )

        recovered = FaissVectorStore(self.database, faiss_root, platform.embedding_provider.model_name())
        self.assertEqual(recovered.index_count, persisted_count + 1)
        self.assertIsNotNone(recovered.get_by_chunk("sqlite-only-ready-vector"))

    def test_faiss_batch_updates_persist_once(self) -> None:
        try:
            import faiss  # noqa: F401
        except ImportError:
            self.skipTest("faiss-cpu is not installed")
        settings = replace(
            self.settings,
            vector_backend="faiss",
            faiss_index_root=Path(self.temporary.name) / "faiss-batch",
            local_embedding_provider="mock",
        )
        platform = PlatformService(settings, self.database)
        platform.initialize()
        store = platform.vector_store
        assert isinstance(store, FaissVectorStore)
        persist_calls = 0
        original_persist = store._persist

        def count_persist() -> None:
            nonlocal persist_calls
            persist_calls += 1
            original_persist()

        store._persist = count_persist  # type: ignore[method-assign]
        first_document_id = ""
        with store.batch_updates():
            for index in range(2):
                imported = platform.import_knowledge(
                    f"批量规则{index}.txt", f"第{index + 1}条 批量向量正文。".encode("utf-8"),
                    {"document_type": "regulations", "title": f"批量规则{index}", "source": "测试"},
                    self.admin,
                    "127.0.0.1",
                )
                if index == 0:
                    first_document_id = imported["document"]["id"]
            platform.rebuild_knowledge(first_document_id, self.admin, "127.0.0.1")
        self.assertEqual(persist_calls, 1)
        self.assertEqual(store.index_count, 2)


if __name__ == "__main__":
    unittest.main()
